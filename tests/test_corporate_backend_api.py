from datetime import date, timedelta
from io import BytesIO
import base64
import hashlib
import hmac
import json
import logging
import os

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.env_loader import load_travel_ai_env
from app.internal_logger import INTERNAL_LOGGER_NAME
from app.main import app
from app.security import create_access_token


def _client(tmp_path, monkeypatch):
    monkeypatch.setenv("TRAVEL_AI_DB_PATH", str(tmp_path / "travel_ai.db"))
    monkeypatch.setenv("TRAVEL_AI_TOKEN_SECRET", "test-token-secret")
    for name in (
        "DEEPSEEK_API_KEY",
        "OPENROUTER_API_KEY",
        "DUFFEL_API_TOKEN",
        "BOOKING_COM_TOKEN",
        "BOOKING_COM_AFFILIATE_ID",
        "RAPIDAPI_BOOKING_KEY",
        "RESEND_API_KEY",
        "RESEND_FROM_EMAIL",
        "RESEND_WEBHOOK_SECRET",
        "EMAIL_FROM",
        "MAIL_FROM",
        "RESEND_FROM",
        "RESEND_WEBHOOK_SIGNING_SECRET",
        "RESEND_SIGNING_SECRET",
    ):
        monkeypatch.delenv(name, raising=False)
    return TestClient(app)


def _headers(email="demo.user@unipro.com", purpose="manage corporate travel request"):
    token, _ = create_access_token(email)
    return {"Authorization": f"Bearer {token}", "X-Travel-Purpose": purpose}


def test_env_loader_reads_resend_values_from_env_file(tmp_path, monkeypatch):
    for name in ("RESEND_API_KEY", "RESEND_FROM_EMAIL", "RESEND_WEBHOOK_SECRET"):
        monkeypatch.delenv(name, raising=False)
    env_file = tmp_path / ".env"
    env_file.write_text(
        "\n".join([
            "RESEND_API_KEY=re_test_env_loader",
            "EMAIL_FROM=Travel Ops <travel@example.com>",
            "RESEND_WEBHOOK_SIGNING_SECRET=whsec_test_env_loader",
        ]),
        encoding="utf-8",
    )

    load_travel_ai_env([env_file])

    assert os.getenv("RESEND_API_KEY") == "re_test_env_loader"
    assert os.getenv("RESEND_FROM_EMAIL") == "Travel Ops <travel@example.com>"
    assert os.getenv("RESEND_WEBHOOK_SECRET") == "whsec_test_env_loader"
    for name in ("RESEND_API_KEY", "RESEND_FROM_EMAIL", "RESEND_WEBHOOK_SECRET", "EMAIL_FROM", "RESEND_WEBHOOK_SIGNING_SECRET"):
        monkeypatch.delenv(name, raising=False)


def _traveler_profile_payload(**overrides):
    payload = {
        "id": "traveler_ananya_shah",
        "name": "Ananya Shah",
        "email": "ananya.shah@orbitex.example",
        "company": "Orbitex",
        "department": None,
        "vip_level": None,
        "status": "Compliant",
        "location": "Mumbai home office",
        "seat_preference": "Window seat",
        "meal_preference": "Vegetarian meal",
        "hotel_preference": "Bandra Kurla Complex",
        "policy_notes": [],
        "loyalty_programs": [
            {"provider": "Air India Flying Returns", "tier": "Silver", "account_ref": "On file"}
        ],
        "documents": [
            {"document_type": "passport", "label": "Passport", "status": "Needs Review", "redacted_value": "Submitted by form"}
        ],
        "recent_trips": [],
        "created_at": "2026-05-24T00:00:00+00:00",
        "updated_at": "2026-05-24T00:00:00+00:00",
    }
    payload.update(overrides)
    return payload


def _corporate_payload(total_budget=2500, passport_expiry=None, visa_expiry=None):
    depart = date.today() + timedelta(days=45)
    ret = depart + timedelta(days=4)
    return {
        "traveller_details": {
            "traveler_name": "Anika Rao",
            "traveler_email": "anika.rao@unipro.com",
            "phone": "+91 98765 43210",
            "employee_id": "E-101",
            "employee_level": "manager",
            "department": "sales",
            "nationality": "India",
            "passport_expiry": str(passport_expiry or ret + timedelta(days=220)),
            "visa_status": "Valid visa on file",
            "visa_expiry": str(visa_expiry or ret + timedelta(days=30)),
            "medical_notes": "none",
            "accessibility_notes": "none",
        },
        "company_details": {
            "company_name": "Unipro",
            "cost_center": "SALES-42",
            "approving_manager": "Travel Manager",
            "approval_manager_email": "manager@unipro.com",
            "policy_tier": "standard",
        },
        "travel_details": {
            "origin": "Hyderabad",
            "destination": "Johannesburg",
            "destination_country": "South Africa",
            "depart_date": str(depart),
            "return_date": str(ret),
            "trip_purpose": "client workshops",
            "meeting_location": "Sandton client office",
            "flexible_dates": True,
            "travelers": 1,
            "cabin": "economy",
        },
        "preferences": {
            "preferred_airline": "Qatar Airways",
            "flight_preference": "short layover",
            "hotel_preference": "business hotel",
            "preferred_hotel_area": "Sandton",
            "hotel_star_rating": "4",
            "past_hotel_preference": "Garden Court",
            "airport_transfer_needed": True,
        },
        "budgets": {"total_budget": total_budget, "currency": "USD", "extra_baggage_notes": "one sample kit"},
        "special_requests": ["late check-in", "quiet room"],
    }


def _workbook_bytes():
    depart = date.today() + timedelta(days=45)
    ret = depart + timedelta(days=4)
    wb = Workbook()
    ws = wb.active
    ws.title = "Travel Requests"
    ws.append(
        [
            "traveler_name",
            "traveler_email",
            "phone",
            "employee_id",
            "employee_level",
            "nationality",
            "origin",
            "destination",
            "destination_country",
            "depart_date",
            "return_date",
            "trip_purpose",
            "meeting_location",
            "flexible_dates",
            "cabin",
            "total_budget",
            "passport_expiry",
            "visa_status",
            "visa_expiry",
            "approval_manager_email",
            "flight_preference",
            "preferred_hotel_area",
            "hotel_star_rating",
            "past_hotel_preference",
            "airport_transfer_needed",
            "medical_notes",
            "accessibility_notes",
            "extra_baggage_notes",
        ]
    )
    ws.append(
        [
            "Anika Rao",
            "anika.rao@unipro.com",
            "+91 98765 43210",
            "E-101",
            "manager",
            "India",
            "Hyderabad",
            "Johannesburg",
            "South Africa",
            depart,
            ret,
            "client workshops",
            "Sandton client office",
            "yes",
            "economy",
            1900,
            ret + timedelta(days=220),
            "Valid visa on file",
            ret + timedelta(days=30),
            "manager@unipro.com",
            "short layover",
            "Sandton",
            "4",
            "Garden Court",
            "yes",
            "none",
            "none",
            "one sample kit",
        ]
    )
    policy = wb.create_sheet("Company Policy")
    policy.append(["allowed_cabins", "max_budget"])
    policy.append(["economy,premium_economy", 1950])
    employees = wb.create_sheet("Employee Details")
    employees.append(["employee_name", "employee_email", "company_name", "department", "employee_level", "passport_number", "passport_expiry", "visa_status", "seat_preference", "meal_preference", "loyalty_airline"])
    employees.append(["Anika Rao", "anika.rao@unipro.com", "Unipro", "Sales", "manager", "Z1234567", ret + timedelta(days=220), "Valid visa on file", "Aisle", "Vegetarian", "Qatar Airways"])
    history = wb.create_sheet("Traveller History")
    history.append(["traveler_email", "preferred_airline", "notes"])
    history.append(["anika.rao@unipro.com", "Qatar Airways", "prefers aisle"])
    visa = wb.create_sheet("Visa Rules")
    visa.append(["from_country", "destination_country", "visa_required"])
    visa.append(["India", "South Africa", "yes"])
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _revised_profile_workbook_bytes():
    depart = date.today() + timedelta(days=45)
    ret = depart + timedelta(days=4)
    wb = Workbook()
    ws = wb.active
    ws.title = "Traveler Profiles"
    ws.append(["Traveler ID", "Name", "Profile Type", "Department", "Base Location", "Home Airport", "Nationality", "Phone", "Email", "Approval Notes"])
    ws.append(["TRV-MGR-002", "Karthik Menon", "Manager", "Sales", "Hyderabad, India", "Hyderabad", "India", "+91 90000 10002", "karthik.menon@company.example", "Department head approval for international trips."])
    preferences = wb.create_sheet("Travel Preferences")
    preferences.append(["Traveler ID", "Profile Type", "Cabin Preference", "Layover Preference", "Seat Preference", "Hotel Tier Preference", "Room Preference", "Meal Preference", "Ground Transport"])
    preferences.append(["TRV-MGR-002", "Manager", "Economy / Premium Economy", "Direct preferred", "Aisle preferred", "4-star business hotel", "Standard king room", "Breakfast included", "Airport taxi"])
    flights = wb.create_sheet("Flight Preferences")
    flights.append(["Traveler ID", "Profile Type", "Preferred Airlines", "Layover Rule", "Timing Preference", "Cabin Rule"])
    flights.append(["TRV-MGR-002", "Manager", "Qatar Airways, Emirates", "Avoid layovers above 4 hours", "Morning arrival", "Economy or premium economy"])
    visas = wb.create_sheet("Visa Passport Records")
    visas.append(["Traveler ID", "Traveler Name", "Nationality", "Passport Number", "Passport Expiry", "Visa / Permit Type", "Country / Region", "Visa Expiry", "Status", "Entry Type"])
    visas.append(["TRV-MGR-002", "Karthik Menon", "India", "Z7654321", ret + timedelta(days=260), "South Africa Business Visa", "South Africa", ret + timedelta(days=45), "Active", "Multiple entry"])
    hotels = wb.create_sheet("Preferred Hotels")
    hotels.append(["City", "Hotel Name", "Hotel Type", "Best Fit Profile", "Area", "Estimated Nightly Range", "Business Comment"])
    hotels.append(["Johannesburg", "Garden Court Sandton City", "4-star business", "Manager", "Sandton", "USD 110-170", "Good balance of cost and location."])
    history = wb.create_sheet("Past Travel History")
    history.append(["Trip ID", "Traveler ID", "Traveler Name", "Profile Type", "Destination City", "Country", "Purpose", "Departure Date", "Return Date", "Airline", "Hotel Name", "Hotel Feedback / Comment", "Preference Inferred"])
    history.append(["PTH-101", "TRV-MGR-002", "Karthik Menon", "Manager", "Johannesburg", "South Africa", "Client renewal", depart - timedelta(days=120), ret - timedelta(days=120), "Qatar Airways", "Garden Court Sandton City", "Good invoice support.", "Repeat booking indicates high fit."])
    insights = wb.create_sheet("Hotel Insights")
    insights.append(["Hotel Name", "City", "Best Fit Profile", "Times Used", "Average Rating", "Recommendation", "Watch-outs"])
    insights.append(["Garden Court Sandton City", "Johannesburg", "Manager", 2, 4.3, "Recommended when location and policy fit.", "No major issue."])
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _policy_pdf_bytes():
    lines = [
        "Company: Company",
        "Policy tier: Standard",
        "Allowed cabins: economy,premium_economy",
        "Maximum trip budget: USD 1950",
        "Hotel nightly limit: USD 180",
        "Approval rule: Department Head and Finance approval for international travel and any exception above policy.",
        "Notes: Use approved hotels close to the meeting location and prefer refundable fares for client trips.",
    ]
    content_lines = ["BT", "/F1 16 Tf", "50 780 Td"]
    for line in lines:
        content_lines.append(f"({line}) Tj")
        content_lines.append("0 -22 Td")
    content_lines.append("ET")
    stream = "\n".join(content_lines).encode("utf-8")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode("ascii"))
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")
    xref_start = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_start}\n%%EOF\n".encode("ascii"))
    return bytes(pdf)


def _pdf_form_bytes():
    lines = [
        "Client Travel Request Form",
        "Traveller name: Priya Menon",
        "Traveller email: priya.menon@orbitex.com",
        "Origin city / airport: Bengaluru",
        "Destination city / airport: Berlin",
        "Departure date: 2026-07-08",
        "Return date: 2026-07-13",
        "Purpose / meeting location: Partner onboarding workshop",
        "Approved budget and currency: 2200 EUR",
        "Flight preference: Morning arrival",
        "Hotel area / star rating: Mitte office area",
        "Special requests: Late check-in",
    ]
    content_lines = ["BT", "/F1 16 Tf", "50 780 Td"]
    for line in lines:
        content_lines.append(f"({line}) Tj")
        content_lines.append("0 -22 Td")
    content_lines.append("ET")
    stream = "\n".join(content_lines).encode("utf-8")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode("ascii"))
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")
    xref_start = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_start}\n%%EOF\n".encode("ascii"))
    return bytes(pdf)


def _component_skip_pdf_form_bytes():
    lines = [
        "Client Travel Request Form",
        "Traveller name: Meera Iyer",
        "Traveller email: meera.iyer@orbitex.com",
        "Origin city / airport: Bengaluru",
        "Destination city / airport: Berlin",
        "Departure date: 2026-07-08",
        "Purpose / meeting location: Partner onboarding workshop",
        "Approved budget and currency: 180000 INR",
        "Include outbound flight: Yes",
        "Include return flight: No",
        "Include hotel: No",
        "Special requests: Flight only, no hotel booking",
    ]
    content_lines = ["BT", "/F1 16 Tf", "50 780 Td"]
    for line in lines:
        content_lines.append(f"({line}) Tj")
        content_lines.append("0 -22 Td")
    content_lines.append("ET")
    stream = "\n".join(content_lines).encode("utf-8")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode("ascii"))
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")
    xref_start = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_start}\n%%EOF\n".encode("ascii"))
    return bytes(pdf)


def test_manual_create_list_and_detail(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    created = client.post("/api/corporate/requests", json=_corporate_payload(), headers=_headers())
    assert created.status_code == 201
    request_id = created.json()["id"]
    assert created.json()["status"] == "New"
    assert created.json()["owner_id"] == "usr_demo"

    listed = client.get("/api/corporate/requests", headers=_headers(purpose="review corporate requests"))
    assert listed.status_code == 200
    assert [row["id"] for row in listed.json()] == [request_id]

    detail = client.get(f"/api/corporate/requests/{request_id}", headers=_headers(purpose="review corporate request detail"))
    assert detail.status_code == 200
    assert detail.json()["travel_details"]["destination"] == "Johannesburg"


def test_agent_deletes_own_corporate_request(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    created = client.post("/api/corporate/requests", json=_corporate_payload(), headers=_headers())
    request_id = created.json()["id"]

    deleted = client.delete(
        f"/api/corporate/requests/{request_id}",
        headers=_headers(purpose="delete incorrect corporate request"),
    )

    assert deleted.status_code == 204
    listed = client.get("/api/corporate/requests", headers=_headers(purpose="review corporate requests"))
    assert listed.json() == []
    detail = client.get(f"/api/corporate/requests/{request_id}", headers=_headers(purpose="review deleted request"))
    assert detail.status_code == 404


def test_agent_marks_critical_issue_for_recovery_queue(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post("/api/corporate/requests", json=_corporate_payload(), headers=_headers())
    request_id = created.json()["id"]

    response = client.post(
        f"/api/corporate/requests/{request_id}/critical-issue",
        json={
            "status": "Urgent",
            "issue": "Flight cancelled - book an alternative from the same origin and adjust hotel dates if needed.",
        },
        headers=_headers(purpose="update urgent travel issue"),
    )

    assert response.status_code == 200
    body = response.json()
    assert body["critical_issue_status"] == "Urgent"
    assert "Flight cancelled" in body["critical_issue"]
    listed = client.get("/api/corporate/requests", headers=_headers(purpose="review corporate requests")).json()
    assert listed[0]["critical_issue_status"] == "Urgent"


def test_excel_import_persists_requests_history_and_visa_rules_without_policy_pdf(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    response = client.post(
        "/api/corporate/upload-excel",
        files={"file": ("corporate_travel.xlsx", _workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_headers(purpose="import corporate travel workbook"),
    )

    assert response.status_code == 200
    body = response.json()
    assert body["request_count"] == 1
    assert body["policy_count"] == 0
    assert body["traveller_history_count"] == 2
    assert body["employee_profile_count"] == 1
    assert body["visa_rule_count"] == 1
    assert body["created_request_ids"]
    detail = client.get(
        f"/api/corporate/requests/{body['created_request_ids'][0]}",
        headers=_headers(purpose="review corporate request detail"),
    ).json()
    assert detail["traveller_details"]["phone"] == "+91 98765 43210"
    assert detail["company_details"]["approval_manager_email"] == "manager@unipro.com"
    assert detail["travel_details"]["meeting_location"] == "Sandton client office"
    assert detail["preferences"]["airport_transfer_needed"] is True
    assert detail["budgets"]["extra_baggage_notes"] == "one sample kit"
    travelers = client.get("/api/travelers", headers=_headers("admin.user@unipro.com", "review traveler roster"))
    assert any(traveler["email"] == "anika.rao@unipro.com" and traveler["seat_preference"] == "Aisle" for traveler in travelers.json())
    agent_travelers = client.get("/api/travelers", headers=_headers("demo.agent@unipro.com", "review traveler roster"))
    assert agent_travelers.status_code == 200
    assert any(traveler["email"] == "anika.rao@unipro.com" for traveler in agent_travelers.json())


def test_revised_profile_workbook_and_policy_pdf_feed_company_pipeline_and_plan_context(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    admin_headers = _headers("admin.user@unipro.com", "import manager company context")

    profile_response = client.post(
        "/api/corporate/upload-excel",
        files={"file": ("Corporate_Travel_Profile_Dataset_Revised.xlsx", _revised_profile_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=admin_headers,
    )

    assert profile_response.status_code == 200
    profile_body = profile_response.json()
    assert profile_body["request_count"] == 0
    assert profile_body["policy_count"] == 0
    assert profile_body["employee_profile_count"] == 1
    assert profile_body["visa_rule_count"] == 1

    policy_response = client.post(
        "/api/corporate/company-policy/upload-pdf",
        files={"file": ("company-policy.pdf", _policy_pdf_bytes(), "application/pdf")},
        headers={**admin_headers, "X-Company-Name": "Company"},
    )

    assert policy_response.status_code == 200
    assert policy_response.json()["policy_count"] == 1
    companies = client.get("/api/corporate/companies", headers=_headers("admin.user@unipro.com", "review company pipeline")).json()
    company = next(item for item in companies if item["company_name"] == "Company")
    assert company["traveler_list_status"] == "Updated"
    assert company["policy_status"] == "Uploaded"
    assert company["traveler_count"] == 1

    payload = _corporate_payload()
    payload["traveller_details"]["traveler_name"] = "Karthik Menon"
    payload["traveller_details"]["traveler_email"] = "karthik.menon@company.example"
    payload["traveller_details"].pop("passport_expiry")
    payload["traveller_details"].pop("visa_expiry")
    payload["traveller_details"].pop("visa_status")
    payload["company_details"]["company_name"] = "Company"
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    planned = client.post(f"/api/corporate/requests/{created.json()['id']}/plan", headers=_headers()).json()

    assert planned["traveller_details"]["passport_expiry"]
    assert planned["traveller_details"]["visa_expiry"]
    assert planned["generated_plan"]["missing_information"] == []
    assert any("Traveller history note" in note for note in planned["generated_plan"]["agent_notes"])


def test_trip_only_request_uses_uploaded_employee_roster_for_documents_and_preferences(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    admin_headers = _headers("admin.user@unipro.com", "import employee travel roster")

    profile_response = client.post(
        "/api/corporate/upload-excel",
        files={"file": ("Corporate_Travel_Profile_Dataset_Revised.xlsx", _revised_profile_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=admin_headers,
    )
    assert profile_response.status_code == 200

    depart = date.today() + timedelta(days=45)
    ret = depart + timedelta(days=4)
    trip_only_payload = {
        "traveller_details": {
            "traveler_name": "Karthik Menon",
            "traveler_email": "karthik.menon@company.example",
            "employee_id": "TRV-MGR-002",
        },
        "company_details": {"company_name": "Company"},
        "travel_details": {
            "origin": "Bengaluru",
            "destination": "Johannesburg",
            "destination_country": "South Africa",
            "depart_date": str(depart),
            "return_date": str(ret),
            "trip_purpose": "Client workshop",
            "meeting_location": "Sandton client office",
            "cabin": "economy",
        },
        "budgets": {"total_budget": 2500, "currency": "USD"},
        "special_requests": ["projector access"],
    }

    created = client.post("/api/corporate/requests", json=trip_only_payload, headers=_headers())
    planned = client.post(f"/api/corporate/requests/{created.json()['id']}/plan", headers=_headers()).json()

    assert planned["traveller_details"]["passport_number"] == "Z7654321"
    assert planned["traveller_details"]["visa_status"] == "Active - South Africa Business Visa"
    assert planned["preferences"]["hotel_preference"]
    assert planned["generated_plan"]["missing_information"] == []


def test_pdf_form_import_starts_automated_pipeline(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    response = client.post(
        "/api/corporate/requests/upload-excel",
        files={"file": ("client_travel_form.pdf", _pdf_form_bytes(), "application/pdf")},
        headers=_headers(purpose="import corporate travel pdf form"),
    )

    assert response.status_code == 200
    body = response.json()
    assert body["request_count"] == 1
    detail = client.get(
        f"/api/corporate/requests/{body['created_request_ids'][0]}",
        headers=_headers(purpose="review corporate request detail"),
    ).json()
    assert detail["status"] == "Processing"
    assert len(detail["generated_plan"]["travel_options"]) == 3
    assert detail["traveller_details"]["traveler_name"] == "Priya Menon"
    assert detail["travel_details"]["destination"] == "Berlin"
    assert detail["budgets"]["currency"] == "EUR"


def test_excel_template_download_contains_request_history_policy_and_visa_sheets(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    response = client.get(
        "/api/corporate/excel-template",
        headers=_headers(purpose="download corporate travel workbook template"),
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["Travel Requests", "Company Policy", "Employee Details", "Traveller History", "Visa Rules"]
    assert [cell.value for cell in workbook["Travel Requests"][1]][:6] == [
        "traveler_name",
        "traveler_email",
        "employee_id",
        "department",
        "nationality",
        "origin",
    ]
    assert "preferred_airline" in [cell.value for cell in workbook["Traveller History"][1]]
    assert "employee_email" in [cell.value for cell in workbook["Employee Details"][1]]
    request_headers = [cell.value for cell in workbook["Travel Requests"][1]]
    for field in [
        "phone",
        "approval_manager_email",
        "employee_level",
        "meeting_location",
        "flexible_dates",
        "flight_preference",
        "preferred_hotel_area",
        "hotel_star_rating",
        "past_hotel_preference",
        "airport_transfer_needed",
        "medical_notes",
        "accessibility_notes",
        "extra_baggage_notes",
    ]:
        assert field in request_headers


def test_client_request_form_pdf_download_is_available_for_mvp_intake(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    response = client.get(
        "/api/corporate/request-form.pdf",
        headers=_headers(purpose="download client travel request form"),
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == 'attachment; filename="client_travel_request_form.pdf"'
    assert response.content.startswith(b"%PDF-1.4")
    assert b"Client Travel Request Form" in response.content
    assert b"Passport expiry" not in response.content
    assert b"Nationality" not in response.content
    assert b"loyalty" in response.content.lower()
    assert b"Synthetic" not in response.content


def test_request_update_resolves_missing_information_without_changing_owner(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    payload = _corporate_payload()
    payload["traveller_details"]["traveler_email"] = None
    payload["traveller_details"]["nationality"] = None
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    request_id = created.json()["id"]
    planned = client.post(f"/api/corporate/requests/{request_id}/plan", headers=_headers()).json()
    assert "traveller_details.traveler_email" in planned["generated_plan"]["missing_information"]
    assert "traveller_details.nationality" in planned["generated_plan"]["missing_information"]

    planned["traveller_details"]["traveler_email"] = "anika.rao@unipro.com"
    updated = client.put(
        f"/api/corporate/requests/{request_id}/plan",
        json={
            "generated_plan": {
                **planned["generated_plan"],
                "missing_information": ["traveller_details.traveler_email"],
            },
            "traveller_details": {"nationality": "Indian"},
        },
        headers=_headers(purpose="resolve missing corporate request information"),
    )

    assert updated.status_code == 200
    assert updated.json()["id"] == request_id
    assert updated.json()["owner_id"] == "usr_demo"
    assert updated.json()["traveller_details"]["nationality"] == "Indian"
    assert "traveller_details.nationality" not in updated.json()["generated_plan"]["missing_information"]
    full_update = client.put(
        f"/api/corporate/requests/{request_id}",
        json={**updated.json(), "traveller_details": {**updated.json()["traveller_details"], "traveler_email": "anika.rao@unipro.com"}},
        headers=_headers(purpose="resolve missing corporate request information"),
    )
    assert full_update.status_code == 200
    replanned = client.post(f"/api/corporate/requests/{request_id}/plan", headers=_headers()).json()
    assert replanned["generated_plan"]["missing_information"] == []


def test_plan_update_persists_core_request_fields_from_agent_guide(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post("/api/corporate/requests", json=_corporate_payload(), headers=_headers())
    request_id = created.json()["id"]
    planned = client.post(f"/api/corporate/requests/{request_id}/plan", headers=_headers()).json()

    updated = client.put(
        f"/api/corporate/requests/{request_id}/plan",
        json={
            "generated_plan": planned["generated_plan"],
            "traveller_details": {"traveler_name": "Ananya Shah", "traveler_email": "ananya.shah@orbitex.example"},
            "company_details": {"company_name": "Orbitex"},
            "travel_details": {
                "origin": "Mumbai",
                "destination": "San Jose",
                "depart_date": "2026-10-26",
                "return_date": "2026-11-25",
                "trip_purpose": "Client meeting",
            },
            "preferences": {"hotel_preference": "Hotel near North First Street office"},
            "special_requests": ["Vegetarian meal"],
        },
        headers=_headers("demo.user@unipro.com", "save guide-edited corporate request fields"),
    )

    assert updated.status_code == 200
    body = updated.json()
    assert body["traveller_details"]["traveler_name"] == "Ananya Shah"
    assert body["traveller_details"]["traveler_email"] == "ananya.shah@orbitex.example"
    assert body["company_details"]["company_name"] == "Orbitex"
    assert body["travel_details"]["origin"] == "Mumbai"
    assert body["travel_details"]["destination"] == "San Jose"
    assert body["travel_details"]["depart_date"] == "2026-10-26"
    assert body["travel_details"]["return_date"] == "2026-11-25"
    assert body["travel_details"]["trip_purpose"] == "Client meeting"
    assert body["preferences"]["hotel_preference"] == "Hotel near North First Street office"
    assert body["special_requests"] == ["Vegetarian meal"]


def test_approval_status_tracking_is_request_data(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post("/api/corporate/requests", json=_corporate_payload(total_budget=1900), headers=_headers())
    request_id = created.json()["id"]

    planned = client.post(f"/api/corporate/requests/{request_id}/plan", headers=_headers()).json()
    assert planned["status"] == "Waiting for Approval"
    assert planned["approval_status"] == "Required"

    planned["approval_status"] = "Rejected"
    tracked = client.put(
        f"/api/corporate/requests/{request_id}",
        json=planned,
        headers=_headers(purpose="track corporate approval status"),
    )
    assert tracked.status_code == 200
    assert tracked.json()["approval_status"] == "Rejected"
    assert tracked.json()["status"] == "Waiting for Approval"


def test_finalize_is_blocked_without_generated_plan_and_agent_review(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post("/api/corporate/requests", json=_corporate_payload(), headers=_headers())
    request_id = created.json()["id"]

    no_plan = client.post(
        f"/api/corporate/requests/{request_id}/finalize",
        json={"agent_reviewed": True},
        headers=_headers(purpose="finalize approved corporate itinerary"),
    )
    assert no_plan.status_code == 400
    assert no_plan.json()["detail"] == "Final itinerary is not ready"

    client.post(f"/api/corporate/requests/{request_id}/plan", headers=_headers())
    no_review = client.post(
        f"/api/corporate/requests/{request_id}/finalize",
        json={"agent_reviewed": False},
        headers=_headers(purpose="finalize approved corporate itinerary"),
    )
    assert no_review.status_code == 400
    assert no_review.json()["detail"] == "Final itinerary is not ready"


def test_finalize_succeeds_when_agent_reviewed_and_approval_received(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post("/api/corporate/requests", json=_corporate_payload(total_budget=1900), headers=_headers())
    request_id = created.json()["id"]
    client.post(f"/api/corporate/requests/{request_id}/plan", headers=_headers())

    blocked = client.post(
        f"/api/corporate/requests/{request_id}/finalize",
        json={"agent_reviewed": True},
        headers=_headers(purpose="finalize approved corporate itinerary"),
    )
    assert blocked.status_code == 400
    assert blocked.json()["detail"] == "Final itinerary is not ready"

    finalized = client.post(
        f"/api/corporate/requests/{request_id}/finalize",
        json={"agent_reviewed": True, "approval_status": "Received"},
        headers=_headers(purpose="finalize approved corporate itinerary"),
    )
    assert finalized.status_code == 200
    assert finalized.json()["status"] == "Finalized"
    assert finalized.json()["approval_status"] == "Received"


def test_final_itinerary_export_is_blocked_until_request_is_finalized(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post("/api/corporate/requests", json=_corporate_payload(), headers=_headers())
    request_id = created.json()["id"]

    unplanned_export = client.get(
        f"/api/corporate/requests/{request_id}/export.xlsx",
        headers=_headers(purpose="download finalized corporate itinerary excel"),
    )
    assert unplanned_export.status_code == 400
    assert unplanned_export.json()["detail"] == "Final itinerary is not ready"

    client.post(f"/api/corporate/requests/{request_id}/plan", headers=_headers())
    planned_export = client.get(
        f"/api/corporate/requests/{request_id}/export.xlsx",
        headers=_headers(purpose="download finalized corporate itinerary excel"),
    )
    assert planned_export.status_code == 400
    assert planned_export.json()["detail"] == "Final itinerary is not ready"
    planned_pdf = client.get(
        f"/api/corporate/requests/{request_id}/export.pdf",
        headers=_headers(purpose="download finalized corporate itinerary pdf"),
    )
    assert planned_pdf.status_code == 400
    assert planned_pdf.json()["detail"] == "Final itinerary is not ready"


def test_blocked_finalization_and_export_attempts_are_audited(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post("/api/corporate/requests", json=_corporate_payload(), headers=_headers())
    request_id = created.json()["id"]

    finalize = client.post(
        f"/api/corporate/requests/{request_id}/finalize",
        json={"agent_reviewed": True},
        headers=_headers(purpose="finalize approved corporate itinerary"),
    )
    export = client.get(
        f"/api/corporate/requests/{request_id}/export.xlsx",
        headers=_headers(purpose="download finalized corporate itinerary excel"),
    )

    assert finalize.status_code == 400
    assert export.status_code == 400
    audit = client.get(
        "/api/admin/audit",
        headers=_headers("admin.user@unipro.com", "review security audit events"),
    )
    assert audit.status_code == 200
    events = audit.json()
    assert any(event["event_type"] == "corporate.finalize.blocked" and event["decision"] == "deny" for event in events)
    assert any(event["event_type"] == "corporate.request_export.blocked" and event["decision"] == "deny" for event in events)
    assert "anika.rao@unipro.com" not in audit.text


def test_final_itinerary_notification_is_blocked_until_request_is_finalized(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post("/api/corporate/requests", json=_corporate_payload(), headers=_headers())
    request_id = created.json()["id"]

    blocked = client.post(
        f"/api/corporate/requests/{request_id}/notifications",
        json={"kind": "final_itinerary", "to": ["traveler@example.com"], "attach_itinerary": True},
        headers=_headers(purpose="send final itinerary notification"),
    )

    assert blocked.status_code == 400
    assert blocked.json()["detail"] == "Final itinerary is not ready"
    audit = client.get(
        "/api/admin/audit",
        headers=_headers("admin.user@unipro.com", "review travel audit events"),
    )
    assert any(event["event_type"] == "corporate.notification.blocked" and event["decision"] == "deny" for event in audit.json())


def test_generate_plan_for_hyderabad_to_johannesburg_checks_documents_budget_and_policy(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    imported = client.post(
        "/api/corporate/upload-excel",
        files={"file": ("corporate_travel.xlsx", _workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_headers(purpose="import corporate travel workbook"),
    )
    request_id = imported.json()["created_request_ids"][0]

    planned = client.post(
        f"/api/corporate/requests/{request_id}/plan",
        headers=_headers(purpose="generate compliant corporate travel plan"),
    )

    assert planned.status_code == 200
    body = planned.json()
    plan = body["generated_plan"]
    assert body["status"] == "Waiting for Approval"
    assert plan["request_summary"].startswith("Anika Rao needs client workshops travel from Hyderabad to Johannesburg")
    assert plan["missing_information"] == []
    assert plan["travel_readiness"]["passport_status"] == "Ready"
    assert plan["travel_readiness"]["visa_status"] == "Ready"
    assert plan["budget_policy_check"]["budget_status"] == "Needs Approval"
    assert plan["budget_policy_check"]["policy_status"] == "Needs Approval"
    assert plan["budget_policy_check"]["approval_required"] is True
    assert [option["option_name"] for option in plan["travel_options"]] == [
        "Best within budget",
        "Fastest route",
        "Comfort-focused option",
    ]
    assert len(plan["hotel_offers"]) == 3
    assert plan["hotel_offers"][0]["name"]
    assert plan["hotel_offers"][0]["image_url"].startswith("/travel-media/")
    assert plan["selected_hotel_offer_id"] == plan["hotel_offers"][0]["id"]
    assert "Traveller history note" in " ".join(plan["agent_notes"])


def test_generate_plan_attaches_selectable_duffel_flight_offers(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setenv("DUFFEL_API_TOKEN", "test-duffel-token")

    import httpx

    class FakeResponse:
        def raise_for_status(self):
            return None

        def json(self):
            return {
                "data": {
                    "offers": [
                        {
                            "id": "off_live_123",
                            "total_amount": "1200",
                            "total_currency": "USD",
                            "expires_at": "2026-05-22T12:00:00Z",
                            "owner": {"name": "Qatar Airways", "iata_code": "QR"},
                            "slices": [
                                {
                                    "segments": [
                                        {
                                            "origin": {"iata_code": "HYD"},
                                            "destination": {"iata_code": "DOH"},
                                            "marketing_carrier": {"name": "Qatar Airways"},
                                            "departing_at": "2026-06-10T03:30:00",
                                            "arriving_at": "2026-06-10T05:30:00",
                                        },
                                        {
                                            "origin": {"iata_code": "DOH"},
                                            "destination": {"iata_code": "JNB"},
                                            "marketing_carrier": {"name": "Qatar Airways"},
                                            "departing_at": "2026-06-10T08:00:00",
                                            "arriving_at": "2026-06-10T15:00:00",
                                        },
                                    ]
                                }
                            ],
                        }
                    ]
                }
            }

    captured = {}

    def fake_post(url, params, headers, json, timeout):
        captured["url"] = url
        captured["params"] = params
        captured["headers"] = headers
        captured["json"] = json
        return FakeResponse()

    monkeypatch.setattr(httpx, "post", fake_post)
    created = client.post("/api/corporate/requests", json=_corporate_payload(total_budget=2500), headers=_headers())
    planned = client.post(f"/api/corporate/requests/{created.json()['id']}/plan", headers=_headers())

    assert planned.status_code == 200
    plan = planned.json()["generated_plan"]
    assert captured["url"] == "https://api.duffel.com/air/offer_requests"
    assert captured["params"]["return_offers"] is True
    assert captured["json"]["data"]["slices"][0]["origin"] == "HYD"
    assert captured["json"]["data"]["slices"][0]["destination"] == "JNB"
    assert plan["flight_offers"][0]["id"] == "off_live_123"
    assert plan["flight_offers"][0]["source"] == "duffel"
    assert plan["selected_flight_offer_id"] == "off_live_123"
    assert plan["travel_options"][0]["flight_offer_id"] == "off_live_123"


def test_generate_plan_uses_duffel_and_booking_mcp_tools(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setenv("TRAVEL_AI_MCP_TOOLS_URL", "http://mcp.local/tools")

    import httpx

    class FakeResponse:
        def __init__(self, payload):
            self.payload = payload

        def raise_for_status(self):
            return None

        def json(self):
            return self.payload

    calls = []

    def fake_post(url, json, timeout):
        calls.append((url, json))
        if url.endswith("/search_flights"):
            text = json_module.dumps([
                {
                    "airline": "Qatar Airways",
                    "airline_iata": "QR",
                    "origin": "HYD",
                    "destination": "JNB",
                    "departure_at": "2026-06-10T03:30:00",
                    "arrival_at": "2026-06-10T15:00:00",
                    "return_departure_at": "2026-06-14T18:00:00",
                    "return_arrival_at": "2026-06-15T08:00:00",
                    "stops": 1,
                    "price_usd": 1180,
                    "currency": "USD",
                }
            ])
        elif url.endswith("/search_hotels"):
            text = json_module.dumps([
                {
                    "name": "Booking Live Sandton Hotel",
                    "address": "Sandton",
                    "star_rating": 4,
                    "price_per_night": 140,
                    "total_price": 560,
                    "currency": "USD",
                    "redirect_url": "https://booking.com/live-sandton",
                }
            ])
        else:
            text = "[]"
        return FakeResponse({"content": [{"type": "text", "text": text}]})

    json_module = json
    monkeypatch.setattr(httpx, "post", fake_post)

    created = client.post("/api/corporate/requests", json=_corporate_payload(total_budget=2400), headers=_headers())
    planned = client.post(f"/api/corporate/requests/{created.json()['id']}/plan", headers=_headers())

    assert planned.status_code == 200
    plan = planned.json()["generated_plan"]
    assert [call[0] for call in calls] == ["http://mcp.local/tools/search_flights", "http://mcp.local/tools/search_hotels"]
    assert calls[0][1]["origin"] == "HYD"
    assert calls[0][1]["destination"] == "JNB"
    assert plan["flight_offers"][0]["provider"] == "duffel-mcp/qr"
    assert plan["flight_offers"][0]["source"] == "duffel"
    assert plan["travel_options"][0]["flight_offer_id"] == plan["flight_offers"][0]["id"]
    assert plan["hotel_offers"][0]["provider"] == "booking.com-mcp"
    assert plan["hotel_offers"][0]["name"] == "Booking Live Sandton Hotel"
    assert plan["hotel_offers"][0]["image_url"].startswith("/travel-media/")
    assert plan["selected_hotel_offer_id"] == plan["hotel_offers"][0]["id"]
    assert "Booking Live Sandton Hotel via booking.com-mcp" in plan["travel_options"][0]["hotel_summary"]
    assert any("Live Duffel flight offers are attached" in note for note in plan["agent_notes"])
    assert any("Live Booking.com hotel offers informed" in note for note in plan["agent_notes"])


def test_generate_plan_uses_booking_demand_api_for_hotel_summary(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setenv("BOOKING_COM_TOKEN", "test-booking-token")
    monkeypatch.setenv("BOOKING_COM_AFFILIATE_ID", "12345")
    monkeypatch.setenv("BOOKING_COM_CITY_IDS", json.dumps({"Johannesburg": -1240260}))
    monkeypatch.setenv("TRAVEL_AI_MCP_TOOLS_URL", "http://mcp.local/tools")

    import httpx

    class FakeResponse:
        def __init__(self, payload):
            self.payload = payload

        def raise_for_status(self):
            return None

        def json(self):
            return self.payload

    calls = []

    def fake_post(url, *args, **kwargs):
        calls.append((url, kwargs))
        if url.endswith("/accommodations/search"):
            return FakeResponse(
                {
                    "data": [
                        {
                            "id": 9981,
                            "name": "Booking Demand Sandton Hotel",
                            "currency": "USD",
                            "price": {"book": "620"},
                            "products": [{"policies": {"cancellation": {"free_cancellation_until": "2026-06-01"}}}],
                        }
                    ]
                }
            )
        return FakeResponse({"content": [{"type": "text", "text": "[]"}]})

    monkeypatch.setattr(httpx, "post", fake_post)

    created = client.post("/api/corporate/requests", json=_corporate_payload(total_budget=2600), headers=_headers())
    planned = client.post(f"/api/corporate/requests/{created.json()['id']}/plan", headers=_headers())

    assert planned.status_code == 200
    plan = planned.json()["generated_plan"]
    booking_call = next(call for call in calls if call[0].endswith("/accommodations/search"))
    assert booking_call[1]["headers"]["X-Affiliate-Id"] == "12345"
    assert booking_call[1]["json"]["city"] == -1240260
    assert "Booking Demand Sandton Hotel via booking.com-demand-api" in plan["travel_options"][0]["hotel_summary"]
    assert plan["hotel_offers"][0]["provider"] == "booking.com-demand-api"
    assert plan["hotel_offers"][0]["name"] == "Booking Demand Sandton Hotel"
    assert plan["selected_hotel_offer_id"] == plan["hotel_offers"][0]["id"]
    assert any("booking.com-demand-api" in note for note in plan["agent_notes"])


def test_generate_plan_uses_openrouter_deepseek_for_structured_ai_draft_when_configured(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")
    monkeypatch.setenv("OPENROUTER_FLASH_MODEL", "openrouter/deepseek/deepseek-v4-flash")
    monkeypatch.setenv("OPENROUTER_PROVIDER_ORDER", "DeepSeek")

    import httpx

    class FakeResponse:
        def raise_for_status(self):
            return None

        def json(self):
            plan = {
                "request_summary": "DeepSeek structured summary for Anika Rao's Johannesburg trip.",
                "missing_information": ["model tried to change this"],
                "travel_readiness": {
                    "passport_status": "Needs Review",
                    "visa_status": "Needs Review",
                    "transit_warning": "Model transit note",
                    "document_notes": ["Model document note"],
                },
                "budget_policy_check": {
                    "budget_status": "Within Budget",
                    "policy_status": "Compliant",
                    "approval_required": False,
                    "approval_reason": "Model tried to change this",
                    "total_budget": 999999,
                    "estimated_cost": 1,
                },
                "travel_options": [
                    {
                        "option_name": "Best within budget",
                        "flight_summary": "DeepSeek low-cost flight wording.",
                        "hotel_summary": "DeepSeek low-cost hotel wording.",
                        "estimated_cost": 1,
                        "pros": ["clear customer wording"],
                        "cons": ["requires agent review"],
                        "policy_status": "Model changed",
                        "recommendation_reason": "DeepSeek recommendation.",
                    },
                    {
                        "option_name": "Fastest route",
                        "flight_summary": "DeepSeek fastest flight wording.",
                        "hotel_summary": "DeepSeek fastest hotel wording.",
                        "estimated_cost": 2,
                        "pros": ["shorter travel day"],
                        "cons": ["higher cost"],
                        "policy_status": "Model changed",
                        "recommendation_reason": "DeepSeek recommendation.",
                    },
                    {
                        "option_name": "Comfort-focused option",
                        "flight_summary": "DeepSeek comfort flight wording.",
                        "hotel_summary": "DeepSeek comfort hotel wording.",
                        "estimated_cost": 3,
                        "pros": ["better rest"],
                        "cons": ["approval likely"],
                        "policy_status": "Model changed",
                        "recommendation_reason": "DeepSeek recommendation.",
                    },
                ],
                "agent_note": "DeepSeek agent note.",
                "agent_notes": ["DeepSeek operational note."],
                "customer_message_draft": "DeepSeek customer message for passport Z1234567.",
                "customer_itinerary_draft": "DeepSeek itinerary draft for passport Z1234567.",
                "approval_status": "Plan Generated",
            }
            return {"choices": [{"message": {"content": json.dumps(plan)}}]}

    captured = {}

    def fake_post(url, headers, json, timeout):
        captured["url"] = url
        captured["headers"] = headers
        captured["json"] = json
        return FakeResponse()

    monkeypatch.setattr(httpx, "post", fake_post)

    payload = _corporate_payload(total_budget=1900)
    payload["traveller_details"]["passport_number"] = "Z1234567"
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    planned = client.post(f"/api/corporate/requests/{created.json()['id']}/plan", headers=_headers())

    assert planned.status_code == 200
    plan = planned.json()["generated_plan"]
    assert captured["url"] == "https://openrouter.ai/api/v1/chat/completions"
    assert captured["headers"]["Authorization"] == "Bearer test-openrouter-key"
    assert captured["json"]["model"] == "deepseek/deepseek-v4-flash"
    assert captured["json"]["provider"] == {"order": ["DeepSeek"], "allow_fallbacks": False}
    assert captured["json"]["response_format"] == {"type": "json_object"}
    assert "Z1234567" not in captured["json"]["messages"][1]["content"]
    assert plan["request_summary"].startswith("DeepSeek structured summary")
    assert "Z1234567" not in plan["customer_message_draft"]
    assert "****4567" in plan["customer_itinerary_draft"]
    assert plan["budget_policy_check"]["budget_status"] == "Needs Approval"
    assert plan["budget_policy_check"]["approval_required"] is True
    assert plan["travel_options"][0]["estimated_cost"] != 1
    assert "OpenRouter generated the narrative draft" in " ".join(plan["agent_notes"])


def test_generate_plan_logs_sanitized_llm_fallback_when_ai_fails(tmp_path, monkeypatch, caplog):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")
    caplog.set_level(logging.WARNING, logger=INTERNAL_LOGGER_NAME)

    import httpx

    def fake_post(url, headers, json, timeout):
        raise RuntimeError("OpenRouter failed for anika.rao@unipro.com passport Z1234567 token sk-live-abcdef1234567890")

    monkeypatch.setattr(httpx, "post", fake_post)

    payload = _corporate_payload()
    payload["traveller_details"]["passport_number"] = "Z1234567"
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    planned = client.post(f"/api/corporate/requests/{created.json()['id']}/plan", headers=_headers())

    assert planned.status_code == 200
    assert "corporate.llm_enhancement.failed" in caplog.text
    assert "anika.rao@unipro.com" not in caplog.text
    assert "Z1234567" not in caplog.text
    assert "sk-live-abcdef1234567890" not in caplog.text
    assert "[redacted-email]" in caplog.text
    assert "[redacted-id]" in caplog.text
    assert "OpenRouter planning enhancement unavailable" in " ".join(planned.json()["generated_plan"]["agent_notes"])


def test_traveler_policy_and_resend_notification_contracts(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post("/api/corporate/requests", json=_corporate_payload(), headers=_headers())
    request_id = created.json()["id"]

    travelers = client.get("/api/travelers", headers=_headers("admin.user@unipro.com", "review traveler roster"))
    assert travelers.status_code == 200
    assert travelers.json()[0]["email"] == "anika.rao@unipro.com"
    traveler_detail = client.get(f"/api/travelers/{travelers.json()[0]['id']}", headers=_headers("admin.user@unipro.com", "review traveler dossier"))
    assert traveler_detail.status_code == 200
    assert traveler_detail.json()["documents"][0]["label"] == "Passport"

    policies = client.get("/api/policies", headers=_headers("admin.user@unipro.com", "review policy center"))
    assert policies.status_code == 200
    policy = policies.json()[0]
    assert policy["id"] == "policy_global_travel_2024"

    versions = client.get(
        f"/api/policies/{policy['id']}/versions",
        headers=_headers("admin.user@unipro.com", "review policy lifecycle"),
    )
    assert versions.status_code == 200

    revision_id = versions.json()[0]["id"]

    approved = client.post(
        f"/api/policies/{policy['id']}/revisions/{revision_id}/approve",
        json={"comment": "Approved for publication."},
        headers=_headers("admin.user@unipro.com", "approve policy revision"),
    )
    assert approved.status_code == 200
    assert approved.json()["revisions"][0]["status"] == "Approved"

    activity = client.get(
        f"/api/policies/{policy['id']}/activity",
        headers=_headers("admin.user@unipro.com", "review policy activity"),
    )
    assert activity.status_code == 200
    assert "approved" in activity.json()[0]["activity"].lower()

    notification = client.post(
        f"/api/corporate/requests/{request_id}/notifications",
        json={"kind": "approval_request", "to": ["manager@unipro.com"], "note": "Please review.", "attach_itinerary": False},
        headers=_headers(purpose="send approval request notification"),
    )
    assert notification.status_code == 200
    assert notification.json()["status"] == "configuration_required"
    assert notification.json()["safe_message"] == "Email provider is not configured. Continue with manual follow-up."
    assert notification.json()["body_text"] == "\n".join([
        "Approval needed for Anika Rao's travel plan",
        "Route: Hyderabad to Johannesburg",
        f"Request: {request_id}",
        "Please review.",
        "No booking, ticketing, or payment has been created by this notification.",
    ])
    assert notification.json()["attachment_names"] == []


def test_agent_can_save_registration_to_roster_and_update_without_duplicate(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    headers = _headers("demo.agent@unipro.com", "save traveler roster profile")

    created = client.put("/api/travelers/traveler_ananya_shah", json=_traveler_profile_payload(), headers=headers)

    assert created.status_code == 200
    assert created.json()["email"] == "ananya.shah@orbitex.example"
    assert created.json()["seat_preference"] == "Window seat"
    travelers = client.get("/api/travelers", headers=_headers("demo.agent@unipro.com", "review traveler roster"))
    matching = [traveler for traveler in travelers.json() if traveler["email"] == "ananya.shah@orbitex.example"]
    assert len(matching) == 1

    updated = client.put(
        "/api/travelers/traveler_ananya_shah",
        json=_traveler_profile_payload(meal_preference="Jain vegetarian meal", hotel_preference="Near Mumbai home office"),
        headers=headers,
    )

    assert updated.status_code == 200
    assert updated.json()["meal_preference"] == "Jain vegetarian meal"
    assert updated.json()["created_at"] == created.json()["created_at"]
    travelers_after_update = client.get("/api/travelers", headers=_headers("demo.agent@unipro.com", "review traveler roster"))
    matching_after_update = [traveler for traveler in travelers_after_update.json() if traveler["email"] == "ananya.shah@orbitex.example"]
    assert len(matching_after_update) == 1
    assert matching_after_update[0]["hotel_preference"] == "Near Mumbai home office"


def test_roster_profile_save_requires_a_purpose_header(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    token, _ = create_access_token("demo.agent@unipro.com")

    response = client.put(
        "/api/travelers/traveler_ananya_shah",
        json=_traveler_profile_payload(),
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 403


def test_resend_notification_and_webhook_store_safe_events(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post("/api/corporate/requests", json=_corporate_payload(), headers=_headers())
    request_id = created.json()["id"]
    client.post(f"/api/corporate/requests/{request_id}/plan", headers=_headers())
    finalized = client.post(
        f"/api/corporate/requests/{request_id}/finalize",
        json={"agent_reviewed": True, "approval_status": "Received"},
        headers=_headers(purpose="finalize approved corporate itinerary"),
    )
    assert finalized.status_code == 200
    monkeypatch.setenv("RESEND_API_KEY", "re_test_key")
    monkeypatch.setenv("RESEND_FROM_EMAIL", "Unipro Travel <travel@example.com>")

    import httpx

    class FakeResponse:
        def raise_for_status(self):
            return None

        def json(self):
            return {"id": "email_123"}

    captured = {}

    def fake_post(url, headers, json, timeout):
        captured["url"] = url
        captured["headers"] = headers
        captured["json"] = json
        return FakeResponse()

    monkeypatch.setattr(httpx, "post", fake_post)

    sent = client.post(
        f"/api/corporate/requests/{request_id}/notifications",
        json={"kind": "final_itinerary", "to": ["traveler@example.com"], "attach_itinerary": True},
        headers=_headers(purpose="send final itinerary notification"),
    )

    assert sent.status_code == 200
    assert sent.json()["status"] == "sent"
    assert captured["url"] == "https://api.resend.com/emails"
    assert captured["headers"]["Authorization"] == "Bearer re_test_key"
    assert captured["headers"]["Idempotency-Key"] == f"travel-ai/{request_id}/final_itinerary"
    assert "Approved itinerary details:" in captured["json"]["text"]
    assert "Outbound flight:" in captured["json"]["text"]
    assert "Hotel check-in starts: 15:00" in captured["json"]["text"]
    assert "Hotel checkout time: 11:00" in captured["json"]["text"]
    assert "Cab service provider: Johannesburg Airport Cars" in captured["json"]["text"]
    assert "Cab pickup: JNB at" in captured["json"]["text"]
    assert captured["json"]["attachments"][0]["filename"] == f"{request_id}-final-itinerary.pdf"
    assert captured["json"]["attachments"][0]["content_type"] == "application/pdf"

    webhook = client.post(
        "/api/mail/resend/webhook",
        json={
            "type": "email.sent",
            "data": {
                "email_id": "email_123",
                "to": ["traveler@example.com"],
                "subject": "Final itinerary",
                "tags": {"request_id": request_id, "kind": "final_itinerary"},
            },
        },
    )
    assert webhook.status_code == 200
    assert webhook.json()["status"] == "received"
    assert webhook.json()["request_id"] == request_id


def test_final_itinerary_pdf_download_is_client_handoff(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post("/api/corporate/requests", json=_corporate_payload(), headers=_headers())
    request_id = created.json()["id"]
    client.post(f"/api/corporate/requests/{request_id}/plan", headers=_headers())
    finalized = client.post(
        f"/api/corporate/requests/{request_id}/finalize",
        json={"agent_reviewed": True, "approval_status": "Received"},
        headers=_headers(purpose="finalize approved corporate itinerary"),
    )
    assert finalized.status_code == 200

    exported = client.get(
        f"/api/corporate/requests/{request_id}/export.pdf",
        headers=_headers(purpose="download finalized corporate itinerary pdf"),
    )

    assert exported.status_code == 200
    assert exported.headers["content-type"] == "application/pdf"
    assert exported.headers["content-disposition"] == f'attachment; filename="{request_id}-final-itinerary.pdf"'
    assert exported.content.startswith(b"%PDF-1.4")
    assert b"Final Travel Itinerary" in exported.content


def test_resend_received_email_with_pdf_attachment_creates_pending_request(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setenv("RESEND_API_KEY", "re_test_key")
    pdf_content = "\n".join([
        "(Traveller name: Inbound Sender) Tj",
        "(Traveller email: inbound.sender@example.com) Tj",
        "(Origin: Hyderabad) Tj",
        "(Destination: San Jose) Tj",
        "(Departure date: 2026-10-26) Tj",
        "(Return date: 2026-11-25) Tj",
        "(Purpose: Client meeting) Tj",
        "(Approved budget and currency: INR 200000) Tj",
    ]).encode("utf-8")

    import httpx

    class FakeResponse:
        def __init__(self, payload=None, content=b""):
            self.payload = payload
            self.content = content

        def raise_for_status(self):
            return None

        def json(self):
            if self.payload is None:
                raise ValueError("No JSON payload")
            return self.payload

    def fake_get(url, headers=None, timeout=30.0):
        if url.endswith("/emails/receiving/inbound_123"):
            return FakeResponse({"id": "inbound_123", "text": "Attached travel form.", "subject": "Travel request form"})
        if url.endswith("/emails/receiving/inbound_123/attachments"):
            return FakeResponse({
                "data": [{
                    "filename": "travel-request.pdf",
                    "content_type": "application/pdf",
                    "content": base64.b64encode(pdf_content).decode("ascii"),
                }]
            })
        if url.endswith("/emails/receiving/inbound_123/attachments/att_123"):
            return FakeResponse({
                "filename": "travel-request.pdf",
                "content_type": "application/pdf",
                "download_url": "https://cdn.example.test/travel-request.pdf",
            })
        if url == "https://cdn.example.test/travel-request.pdf":
            return FakeResponse(content=pdf_content)
        raise AssertionError(f"unexpected URL {url}")

    monkeypatch.setattr(httpx, "get", fake_get)

    webhook = client.post(
        "/api/mail/resend/webhook",
        json={
            "type": "email.received",
            "data": {
                "email_id": "inbound_123",
                "from": "Inbound Sender <inbound.sender@example.com>",
                "to": ["travel-intake@example.com"],
                "subject": "Travel request form",
                "attachments": [{"id": "att_123", "filename": "travel-request.pdf", "content_type": "application/pdf"}],
            },
        },
    )

    assert webhook.status_code == 200
    assert webhook.json()["status"] == "received"
    assert webhook.json()["request_id"]
    assert "entered the automated pipeline" in webhook.json()["safe_message"]

    requests = client.get("/api/corporate/requests", headers=_headers("demo.agent@unipro.com", "review corporate travel request queue"))
    created = [item for item in requests.json() if item["id"] == webhook.json()["request_id"]][0]
    assert created["status"] == "Processing"
    assert created["traveller_details"]["traveler_name"] == "Inbound Sender"
    assert created["travel_details"]["origin"] == "Hyderabad"
    assert created["travel_details"]["destination"] == "San Jose"
    assert created["budgets"]["total_budget"] == 200000


def test_resend_received_email_with_multiple_complete_travelers_creates_complete_requests(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setenv("RESEND_API_KEY", "re_test_key")
    text = """
Traveller name: Priya Menon
Traveller email: priya.menon@example.com
Employee ID: PRIYA-001
Department: Sales
Nationality: Indian
Passport expiry: 2030-01-15
Origin city / airport: Hyderabad
Destination city / airport: San Jose
Departure date: 2026-10-26
Return date: 2026-11-25
Purpose / meeting location: Client meeting
Meeting location: North First Street office
Cabin: economy
Approved budget and currency: 150000 INR
Flight preference: Nonstop or shortest route
Office or location preference: Hotel near North First Street office in San Jose
Airport transfer needed: Yes
Special requests: Vegetarian meal

Traveller name: Arjun Mehta
Traveller email: arjun.mehta@example.com
Employee ID: ARJUN-002
Department: Engineering
Nationality: Indian
Passport expiry: 2031-02-20
Origin city / airport: Bengaluru
Destination city / airport: Berlin
Departure date: 2026-09-02
Return date: 2026-09-08
Purpose / meeting location: Product workshop
Meeting location: Mitte partner office
Cabin: economy
Approved budget and currency: 2400 EUR
Flight preference: Morning arrival
Office or location preference: Hotel near Mitte partner office in Berlin
Airport transfer needed: Yes
Special requests: Quiet room

Traveller name: Nisha Iyer
Traveller email: nisha.iyer@example.com
Employee ID: NISHA-003
Department: Partnerships
Nationality: Indian
Passport expiry: 2032-03-25
Origin city / airport: Chennai
Destination city / airport: Singapore
Departure date: 2026-08-12
Return date: 2026-08-16
Purpose / meeting location: Renewal meeting
Meeting location: Marina Bay office
Cabin: economy
Approved budget and currency: 1800 USD
Flight preference: Morning departure
Office or location preference: Hotel near Marina Bay office in Singapore
Airport transfer needed: No
Special requests: Aisle seat
"""

    import httpx

    class FakeResponse:
        def raise_for_status(self):
            return None

        def json(self):
            return {"id": "multi_sales_team", "text": text, "subject": "October leadership travel batch"}

    monkeypatch.setattr(httpx, "get", lambda *args, **kwargs: FakeResponse())

    webhook = client.post(
        "/api/mail/resend/webhook",
        json={
            "type": "email.received",
            "data": {
                "email_id": "multi_sales_team",
                "to": ["travel-intake@example.com"],
                "subject": "October leadership travel batch",
            },
        },
    )

    assert webhook.status_code == 200
    assert webhook.json()["status"] == "received"
    assert "3 requests entered the automated pipeline" in webhook.json()["safe_message"]

    requests = client.get("/api/corporate/requests", headers=_headers("demo.agent@unipro.com", "review corporate travel request queue")).json()
    created = {item["traveller_details"]["traveler_email"]: item for item in requests if str(item["traveller_details"]["traveler_email"]) in {"priya.menon@example.com", "arjun.mehta@example.com", "nisha.iyer@example.com"}}
    assert set(created) == {"priya.menon@example.com", "arjun.mehta@example.com", "nisha.iyer@example.com"}
    assert created["priya.menon@example.com"]["preferences"]["hotel_preference"] == "Hotel near North First Street office in San Jose"
    assert created["arjun.mehta@example.com"]["preferences"]["hotel_preference"] == "Hotel near Mitte partner office in Berlin"
    assert created["nisha.iyer@example.com"]["preferences"]["hotel_preference"] == "Hotel near Marina Bay office in Singapore"
    for request in created.values():
        planned = client.post(f"/api/corporate/requests/{request['id']}/plan", headers=_headers("demo.agent@unipro.com", "generate complete ideal travel plan")).json()
        assert planned["generated_plan"]["missing_information"] == []


def test_resend_webhook_requires_valid_svix_signature_when_secret_is_configured(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    secret_bytes = b"travel-ai-webhook-secret"
    monkeypatch.setenv("RESEND_WEBHOOK_SECRET", "whsec_" + base64.b64encode(secret_bytes).decode("ascii"))
    event = {
        "type": "email.sent",
        "data": {
            "email_id": "email_456",
            "to": ["traveler@example.com"],
            "subject": "Approval request",
            "tags": {"request_id": "corp_req_123", "kind": "approval_request"},
        },
    }
    payload = json.dumps(event, separators=(",", ":")).encode("utf-8")

    missing_headers = client.post("/api/mail/resend/webhook", content=payload, headers={"Content-Type": "application/json"})
    assert missing_headers.status_code == 400

    svix_id = "msg_123"
    timestamp = "1716210000"
    signed_payload = f"{svix_id}.{timestamp}.{payload.decode('utf-8')}".encode("utf-8")
    signature = base64.b64encode(hmac.new(secret_bytes, signed_payload, hashlib.sha256).digest()).decode("ascii")

    accepted = client.post(
        "/api/mail/resend/webhook",
        content=payload,
        headers={
            "Content-Type": "application/json",
            "svix-id": svix_id,
            "svix-timestamp": timestamp,
            "svix-signature": f"v1,{signature}",
        },
    )

    assert accepted.status_code == 200
    assert accepted.json()["status"] == "received"
    assert accepted.json()["request_id"] == "corp_req_123"


def test_missing_info_and_passport_visa_blocking_detection(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    depart = date.today() + timedelta(days=30)
    ret = depart + timedelta(days=4)
    payload = {
        "traveller_details": {
            "traveler_name": "Anika Rao",
            "nationality": "India",
            "passport_expiry": str(ret - timedelta(days=1)),
        },
        "travel_details": {
            "origin": "Hyderabad",
            "destination": "Johannesburg",
            "destination_country": "South Africa",
            "depart_date": str(depart),
            "return_date": str(ret),
        },
        "budgets": {"total_budget": 3000},
    }
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    request_id = created.json()["id"]

    planned = client.post(f"/api/corporate/requests/{request_id}/plan", headers=_headers())

    plan = planned.json()["generated_plan"]
    assert planned.json()["status"] == "Missing Info"
    assert "traveller_details.traveler_email" in plan["missing_information"]
    assert plan["travel_readiness"]["passport_status"] == "Blocking Issue"
    assert plan["travel_readiness"]["visa_status"] == "Needs Review"


def test_missing_required_visa_and_expired_visa_are_blocking(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    client.post(
        "/api/corporate/upload-excel",
        files={"file": ("corporate_travel.xlsx", _workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_headers(purpose="import corporate travel workbook"),
    )
    detail = _corporate_payload()
    detail["traveller_details"]["traveler_name"] = "No Visa Traveler"
    detail["traveller_details"]["traveler_email"] = "no.visa@example.com"
    detail["traveller_details"]["employee_id"] = "E-NO-VISA"
    detail["traveller_details"]["visa_status"] = None
    detail["traveller_details"]["visa_expiry"] = None
    created = client.post("/api/corporate/requests", json=detail, headers=_headers(purpose="create copied corporate request"))

    planned = client.post(f"/api/corporate/requests/{created.json()['id']}/plan", headers=_headers())

    assert planned.json()["generated_plan"]["travel_readiness"]["visa_status"] == "Blocking Issue"

    detail["traveller_details"]["visa_status"] = "Valid visa on file"
    detail["traveller_details"]["visa_expiry"] = detail["travel_details"]["depart_date"]
    expired = client.post("/api/corporate/requests", json=detail, headers=_headers(purpose="create expired visa corporate request"))
    expired_plan = client.post(f"/api/corporate/requests/{expired.json()['id']}/plan", headers=_headers())

    assert expired_plan.json()["generated_plan"]["travel_readiness"]["visa_status"] == "Blocking Issue"


def test_corporate_llm_payload_uses_popia_tokens_instead_of_raw_pii(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")
    payload = _corporate_payload()
    payload["traveller_details"]["passport_number"] = "Z1234567"
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    request_id = created.json()["id"]

    import httpx

    captured = {}

    class FakeResponse:
        def raise_for_status(self):
            return None

        def json(self):
            user_message = next(item["content"] for item in captured["json"]["messages"] if item["role"] == "user")
            outbound_payload = json.loads(user_message)
            return {"choices": [{"message": {"content": json.dumps(outbound_payload["validated_base_plan"])}}]}

    def fake_post(url, headers, json, timeout):
        captured["url"] = url
        captured["headers"] = headers
        captured["json"] = json
        captured["timeout"] = timeout
        return FakeResponse()

    monkeypatch.setattr(httpx, "post", fake_post)

    planned = client.post(f"/api/corporate/requests/{request_id}/plan", headers=_headers())

    assert planned.status_code == 200
    outbound = json.dumps(captured["json"]["messages"])
    for raw_value in (
        "Anika Rao",
        "anika.rao@unipro.com",
        "+91 98765 43210",
        "E-101",
        "manager@unipro.com",
        "Z1234567",
    ):
        assert raw_value not in outbound
    assert "PERSON_" in outbound
    assert "EMAIL_" in outbound
    assert "PHONE_" in outbound
    assert "EMPLOYEE_ID_" in outbound
    assert "PASSPORT_" in outbound
    assert planned.json()["generated_plan"]["request_summary"].startswith("Anika Rao needs")


def test_request_form_upload_runs_automatic_pipeline_and_prepares_client_review_link(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    uploaded = client.post(
        "/api/corporate/requests/upload-excel",
        files={"file": ("travel-request.pdf", _pdf_form_bytes(), "application/pdf")},
        headers=_headers(purpose="import emailed client travel request form"),
    )

    assert uploaded.status_code == 200
    request_id = uploaded.json()["created_request_ids"][0]
    detail = client.get(f"/api/corporate/requests/{request_id}", headers=_headers()).json()
    assert detail["status"] == "Processing"
    assert len(detail["generated_plan"]["travel_options"]) == 3
    assert detail["client_review"]["status"] == "Sent"
    assert detail["client_review"]["review_url"]
    assert len(detail["generated_plan"]["ground_transfer_offers"]) == 3

    for option_index in range(1, 4):
        option_pdf = client.get(
            f"/api/corporate/requests/{request_id}/options/{option_index}.pdf",
            headers=_headers(purpose="download client itinerary option pdf"),
        )
        assert option_pdf.status_code == 200
        assert option_pdf.content.startswith(b"%PDF")

    audit = client.get(
        "/api/admin/audit",
        headers=_headers("admin.user@unipro.com", "review automated travel pipeline logs"),
    ).json()
    event_types = {event["event_type"] for event in audit}
    assert "corporate.pipeline.started" in event_types
    assert "corporate.pipeline.options_sent" in event_types
    assert "corporate.pipeline.review_link_sent" in event_types


def test_pipeline_events_are_available_as_json_and_server_sent_stream(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    uploaded = client.post(
        "/api/corporate/requests/upload-excel",
        files={"file": ("travel-request.pdf", _pdf_form_bytes(), "application/pdf")},
        headers=_headers(purpose="import emailed client travel request form"),
    )
    request_id = uploaded.json()["created_request_ids"][0]

    events = client.get(
        f"/api/corporate/requests/{request_id}/pipeline-events",
        headers=_headers(purpose="watch corporate travel pipeline"),
    )

    assert events.status_code == 200
    stages = {event["stage"] for event in events.json()}
    assert "request_state" in stages
    assert "review_link_email" in stages
    assert "email.approval_request" in stages

    with client.stream(
        "GET",
        f"/api/corporate/requests/{request_id}/pipeline-events/stream",
        headers=_headers(purpose="watch corporate travel pipeline stream"),
    ) as stream:
        body = "".join(stream.iter_text())

    assert stream.status_code == 200
    assert "event: pipeline.event" in body
    assert "event: pipeline.complete" in body
    assert "review_link_email" in body


def test_signed_client_review_link_allows_approval_without_exposing_internal_pii(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    payload = _corporate_payload()
    payload["traveller_details"]["passport_number"] = "Z1234567"
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    request_id = created.json()["id"]
    processed = client.post(f"/api/corporate/requests/{request_id}/pipeline", headers=_headers()).json()
    token = processed["client_review"]["review_url"].rsplit("/", 1)[-1]

    tampered = client.get(f"/api/corporate/reviews/{token[:-1]}x")
    review = client.get(f"/api/corporate/reviews/{token}")

    assert tampered.status_code == 401
    assert review.status_code == 200
    review_text = json.dumps(review.json())
    assert "E-101" not in review_text
    assert "+91" not in review_text
    assert "Z1234567" not in review_text
    assert len(review.json()["options"]) == 3
    assert "transfer_summary" in review.json()["options"][0]

    approved = client.post(
        f"/api/corporate/reviews/{token}",
        json={"action": "approve", "selected_option_index": 2},
    )

    assert approved.status_code == 200
    assert approved.json()["status"] == "Approved"
    detail = client.get(f"/api/corporate/requests/{request_id}", headers=_headers()).json()
    assert detail["status"] == "Completed"
    assert detail["approval_status"] == "Received"
    assert detail["client_review"]["selected_option_index"] == 2
    itinerary = detail["generated_plan"]["customer_itinerary_draft"]
    assert "Airport transfer:" in itinerary
    assert "Outbound flight departure / landing:" in itinerary
    assert "Hotel check-in starts: 15:00" in itinerary
    assert "Hotel checkout time: 11:00" in itinerary
    assert "Cab service provider:" in itinerary


def test_client_review_edit_requests_regenerate_three_options_and_stop_at_limit(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post("/api/corporate/requests", json=_corporate_payload(), headers=_headers())
    request_id = created.json()["id"]
    processed = client.post(f"/api/corporate/requests/{request_id}/pipeline", headers=_headers()).json()
    token = processed["client_review"]["review_url"].rsplit("/", 1)[-1]

    for expected_round in (1, 2, 3):
        edited = client.post(
            f"/api/corporate/reviews/{token}",
            json={"action": "request_edits", "edit_request_text": f"Round {expected_round}: move hotel closer to office and keep airport pickup."},
        )
        assert edited.status_code == 200
        assert edited.json()["revision_round"] == expected_round
        assert len(edited.json()["options"]) == 3
        detail = client.get(f"/api/corporate/requests/{request_id}", headers=_headers()).json()
        assert detail["client_review"]["revision_round"] == expected_round
        assert detail["client_review"]["change_summary"]
        token = detail["client_review"]["review_url"].rsplit("/", 1)[-1]

    blocked = client.post(
        f"/api/corporate/reviews/{token}",
        json={"action": "request_edits", "edit_request_text": "One more change after the limit."},
    )

    assert blocked.status_code == 200
    assert blocked.json()["status"] == "Agent Review Required"
    detail = client.get(f"/api/corporate/requests/{request_id}", headers=_headers()).json()
    assert detail["client_review"]["status"] == "Agent Review Required"


def test_ground_transfer_poc_uses_varied_synthetic_options_without_provider_call(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setenv("AMADEUS_CLIENT_ID", "amadeus-client")
    monkeypatch.setenv("AMADEUS_CLIENT_SECRET", "amadeus-secret")

    def fail_provider_call(*args, **kwargs):
        raise AssertionError("Ground transfer POC should not call live Amadeus")

    monkeypatch.setattr("httpx.post", fail_provider_call)
    payload = _corporate_payload(total_budget=3500)
    payload["travel_details"]["destination"] = "Johannesburg"
    payload["travel_details"]["destination_country"] = "South Africa"
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    planned = client.post(f"/api/corporate/requests/{created.json()['id']}/plan", headers=_headers())

    assert planned.status_code == 200
    plan = planned.json()["generated_plan"]
    transfers = plan["ground_transfer_offers"]
    assert len(transfers) == 3
    assert {transfer["source"] for transfer in transfers} == {"synthetic"}
    assert transfers[0]["vehicle_type"] == "Business sedan"
    assert transfers[1]["service_type"] == "MEET_AND_GREET"
    assert transfers[2]["vehicle_type"] == "Executive SUV"
    assert len({transfer["provider"] for transfer in transfers}) == 3
    assert plan["travel_options"][0]["ground_transfer_offer_id"] == transfers[0]["id"]
    assert any("Synthetic airport transfer options" in note for note in plan["agent_notes"])


def test_component_skip_form_plans_only_requested_segments_with_real_currency_conversion(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setenv("TRAVEL_AI_MCP_TOOLS_URL", "http://mcp-tools.test/tools")

    def fake_post(url, json, timeout):
        assert url == "http://mcp-tools.test/tools/convert_currency"
        assert json["from_currency"] == "USD"
        assert json["to_currency"] == "INR"

        class Response:
            def raise_for_status(self):
                return None

            def json(self):
                return {"converted_amount": json["amount"] * 90, "rate": 90}

        return Response()

    monkeypatch.setattr("app.currency.httpx.post", fake_post)
    payload = _corporate_payload(total_budget=180000)
    payload["travel_details"]["return_date"] = None
    payload["travel_details"]["include_return_flight"] = False
    payload["travel_details"]["include_hotel"] = False
    payload["budgets"]["currency"] = "INR"
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())

    planned = client.post(f"/api/corporate/requests/{created.json()['id']}/plan", headers=_headers())

    assert planned.status_code == 200
    body = planned.json()
    plan = body["generated_plan"]
    assert "travel_details.return_date" not in plan["missing_information"]
    assert plan["hotel_offers"] == []
    assert plan["selected_hotel_offer_id"] is None
    assert plan["travel_options"][0]["estimated_cost"] > 50000
    assert plan["budget_policy_check"]["estimated_cost"] == plan["travel_options"][0]["estimated_cost"]
    assert all("Hotel excluded" in option["hotel_summary"] for option in plan["travel_options"])
    assert all("Return flight excluded" in option["flight_summary"] for option in plan["travel_options"])
    assert any("MCP currency conversion" in note for note in plan["agent_notes"])


def test_pdf_request_form_component_flags_drive_automatic_pipeline(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    uploaded = client.post(
        "/api/corporate/requests/upload-excel",
        files={"file": ("travel-request.pdf", _component_skip_pdf_form_bytes(), "application/pdf")},
        headers=_headers(purpose="import emailed client travel request form"),
    )

    assert uploaded.status_code == 200
    request_id = uploaded.json()["created_request_ids"][0]
    detail = client.get(f"/api/corporate/requests/{request_id}", headers=_headers()).json()
    assert detail["status"] == "Processing"
    assert detail["travel_details"]["include_outbound_flight"] is True
    assert detail["travel_details"]["include_return_flight"] is False
    assert detail["travel_details"]["include_hotel"] is False
    assert detail["generated_plan"]["hotel_offers"] == []
    assert "travel_details.return_date" not in detail["generated_plan"]["missing_information"]


def test_automatic_pipeline_moves_incomplete_request_to_pending_details(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    created = client.post(
        "/api/corporate/requests",
        json={
            "traveller_details": {"traveler_name": "Missing Fields", "traveler_email": "missing@example.com"},
            "travel_details": {"destination": "Berlin"},
        },
        headers=_headers(),
    )
    request_id = created.json()["id"]

    processed = client.post(f"/api/corporate/requests/{request_id}/pipeline", headers=_headers())

    assert processed.status_code == 200
    body = processed.json()
    assert body["status"] == "Pending Details"
    assert body["generated_plan"]["missing_information"] == [
        "travel_details.origin",
        "travel_details.depart_date",
        "travel_details.return_date",
    ]


def test_client_approval_completes_request_and_queues_final_mail(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    payload = _corporate_payload()
    payload["traveller_details"]["passport_number"] = "Z1234567"
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    request_id = created.json()["id"]
    processed = client.post(f"/api/corporate/requests/{request_id}/pipeline", headers=_headers()).json()
    selected_option_name = processed["generated_plan"]["travel_options"][1]["option_name"]

    approved = client.post(
        f"/api/corporate/requests/{request_id}/client-approval",
        json={"selected_option_name": selected_option_name},
        headers=_headers(purpose="record client itinerary approval"),
    )

    assert approved.status_code == 200
    body = approved.json()
    assert body["status"] == "Completed"
    assert body["approval_status"] == "Received"
    assert body["generated_plan"]["customer_itinerary_draft"]

    exported = client.get(
        f"/api/corporate/requests/{request_id}/export.pdf",
        headers=_headers(purpose="download completed corporate itinerary pdf"),
    )
    assert exported.status_code == 200
    assert exported.content.startswith(b"%PDF")

    audit = client.get(
        "/api/admin/audit",
        headers=_headers("admin.user@unipro.com", "review automated travel pipeline logs"),
    ).json()
    event_types = {event["event_type"] for event in audit}
    assert "corporate.pipeline.client_approved" in event_types
    assert "corporate.pipeline.final_sent" in event_types


def test_travel_agent_can_review_pipeline_audit_and_email_events(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    payload = _corporate_payload()
    payload["traveller_details"]["passport_number"] = "Z1234567"
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    request_id = created.json()["id"]
    client.post(f"/api/corporate/requests/{request_id}/pipeline", headers=_headers())

    audit = client.get(
        "/api/admin/audit",
        headers=_headers("demo.agent@unipro.com", "review automated travel pipeline logs"),
    )
    events = client.get(
        "/api/corporate/admin/email-events",
        headers=_headers("demo.agent@unipro.com", "review corporate travel email audit events"),
    )

    assert audit.status_code == 200
    assert any(event["event_type"] == "corporate.pipeline.options_sent" for event in audit.json())
    assert events.status_code == 200
    body = events.json()
    assert body[0]["request_id"] == request_id
    assert body[0]["kind"] == "approval_request"
    assert body[0]["status"] in {"sent", "configuration_required", "failed"}
    assert body[0]["body_text"]
    assert f"Request: {request_id}" in body[0]["body_text"]
    assert "Review dashboard:" in body[0]["body_text"]
    assert body[0]["attachment_names"] == []
    assert "passport" not in body[0]["safe_message"].lower()


def test_resend_received_email_approval_reply_completes_pipeline(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setenv("RESEND_API_KEY", "re_test_key")
    payload = _corporate_payload()
    payload["traveller_details"]["passport_number"] = "Z1234567"
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    request_id = created.json()["id"]
    processed = client.post(f"/api/corporate/requests/{request_id}/pipeline", headers=_headers()).json()
    assert processed["status"] == "Processing"

    import httpx

    class FakeResponse:
        def raise_for_status(self):
            return None

        def json(self):
            return {
                "id": "approval_reply_123",
                "text": f"Approved option 2 for {request_id}. Please proceed with the final itinerary.",
                "subject": f"Approved itinerary {request_id}",
            }

    monkeypatch.setattr(httpx, "get", lambda *args, **kwargs: FakeResponse())

    webhook = client.post(
        "/api/mail/resend/webhook",
        json={
            "type": "email.received",
            "data": {
                "email_id": "approval_reply_123",
                "to": ["travel-intake@example.com"],
                "subject": f"Approved itinerary {request_id}",
            },
        },
    )

    assert webhook.status_code == 200
    assert webhook.json()["request_id"] == request_id
    assert "approved itinerary" in webhook.json()["safe_message"].lower()
    detail = client.get(f"/api/corporate/requests/{request_id}", headers=_headers()).json()
    assert detail["status"] == "Completed"
    assert detail["approval_status"] == "Received"


def test_resend_received_selected_option_pdf_completes_pipeline(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setenv("RESEND_API_KEY", "re_test_key")
    payload = _corporate_payload()
    payload["traveller_details"]["passport_number"] = "Z1234567"
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    request_id = created.json()["id"]
    processed = client.post(f"/api/corporate/requests/{request_id}/pipeline", headers=_headers()).json()
    assert processed["status"] == "Processing"
    option_pdf = client.get(
        f"/api/corporate/requests/{request_id}/options/2.pdf",
        headers=_headers(purpose="download selected itinerary option pdf"),
    ).content

    import httpx

    class FakeResponse:
        def __init__(self, payload=None, content=b""):
            self.payload = payload
            self.content = content

        def raise_for_status(self):
            return None

        def json(self):
            if self.payload is None:
                raise ValueError("No JSON payload")
            return self.payload

    def fake_get(url, headers=None, timeout=30.0):
        if url.endswith("/emails/receiving/selected_pdf_123"):
            return FakeResponse({"id": "selected_pdf_123", "text": "", "subject": "Selected itinerary option"})
        if url.endswith("/emails/receiving/selected_pdf_123/attachments"):
            return FakeResponse({
                "data": [{
                    "filename": f"{request_id}-option-2.pdf",
                    "content_type": "application/pdf",
                    "content": base64.b64encode(option_pdf).decode("ascii"),
                }]
            })
        raise AssertionError(f"unexpected URL {url}")

    monkeypatch.setattr(httpx, "get", fake_get)

    webhook = client.post(
        "/api/mail/resend/webhook",
        json={
            "type": "email.received",
            "data": {
                "email_id": "selected_pdf_123",
                "to": ["travel-intake@example.com"],
                "subject": "Selected itinerary option",
                "attachments": [{"filename": f"{request_id}-option-2.pdf", "content_type": "application/pdf"}],
            },
        },
    )

    assert webhook.status_code == 200
    assert webhook.json()["request_id"] == request_id
    assert "selected itinerary pdf" in webhook.json()["safe_message"].lower()
    detail = client.get(f"/api/corporate/requests/{request_id}", headers=_headers()).json()
    assert detail["status"] == "Completed"
    assert "Selected option: Fastest route" in detail["generated_plan"]["customer_itinerary_draft"]


def test_edit_plan_and_finalize_flow(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    payload = _corporate_payload()
    payload["traveller_details"]["passport_number"] = "Z1234567"
    created = client.post("/api/corporate/requests", json=payload, headers=_headers())
    request_id = created.json()["id"]
    planned = client.post(f"/api/corporate/requests/{request_id}/plan", headers=_headers()).json()
    plan = planned["generated_plan"]
    plan["agent_note"] = "Edited by agent before customer send."
    plan["customer_message_draft"] = "Please verify passport Z1234567 before departure."
    plan["customer_itinerary_draft"] = "Final itinerary for passport Z1234567."

    edited = client.put(
        f"/api/corporate/requests/{request_id}/plan",
        json={"generated_plan": plan},
        headers=_headers(purpose="edit generated corporate travel plan"),
    )
    assert edited.status_code == 200
    assert edited.json()["generated_plan"]["agent_note"] == "Edited by agent before customer send."
    assert "Z1234567" not in edited.json()["generated_plan"]["customer_message_draft"]
    assert "****4567" in edited.json()["generated_plan"]["customer_message_draft"]

    finalized = client.post(
        f"/api/corporate/requests/{request_id}/finalize",
        json={"agent_reviewed": True, "approval_status": "Received"},
        headers=_headers(purpose="finalize approved corporate itinerary"),
    )
    assert finalized.status_code == 200
    assert finalized.json()["status"] == "Finalized"

    exported = client.get(
        f"/api/corporate/requests/{request_id}/export.xlsx",
        headers=_headers(purpose="download finalized corporate itinerary excel"),
    )
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), data_only=True)
    assert workbook.sheetnames == ["Request", "Recommended Plans", "Final Itinerary"]
    assert workbook["Request"]["B2"].value == request_id
    assert workbook["Request"]["B5"].value == "****4567"
    assert "Z1234567" not in " ".join(str(cell.value or "") for row in workbook["Request"].iter_rows() for cell in row)
    assert "Z1234567" not in " ".join(str(cell.value or "") for row in workbook["Final Itinerary"].iter_rows() for cell in row)
    assert workbook["Final Itinerary"]["B4"].value == "Edited by agent before customer send."

    summary = client.get(
        "/api/corporate/admin/summary",
        headers=_headers("admin.user@unipro.com", "review corporate travel program summary"),
    )
    assert summary.status_code == 200
    assert summary.json()["finalized"] == 1
    assert summary.json()["approval_required"] == 0
    assert summary.json()["visa_issues"] == 0
    assert summary.json()["common_destinations"] == [{"destination": "Johannesburg", "count": 1}]
