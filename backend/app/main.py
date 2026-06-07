import base64
import hashlib
import hmac
import asyncio
import json
import logging
import os
import re
import zipfile
from contextlib import asynccontextmanager
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from email.utils import parseaddr
from io import BytesIO
import time
from typing import Any
from uuid import uuid4
from xml.etree import ElementTree
from xml.sax.saxutils import escape as xml_escape

import httpx
from fastapi import Depends, FastAPI, File, Header, HTTPException, Request, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse

from .agent import chat_with_travel_ai, generate_corporate_plan_with_request, plan_trip
from .budgeting import budget_with_effective_total
from .currency import convert_final_amount, convert_planning_amount, origin_city_currency
from .env_loader import load_travel_ai_env
from .internal_logger import INTERNAL_LOGGER_NAME, log_internal_issue
from .mailer import send_resend_notification
from .request_rules import request_with_component_dependencies
from .models import (
    AdminSummary,
    AuditEvent,
    AuthContext,
    AuthTokenRequest,
    AuthTokenResponse,
    ChatRequest,
    ChatResponse,
    CompanyPipelineStatus,
    CorporatePipelineEvent,
    CompanyPolicyImportResponse,
    CompanyDetails,
    BudgetPolicyCheck,
    CorporateAdminSummary,
    CorporateClientReview,
    CorporateClientReviewEvent,
    CorporateClientApprovalRequest,
    CorporateCriticalIssueUpdate,
    CorporateGroundTransferOffer,
    CorporateReviewFlight,
    CorporateReviewHotel,
    CorporateReviewOption,
    CorporateReviewResponse,
    CorporateReviewSubmitRequest,
    CorporateReviewTransfer,
    CorporateTravelPlan,
    CorporateFinalizeRequest,
    CorporateImportResponse,
    CorporatePlanUpdate,
    CorporateRequestStatusUpdate,
    CorporateTravelRequest,
    CurrencyConversionRequest,
    CurrencyConversionResponse,
    EmailEvent,
    NotificationRequest,
    PolicyActivityEvent,
    PolicyGroup,
    PolicyRevision,
    PlanResponse,
    ResendWebhookEvent,
    RevisionActionRequest,
    TravelBudget,
    TravelDetails,
    TravelOption,
    TravelPreferences,
    TravelReadiness,
    TravelerProfile,
    TravellerDetails,
    TravelRequest,
    Trip,
)
from .security import (
    SecurityError,
    authorize_traveler_agent_chain,
    create_access_token,
    demo_auth_context,
    ensure_any_scope,
    ensure_scope,
    mask_sensitive_customer_text,
    public_error_message,
    require_purpose,
    resolve_demo_login_identity,
    verify_access_token,
)
from .store import TravelStore

load_travel_ai_env()


def cors_allowed_origins() -> list[str]:
    raw = os.getenv("FRONTEND_ORIGIN") or os.getenv("TRAVEL_AI_ALLOWED_ORIGINS") or "http://127.0.0.1:3100,http://localhost:3100"
    return [origin.strip() for origin in raw.split(",") if origin.strip()]


def validate_production_environment() -> None:
    if os.getenv("TRAVEL_AI_ENV") != "production":
        return
    missing = [
        name for name in ("TRAVEL_AI_TOKEN_SECRET", "TRAVEL_AI_ENCRYPTION_KEY", "TRAVEL_AI_DB_PATH")
        if not os.getenv(name)
    ]
    if not (os.getenv("FRONTEND_ORIGIN") or os.getenv("TRAVEL_AI_ALLOWED_ORIGINS")):
        missing.append("FRONTEND_ORIGIN or TRAVEL_AI_ALLOWED_ORIGINS")
    if missing:
        raise RuntimeError(f"Missing production Travel AI environment values: {', '.join(missing)}")


@asynccontextmanager
async def lifespan(app: FastAPI):
    validate_production_environment()
    yield


app = FastAPI(title="Travel AI Backend", lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_allowed_origins(),
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


CLIENT_REVIEW_OPTION_COUNT = 3
CLIENT_REVIEW_MAX_REVISIONS = 3
CLIENT_REVIEW_TOKEN_TTL_HOURS = 72
CLIENT_REVIEW_TOKEN_VERSION = "travel-ai-client-review-v1"
CLIENT_REVIEW_DEFAULT_BASE_URL = "http://127.0.0.1:3100"
CLIENT_REVIEW_HOTEL_SKIP_TERMS = ("hotel", "lodging", "accommodation")
CLIENT_REVIEW_TRANSFER_SKIP_TERMS = ("cab", "cabs", "taxi", "airport transfer", "ground transfer", "transfer")
CLIENT_REVIEW_EDIT_MONTHS = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}
CLIENT_REVIEW_EDIT_MONTH_PATTERN = "|".join(sorted(CLIENT_REVIEW_EDIT_MONTHS, key=len, reverse=True))
CLIENT_REVIEW_EDIT_DATE_TOKEN_PATTERN = (
    rf"(?:\d{{4}}-\d{{1,2}}-\d{{1,2}}|"
    rf"(?:{CLIENT_REVIEW_EDIT_MONTH_PATTERN})\.?\s+\d{{1,2}},?\s+\d{{4}}|"
    rf"\d{{1,2}}\s+(?:{CLIENT_REVIEW_EDIT_MONTH_PATTERN})\.?,?\s+\d{{4}})"
)
CLIENT_REVIEW_PROTECTED_EDIT_TERMS = (
    "budget",
    "daily budget",
    "flight budget",
    "employee band",
    "approval band",
    "band",
    "policy tier",
    "cabin",
    "business class",
    "first class",
)
CLIENT_REVIEW_EDIT_VALUE_LOOKAHEAD = (
    r"(?=$|[.;!?]\s|\n|\s+(?:origin|destination|destination country|country|depart(?:ure)? date|return date|"
    r"travel dates?|meeting location|client office|office|trip purpose|purpose|preferred airline|airline|"
    r"flight preference|hotel preference|hotel area|meal preference|seat preference|timing preference|"
    r"traveler name|traveller name|traveler email|traveller email|phone|mobile|department|nationality|"
    r"passport expiry|visa expiry|visa status|budget|employee band|approval band|policy tier|cabin)\s*(?:to|as|is|are|:|-))"
)
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ISO_DATETIME_PATTERN = re.compile(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(?::\d{2})?")
DISPLAY_CURRENCY_CODES = {"USD", "INR", "EUR", "GBP", "CAD", "AUD", "JPY", "ZAR"}
PIPELINE_STREAM_MAX_FOLLOW_SECONDS = 15
PIPELINE_STREAM_POLL_SECONDS = 1.0
NEW_TRAVELER_PROFILE_WARNING = (
    "No matching traveler profile was found in the uploaded roster. "
    "Ask the traveler to register or update their profile for better results."
)
COUNTRY_ONLY_ROUTE_KEYS = {
    "australia",
    "canada",
    "china",
    "germany",
    "india",
    "japan",
    "southafrica",
    "uae",
    "uk",
    "unitedarabemirates",
    "unitedkingdom",
    "unitedstates",
    "unitedstatesofamerica",
    "usa",
}
COUNTRY_ROUTE_EXAMPLES = {
    "china": "Shanghai or Beijing",
    "india": "Hyderabad, Bengaluru, or Delhi",
    "unitedstates": "San Jose, New York, or Chicago",
    "unitedstatesofamerica": "San Jose, New York, or Chicago",
    "usa": "San Jose, New York, or Chicago",
    "unitedkingdom": "London",
    "uk": "London",
    "unitedarabemirates": "Dubai",
    "uae": "Dubai",
    "southafrica": "Johannesburg",
    "germany": "Berlin or Munich",
    "japan": "Tokyo",
}


class TravelFormValidationError(ValueError):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("; ".join(errors))

    @property
    def user_message(self) -> str:
        return "; ".join(self.errors)


BAND_POLICY_GUESS_ROWS = {
    "employee": {
        "policy_tier": "employee_estimate",
        "allowed_cabins": "economy",
        "hotel_tier": "Employee tier: 2/3-star hotels",
        "notes": "Estimated from submitted band because no uploaded Company Policy row matched.",
    },
    "manager": {
        "policy_tier": "manager_estimate",
        "allowed_cabins": "economy,premium_economy",
        "hotel_tier": "Manager tier: 3/4-star hotels",
        "notes": "Estimated from submitted band because no uploaded Company Policy row matched.",
    },
    "senior": {
        "policy_tier": "senior_estimate",
        "allowed_cabins": "economy,premium_economy,business",
        "hotel_tier": "CEO tier: 4/5-star hotels",
        "notes": "Estimated from submitted band because no uploaded Company Policy row matched.",
    },
    "executive": {
        "policy_tier": "executive_estimate",
        "allowed_cabins": "premium_economy,business,first",
        "hotel_tier": "CEO tier: 4/5-star hotels",
        "notes": "Estimated from submitted band because no uploaded Company Policy row matched.",
    },
}
PIPELINE_STAGE_BY_EVENT_TYPE = {
    "corporate.pipeline.started": "intake",
    "corporate.pipeline.pending_details": "intake",
    "corporate.pipeline.profile_resolved": "profile_context",
    "corporate.pipeline.planning_started": "provider_and_ai_planning",
    "corporate.pipeline.plan_generated": "plan_generated",
    "corporate.pipeline.options_blocked": "option_email",
    "corporate.pipeline.options_sent": "option_email",
    "corporate.pipeline.review_link_sent": "review_link_email",
    "corporate.pipeline.final_blocked": "final_itinerary",
    "corporate.pipeline.final_sent": "final_itinerary",
}


@dataclass(frozen=True)
class PlanningReferenceContext:
    policy_rows: list[dict[str, object]]
    history_rows: list[dict[str, object]]
    visa_rows: list[dict[str, object]]
    traveler_matches: list[dict[str, object]]
    enriched_request: CorporateTravelRequest
    profile_warning: str | None
    band_policy_note: str | None


def require_auth(authorization: str | None = Header(default=None)) -> AuthContext:
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Unauthorized")
    try:
        return verify_access_token(authorization.removeprefix("Bearer ").strip())
    except SecurityError as exc:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Unauthorized") from exc


def get_store() -> TravelStore:
    return TravelStore()


def audit_security_decision(
    store: TravelStore,
    context: AuthContext | None,
    event_type: str,
    message: str,
    purpose: str | None,
    decision: str,
    trip_id: str | None = None,
) -> None:
    store.save_audit_event(
        AuditEvent(
            trip_id=trip_id,
            actor_id=context.user_id if context else None,
            event_type=event_type,
            message=message,
            purpose=purpose,
            decision=decision,  # type: ignore[arg-type]
        )
    )


def _save_pipeline_stage(
    store: TravelStore,
    request: CorporateTravelRequest,
    context: AuthContext,
    purpose: str,
    event_type: str,
    message: str,
    *,
    decision: str = "allow",
) -> CorporateTravelRequest:
    saved = store.save_corporate_request(request)
    audit_security_decision(store, context, event_type, message, purpose, decision, saved.id)
    return saved


def protected_purpose(purpose: str | None, store: TravelStore, context: AuthContext, event_type: str) -> str:
    try:
        return require_purpose(purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, event_type, str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc


def _intake_required_missing(request: CorporateTravelRequest) -> list[str]:
    missing: list[str] = []
    traveller = request.traveller_details
    travel = request.travel_details
    if not traveller.traveler_name:
        missing.append("traveller_details.traveler_name")
    if not _delivery_email_for_request(request):
        missing.append("requester_email")
    if not travel.origin:
        missing.append("travel_details.origin")
    if not travel.destination:
        missing.append("travel_details.destination")
    if not travel.depart_date:
        missing.append("travel_details.depart_date")
    if (travel.include_return_flight or travel.include_hotel) and not travel.return_date:
        missing.append("travel_details.return_date")
    return missing


def _pending_details_plan(request: CorporateTravelRequest, missing: list[str]) -> CorporateTravelPlan:
    return CorporateTravelPlan(
        request_summary=_request_summary_for_pending_details(request),
        missing_information=missing,
        travel_readiness=TravelReadiness(
            passport_status="Needs Review",
            visa_status="Needs Review",
            transit_warning="Pending itinerary details.",
            document_notes=["Only itinerary-critical missing fields block the automated pipeline."],
        ),
        budget_policy_check=BudgetPolicyCheck(
            budget_status="Not Applied",
            policy_status="Needs Review",
            approval_required=False,
            approval_reason="Waiting for itinerary-critical details.",
            total_budget=None,
            estimated_cost=1,
        ),
        travel_options=[
            TravelOption(
                option_name="Best tier fit",
                flight_summary="Pending itinerary-critical details.",
                hotel_summary="Pending itinerary-critical details.",
                estimated_cost=1,
                policy_status="Pending Details",
                recommendation_reason="Complete the required fields before itinerary generation.",
            ),
            TravelOption(
                option_name="Fastest route",
                flight_summary="Pending itinerary-critical details.",
                hotel_summary="Pending itinerary-critical details.",
                estimated_cost=1,
                policy_status="Pending Details",
                recommendation_reason="Complete the required fields before itinerary generation.",
            ),
            TravelOption(
                option_name="Comfort-focused option",
                flight_summary="Pending itinerary-critical details.",
                hotel_summary="Pending itinerary-critical details.",
                estimated_cost=1,
                policy_status="Pending Details",
                recommendation_reason="Complete the required fields before itinerary generation.",
            ),
        ],
        agent_note="Automated intake paused for missing itinerary-critical details.",
        agent_notes=["Intake completeness agent paused the workflow before planning."],
        customer_message_draft="Please complete the missing travel details so we can prepare itinerary options.",
        customer_itinerary_draft="Itinerary generation is pending required travel details.",
        approval_status="Pending Details",
    )


def _request_summary_for_pending_details(request: CorporateTravelRequest) -> str:
    traveller = request.traveller_details.traveler_name or "Traveller pending"
    origin = request.travel_details.origin or "origin pending"
    destination = request.travel_details.destination or "destination pending"
    return f"{traveller} request from {origin} to {destination} is waiting for itinerary-critical details."


def _elapsed_ms(started_at: float) -> int:
    return max(0, int(round((time.perf_counter() - started_at) * 1000)))


def _request_with_effective_budget(request: CorporateTravelRequest) -> CorporateTravelRequest:
    budget = budget_with_effective_total(request.budgets, request.travel_details)
    if budget == request.budgets:
        return request
    return request.model_copy(update={"budgets": budget})


def _planning_reference_context(store: TravelStore, request: CorporateTravelRequest) -> PlanningReferenceContext:
    policy_rows = store.list_reference_rows("company_policy")
    history_rows = store.list_reference_rows("traveller_history")
    visa_rows = store.list_reference_rows("visa_rules")
    traveler_matches = _matching_traveler_rows(request, history_rows)
    profile_warning = _traveler_profile_warning(history_rows, traveler_matches)
    profile_enriched = _enrich_request_from_reference_data(request, history_rows, visa_rows, traveler_matches)
    policy_enriched, matched_policy_rows, band_policy_note = _request_with_band_policy(profile_enriched, policy_rows)
    enriched = _request_with_effective_budget(policy_enriched)
    return PlanningReferenceContext(
        policy_rows=matched_policy_rows or policy_rows,
        history_rows=history_rows,
        visa_rows=visa_rows,
        traveler_matches=traveler_matches,
        enriched_request=enriched,
        profile_warning=profile_warning,
        band_policy_note=band_policy_note,
    )


def _log_corporate_pipeline_timing(
    request: CorporateTravelRequest,
    stage_times: dict[str, int],
    total_ms: int,
    *,
    pending_details: bool,
) -> None:
    log_internal_issue(
        logging.getLogger(INTERNAL_LOGGER_NAME),
        "corporate.pipeline.timing",
        "Corporate travel pipeline timing.",
        level=logging.WARNING,
        request_id=request.id,
        total_ms=total_ms,
        pending_details=pending_details,
        **stage_times,
    )


def _run_automated_pipeline(
    store: TravelStore,
    request: CorporateTravelRequest,
    context: AuthContext,
    purpose: str,
) -> CorporateTravelRequest:
    total_started_at = time.perf_counter()
    stage_times: dict[str, int] = {}
    audit_security_decision(store, context, "corporate.pipeline.started", "Automated travel pipeline started.", purpose, "allow", request.id)
    stage_started_at = time.perf_counter()
    missing = _intake_required_missing(request)
    stage_times["intake_ms"] = _elapsed_ms(stage_started_at)
    if missing:
        updated = request.model_copy(update={
            "status": "Pending Details",
            "generated_plan": _pending_details_plan(request, missing),
            "updated_at": datetime.now(timezone.utc),
        })
        saved = store.save_corporate_request(updated)
        audit_security_decision(store, context, "corporate.pipeline.pending_details", f"Missing required details: {', '.join(missing)}.", purpose, "allow", saved.id)
        _log_corporate_pipeline_timing(
            saved,
            stage_times,
            _elapsed_ms(total_started_at),
            pending_details=True,
        )
        return saved

    stage_started_at = time.perf_counter()
    planning = _planning_reference_context(store, request)
    stage_times["reference_context_ms"] = _elapsed_ms(stage_started_at)
    profile_stage_message = (
        "Uploaded traveler profile matched and request context was enriched."
        if planning.traveler_matches
        else "No uploaded traveler profile matched; planning will continue with submitted form details."
    )
    staged_request = _save_pipeline_stage(
        store,
        planning.enriched_request.model_copy(update={"status": "Processing", "updated_at": datetime.now(timezone.utc)}),
        context,
        purpose,
        "corporate.pipeline.profile_resolved",
        profile_stage_message,
    )
    planning = PlanningReferenceContext(
        policy_rows=planning.policy_rows,
        history_rows=planning.history_rows,
        visa_rows=planning.visa_rows,
        traveler_matches=planning.traveler_matches,
        enriched_request=staged_request,
        profile_warning=planning.profile_warning,
        band_policy_note=planning.band_policy_note,
    )

    stage_started_at = time.perf_counter()
    audit_security_decision(store, context, "corporate.pipeline.planning_started", "Provider search and itinerary planning started.", purpose, "allow", staged_request.id)
    planned_request, raw_plan = generate_corporate_plan_with_request(
        planning.enriched_request,
        policy_rows=planning.policy_rows,
        traveller_history_rows=planning.history_rows,
        visa_rule_rows=planning.visa_rows,
    )
    plan = _customer_safe_plan(raw_plan)
    stage_times["plan_ms"] = _elapsed_ms(stage_started_at)
    plan = _plan_with_profile_warning(plan, planning.profile_warning)
    plan = _plan_with_agent_note(plan, planning.band_policy_note)
    plan = plan.model_copy(update={"missing_information": [], "approval_status": "Processing"})
    approval_status = "Required"
    processing = planned_request.model_copy(update={
        "generated_plan": plan,
        "status": "Processing",
        "approval_status": approval_status,
        "updated_at": datetime.now(timezone.utc),
    })
    saved = _save_pipeline_stage(
        store,
        processing,
        context,
        purpose,
        "corporate.pipeline.plan_generated",
        "Itinerary planning agent generated three client options.",
    )
    stage_started_at = time.perf_counter()
    saved = _send_option_package_agent(store, saved, context, purpose)
    stage_times["option_email_ms"] = _elapsed_ms(stage_started_at)
    _log_corporate_pipeline_timing(
        saved,
        stage_times,
        _elapsed_ms(total_started_at),
        pending_details=False,
    )
    return saved


def _delivery_email_for_request(request: CorporateTravelRequest) -> str | None:
    return request.requester_email or request.traveller_details.traveler_email


def _send_option_package_agent(store: TravelStore, request: CorporateTravelRequest, context: AuthContext, purpose: str) -> CorporateTravelRequest:
    to_email = _delivery_email_for_request(request)
    if not to_email:
        audit_security_decision(store, context, "corporate.pipeline.options_blocked", "Client email is missing; option PDFs were not sent.", purpose, "deny", request.id)
        return request
    review_ready = _prepare_client_review_link(request, change_summary=None)
    saved_request = store.save_corporate_request(review_ready)
    event = send_resend_notification(
        saved_request.id,
        saved_request,
        NotificationRequest(
            kind="approval_request",
            to=[to_email],
            note="Please review the three itinerary options in the dashboard link and approve one option or request edits.",
            review_url=saved_request.client_review.review_url if saved_request.client_review else None,
            review_round=saved_request.client_review.revision_round if saved_request.client_review else 0,
        ),
    )
    saved = store.save_email_event(event)
    audit_security_decision(store, context, "corporate.pipeline.options_sent", saved.safe_message, purpose, "allow", saved_request.id)
    audit_security_decision(store, context, "corporate.pipeline.review_link_sent", saved.safe_message, purpose, "allow", saved_request.id)
    return saved_request


def _send_final_itinerary_agent(store: TravelStore, request: CorporateTravelRequest, context: AuthContext, purpose: str) -> EmailEvent | None:
    to_email = _delivery_email_for_request(request)
    if not to_email:
        audit_security_decision(store, context, "corporate.pipeline.final_blocked", "Client email is missing; final itinerary was not sent.", purpose, "deny", request.id)
        return None
    event = send_resend_notification(
        request.id,
        request,
        NotificationRequest(
            kind="final_itinerary",
            to=[to_email],
            note="Your approved final itinerary is attached.",
            attach_itinerary=True,
        ),
        _corporate_request_pdf_bytes(request),
        f"{request.id}-final-itinerary.pdf",
        "application/pdf",
    )
    saved = store.save_email_event(event)
    audit_security_decision(store, context, "corporate.pipeline.final_sent", saved.safe_message, purpose, "allow", request.id)
    return saved


def _prepare_client_review_link(request: CorporateTravelRequest, change_summary: str | None) -> CorporateTravelRequest:
    if not request.generated_plan or not request.generated_plan.travel_options:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Client review options are not ready")
    if request.status == "Cancelled" or (request.client_review and request.client_review.status == "Cancelled"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Cancelled itinerary cannot be sent for review")
    now = datetime.now(timezone.utc)
    revision_round = request.client_review.revision_round if request.client_review else 0
    token_id = uuid4().hex
    expires_at = now + timedelta(hours=CLIENT_REVIEW_TOKEN_TTL_HOURS)
    token = _sign_client_review_token(request.id, token_id, revision_round, expires_at)
    review = CorporateClientReview(
        status="Sent",
        token_id=token_id,
        review_url=_review_url(token),
        selected_option_index=None,
        edit_request_text=None,
        revision_round=revision_round,
        change_summary=change_summary,
        expires_at=expires_at,
        sent_at=now,
        submitted_at=None,
    )
    event = CorporateClientReviewEvent(
        action="sent",
        revision_round=revision_round,
        change_summary=change_summary,
        created_at=now,
    )
    return request.model_copy(update={
        "client_review": review,
        "client_review_history": [*request.client_review_history, event],
        "updated_at": now,
    })


def _sign_client_review_token(request_id: str, token_id: str, revision_round: int, expires_at: datetime) -> str:
    payload = {
        "version": CLIENT_REVIEW_TOKEN_VERSION,
        "request_id": request_id,
        "token_id": token_id,
        "revision_round": revision_round,
        "exp": int(expires_at.timestamp()),
    }
    encoded_payload = _urlsafe_b64encode(json.dumps(payload, separators=(",", ":"), sort_keys=True).encode("utf-8"))
    signature = hmac.new(_client_review_secret().encode("utf-8"), encoded_payload.encode("ascii"), hashlib.sha256).digest()
    return f"{encoded_payload}.{_urlsafe_b64encode(signature)}"


def _decode_client_review_token(token: str) -> dict[str, object]:
    try:
        encoded_payload, encoded_signature = token.split(".", 1)
        expected = hmac.new(_client_review_secret().encode("utf-8"), encoded_payload.encode("ascii"), hashlib.sha256).digest()
        if not hmac.compare_digest(_urlsafe_b64encode(expected), encoded_signature):
            raise ValueError("Invalid signature")
        payload = json.loads(_urlsafe_b64decode(encoded_payload))
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Review link is invalid") from exc
    if payload.get("version") != CLIENT_REVIEW_TOKEN_VERSION:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Review link is invalid")
    exp = int(payload.get("exp") or 0)
    if exp <= int(datetime.now(timezone.utc).timestamp()):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Review link has expired")
    return payload


def _load_request_from_review_token(store: TravelStore, token: str) -> CorporateTravelRequest:
    payload = _decode_client_review_token(token)
    request_id = str(payload.get("request_id") or "")
    request = store.get_corporate_request(request_id)
    if not request or not request.generated_plan or not request.client_review:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Review link is not available")
    token_id = str(payload.get("token_id") or "")
    revision_round = int(payload.get("revision_round") or 0)
    if request.client_review.token_id != token_id or request.client_review.revision_round != revision_round:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Review link is no longer active")
    if request.client_review.expires_at and request.client_review.expires_at <= datetime.now(timezone.utc):
        expired = request.model_copy(update={
            "client_review": request.client_review.model_copy(update={"status": "Expired"}),
            "updated_at": datetime.now(timezone.utc),
        })
        store.save_corporate_request(expired)
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Review link has expired")
    return request


def _client_review_secret() -> str:
    configured = os.getenv("TRAVEL_AI_REVIEW_TOKEN_SECRET") or os.getenv("TRAVEL_AI_TOKEN_SECRET")
    if configured:
        return configured
    if os.getenv("TRAVEL_AI_ENV") == "production":
        raise RuntimeError("TRAVEL_AI_REVIEW_TOKEN_SECRET or TRAVEL_AI_TOKEN_SECRET is required in production")
    return "travel-ai-local-review-secret"


def _review_url(token: str) -> str:
    return f"{_review_public_base_url().rstrip('/')}/review/{token}"


def _review_public_base_url() -> str:
    configured = os.getenv("TRAVEL_AI_REVIEW_BASE_URL") or os.getenv("TRAVEL_AI_PUBLIC_BASE_URL")
    if configured:
        return configured
    origin = (os.getenv("FRONTEND_ORIGIN") or "").split(",", 1)[0].strip()
    return origin or CLIENT_REVIEW_DEFAULT_BASE_URL


def _urlsafe_b64encode(value: bytes) -> str:
    return base64.urlsafe_b64encode(value).decode("ascii").rstrip("=")


def _urlsafe_b64decode(value: str) -> bytes:
    padded = value + "=" * (-len(value) % 4)
    return base64.urlsafe_b64decode(padded.encode("ascii"))


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/api/auth/demo-login", response_model=AuthTokenResponse)
def demo_login(request: AuthTokenRequest) -> AuthTokenResponse:
    try:
        email = resolve_demo_login_identity(request.email, request.username, request.password)
    except SecurityError as exc:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Unauthorized") from exc
    token, context = create_access_token(email)
    return AuthTokenResponse(access_token=token, expires_at=context.token_expires_at, user=context)


@app.post("/api/agent/plan", response_model=PlanResponse)
def create_plan(
    request: TravelRequest,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> PlanResponse:
    purpose = protected_purpose(x_travel_purpose, store, context, "agent.plan.denied")
    try:
        authorize_traveler_agent_chain(context, purpose)
        response = plan_trip(request, owner_id=context.user_id, owner_department=context.department)
        for event in response.audit_events:
            event.trip_id = event.trip_id or response.trip.id
            event.actor_id = context.user_id
            event.purpose = purpose
            event.decision = "allow"
            store.save_audit_event(event)
        audit_security_decision(store, context, "agent.plan.allowed", "Traveler planning agent chain authorized.", purpose, "allow", response.trip.id)
        return response
    except SecurityError as exc:
        audit_security_decision(store, context, "agent.plan.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail=public_error_message(exc)) from exc


@app.post("/api/agent/chat", response_model=ChatResponse)
def chat_with_agent(
    request: ChatRequest,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> ChatResponse:
    purpose = protected_purpose(x_travel_purpose, store, context, "agent.chat.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
        response = chat_with_travel_ai(request, context, purpose)
        for event in response.audit_events:
            store.save_audit_event(event)
        audit_security_decision(store, context, "agent.chat.allowed", "Traveler assistant chat completed.", purpose, "allow", request.trip.id if request.trip else None)
        return response
    except SecurityError as exc:
        audit_security_decision(store, context, "agent.chat.denied", str(exc), purpose, "deny", request.trip.id if request.trip else None)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    except Exception as exc:
        audit_security_decision(store, context, "agent.chat.failed", public_error_message(exc), purpose, "deny", request.trip.id if request.trip else None)
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="Travel assistant is unavailable") from exc


@app.post("/api/tools/currency-conversion", response_model=CurrencyConversionResponse)
def convert_currency(
    request: CurrencyConversionRequest,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CurrencyConversionResponse:
    purpose = protected_purpose(x_travel_purpose, store, context, "tools.currency.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "tools.currency.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    converted = convert_final_amount(request.amount_usd, "USD", request.to_currency)
    audit_security_decision(store, context, "tools.currency.converted", "Currency conversion completed for travel estimate.", purpose, "allow")
    return converted


@app.get("/api/trips", response_model=list[Trip])
def list_trips(
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> list[Trip]:
    purpose = protected_purpose(x_travel_purpose, store, context, "trips.list.denied")
    try:
        ensure_scope(context, "self:trips", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "trips.list.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    audit_security_decision(store, context, "trips.list.allowed", "Traveler scoped trip history read.", purpose, "allow")
    return store.list_trips_for_owner(context.user_id)


@app.post("/api/trips", response_model=Trip, status_code=status.HTTP_201_CREATED)
def create_trip(
    request: TravelRequest,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> Trip:
    purpose = protected_purpose(x_travel_purpose, store, context, "trip.create.denied")
    try:
        authorize_traveler_agent_chain(context, purpose)
        response = plan_trip(request, owner_id=context.user_id, owner_department=context.department)
    except SecurityError as exc:
        audit_security_decision(store, context, "trip.create.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    events = [
        *response.audit_events,
        AuditEvent(
            trip_id=response.trip.id,
            actor_id=context.user_id,
            event_type="trip.created",
            message="Trip draft saved.",
            purpose=purpose,
            decision="allow",
        ),
    ]
    saved = store.save_trip(response.trip, events)
    audit_security_decision(store, context, "trip.create.allowed", "Traveler scoped trip draft saved.", purpose, "allow", saved.id)
    return saved


@app.get("/api/admin/summary", response_model=AdminSummary)
def admin_summary(
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> AdminSummary:
    purpose = protected_purpose(x_travel_purpose, store, context, "admin.summary.denied")
    try:
        ensure_scope(context, "admin:summary", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "admin.summary.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    audit_security_decision(store, context, "admin.summary.allowed", "Aggregate travel program summary read.", purpose, "allow")
    return store.admin_summary()


@app.get("/api/admin/audit", response_model=list[AuditEvent])
def admin_audit(
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> list[AuditEvent]:
    purpose = protected_purpose(x_travel_purpose, store, context, "admin.audit.denied")
    try:
        ensure_any_scope(context, {"admin:audit", "travel:plan"}, purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "admin.audit.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    audit_security_decision(store, context, "admin.audit.allowed", "Security audit log read.", purpose, "allow")
    return [event.model_copy(update={"actor_id": None}) for event in store.list_audit_events()]


@app.get("/api/corporate/requests", response_model=list[CorporateTravelRequest])
def list_corporate_requests(
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> list[CorporateTravelRequest]:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.requests.list.denied")
    try:
        ensure_any_scope(context, {"self:trips", "admin:summary"}, purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.requests.list.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    audit_security_decision(store, context, "corporate.requests.list.allowed", "Corporate requests read.", purpose, "allow")
    if "admin:summary" in context.scopes:
        return store.list_corporate_requests()
    return store.list_corporate_requests_for_owner(context.user_id)


@app.post("/api/corporate/requests", response_model=CorporateTravelRequest, status_code=status.HTTP_201_CREATED)
def create_corporate_request(
    request: CorporateTravelRequest,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CorporateTravelRequest:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.request.create.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.request.create.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    now = datetime.now(timezone.utc)
    try:
        _validate_corporate_request_route(request)
    except TravelFormValidationError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=exc.user_message) from exc
    normalized_request = request_with_component_dependencies(_request_with_effective_budget(request))
    saved = normalized_request.model_copy(
        update={
            "id": f"corp_req_{uuid4().hex[:12]}",
            "owner_id": context.user_id,
            "owner_department": context.department,
            "requester_email": normalized_request.requester_email or normalized_request.traveller_details.traveler_email or context.email,
            "created_at": normalized_request.created_at or now,
            "updated_at": now,
        }
    )
    saved = store.save_corporate_request(saved)
    audit_security_decision(store, context, "corporate.request.create.allowed", "Corporate request created.", purpose, "allow", saved.id)
    return saved


@app.get("/api/corporate/requests/{request_id}", response_model=CorporateTravelRequest)
def get_corporate_request(
    request_id: str,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CorporateTravelRequest:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.request.read.denied")
    request = _load_authorized_corporate_request(store, context, request_id, purpose)
    audit_security_decision(store, context, "corporate.request.read.allowed", "Corporate request detail read.", purpose, "allow", request.id)
    return request


@app.put("/api/corporate/requests/{request_id}", response_model=CorporateTravelRequest)
def update_corporate_request(
    request_id: str,
    request_update: CorporateTravelRequest,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CorporateTravelRequest:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.request.update.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.request.update.denied", str(exc), purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    existing = _load_authorized_corporate_request(store, context, request_id, purpose)
    normalized_update = request_with_component_dependencies(_request_with_effective_budget(request_update))
    updated = normalized_update.model_copy(
        update={
            "id": existing.id,
            "owner_id": existing.owner_id,
            "owner_department": existing.owner_department,
            "created_at": existing.created_at,
            "updated_at": datetime.now(timezone.utc),
        }
    )
    saved = store.save_corporate_request(updated)
    audit_security_decision(store, context, "corporate.request.update.allowed", "Corporate request updated.", purpose, "allow", saved.id)
    return saved


@app.delete("/api/corporate/requests/{request_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_corporate_request(
    request_id: str,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> None:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.request.delete.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.request.delete.denied", str(exc), purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    request = _load_authorized_corporate_request(store, context, request_id, purpose)
    store.delete_corporate_request(request.id)
    audit_security_decision(store, context, "corporate.request.delete.allowed", "Corporate request deleted.", purpose, "allow", request.id)


@app.post("/api/corporate/requests/{request_id}/plan", response_model=CorporateTravelRequest)
def generate_corporate_request_plan(
    request_id: str,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CorporateTravelRequest:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.plan.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.plan.denied", str(exc), purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    request = _load_authorized_corporate_request(store, context, request_id, purpose)
    try:
        _validate_corporate_request_route(request)
    except TravelFormValidationError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=exc.user_message) from exc
    planning = _planning_reference_context(store, request)
    planned_request, plan = generate_corporate_plan_with_request(
        planning.enriched_request,
        policy_rows=planning.policy_rows,
        traveller_history_rows=planning.history_rows,
        visa_rule_rows=planning.visa_rows,
    )
    plan = _plan_with_profile_warning(plan, planning.profile_warning)
    approval_status = "Required" if plan.budget_policy_check.approval_required else "Not Required"
    updated = planned_request.model_copy(
        update={
            "generated_plan": _customer_safe_plan(plan),
            "status": plan.approval_status,
            "approval_status": approval_status,
            "updated_at": datetime.now(timezone.utc),
        }
    )
    saved = store.save_corporate_request(updated)
    audit_security_decision(store, context, "corporate.plan.allowed", "Corporate travel plan generated.", purpose, "allow", saved.id)
    return saved


@app.post("/api/corporate/requests/{request_id}/pipeline", response_model=CorporateTravelRequest)
def run_corporate_request_pipeline(
    request_id: str,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CorporateTravelRequest:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.pipeline.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.pipeline.denied", str(exc), purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    request = _load_authorized_corporate_request(store, context, request_id, purpose)
    try:
        _validate_corporate_request_route(request)
    except TravelFormValidationError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=exc.user_message) from exc
    return _run_automated_pipeline(store, request, context, purpose)


@app.put("/api/corporate/requests/{request_id}/plan", response_model=CorporateTravelRequest)
def update_corporate_request_plan(
    request_id: str,
    request_update: CorporatePlanUpdate,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CorporateTravelRequest:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.plan.update.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.plan.update.denied", str(exc), purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    existing = _load_authorized_corporate_request(store, context, request_id, purpose)
    traveller_details = existing.traveller_details
    if request_update.traveller_details:
        traveller_details = traveller_details.model_copy(update=request_update.traveller_details.model_dump(exclude_unset=True))
    company_details = existing.company_details
    if request_update.company_details:
        company_details = company_details.model_copy(update=request_update.company_details.model_dump(exclude_unset=True))
    travel_details = existing.travel_details
    if request_update.travel_details:
        travel_details = travel_details.model_copy(update=request_update.travel_details.model_dump(exclude_unset=True))
    preferences = existing.preferences
    if request_update.preferences:
        preferences = preferences.model_copy(update=request_update.preferences.model_dump(exclude_unset=True))
    route_check = existing.model_copy(update={"travel_details": travel_details})
    try:
        _validate_corporate_request_route(route_check)
    except TravelFormValidationError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=exc.user_message) from exc
    updated = existing.model_copy(
        update={
            "generated_plan": _customer_safe_plan(request_update.generated_plan),
            "budgets": request_update.budgets or existing.budgets,
            "traveller_details": traveller_details,
            "company_details": company_details,
            "travel_details": travel_details,
            "preferences": preferences,
            "special_requests": request_update.special_requests if request_update.special_requests is not None else existing.special_requests,
            "status": request_update.generated_plan.approval_status,
            "updated_at": datetime.now(timezone.utc),
        }
    )
    updated = _request_with_effective_budget(updated)
    saved = store.save_corporate_request(updated)
    audit_security_decision(store, context, "corporate.plan.update.allowed", "Corporate travel plan edited.", purpose, "allow", saved.id)
    return saved


@app.post("/api/corporate/requests/{request_id}/status", response_model=CorporateTravelRequest)
def update_corporate_request_status(
    request_id: str,
    status_update: CorporateRequestStatusUpdate,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CorporateTravelRequest:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.status.denied")
    try:
        ensure_any_scope(context, {"travel:plan", "approval:read"}, purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.status.denied", str(exc), purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    existing = _load_authorized_corporate_request(store, context, request_id, purpose)
    updated = existing.model_copy(update={"status": status_update.status, "updated_at": datetime.now(timezone.utc)})
    saved = store.save_corporate_request(updated)
    audit_security_decision(store, context, "corporate.status.allowed", "Corporate request status updated.", purpose, "allow", saved.id)
    return saved


@app.post("/api/corporate/requests/{request_id}/critical-issue", response_model=CorporateTravelRequest)
def update_corporate_critical_issue(
    request_id: str,
    issue_update: CorporateCriticalIssueUpdate,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CorporateTravelRequest:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.critical_issue.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.critical_issue.denied", str(exc), purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    existing = _load_authorized_corporate_request(store, context, request_id, purpose)
    clean_issue = (issue_update.issue or "").strip() or None
    clean_status = issue_update.status
    if clean_status == "None":
        clean_issue = None
    elif clean_status == "Urgent" and not clean_issue:
        clean_issue = "Urgent travel disruption"
    updated = existing.model_copy(
        update={
            "critical_issue": clean_issue,
            "critical_issue_status": clean_status,
            "updated_at": datetime.now(timezone.utc),
        }
    )
    saved = store.save_corporate_request(updated)
    audit_security_decision(store, context, "corporate.critical_issue.allowed", "Corporate critical issue updated.", purpose, "allow", saved.id)
    return saved


@app.post("/api/corporate/requests/{request_id}/finalize", response_model=CorporateTravelRequest)
def finalize_corporate_request(
    request_id: str,
    finalize_request: CorporateFinalizeRequest,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CorporateTravelRequest:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.finalize.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.finalize.denied", str(exc), purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    existing = _load_authorized_corporate_request(store, context, request_id, purpose)
    approval_status = finalize_request.approval_status or existing.approval_status
    approval_required = bool(existing.generated_plan and existing.generated_plan.budget_policy_check.approval_required)
    if (
        not existing.generated_plan
        or not finalize_request.agent_reviewed
        or existing.status == "Cancelled"
        or (approval_required and approval_status != "Received")
    ):
        audit_security_decision(store, context, "corporate.finalize.blocked", "Final itinerary is not ready.", purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Final itinerary is not ready")
    updated = existing.model_copy(
        update={
            "status": "Finalized",
            "approval_status": approval_status,
            "updated_at": datetime.now(timezone.utc),
        }
    )
    saved = store.save_corporate_request(updated)
    audit_security_decision(store, context, "corporate.finalize.allowed", "Corporate final itinerary generated.", purpose, "allow", saved.id)
    return saved


@app.post("/api/corporate/requests/{request_id}/client-approval", response_model=CorporateTravelRequest)
def record_client_itinerary_approval(
    request_id: str,
    approval: CorporateClientApprovalRequest,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CorporateTravelRequest:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.client_approval.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.client_approval.denied", str(exc), purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    existing = _load_authorized_corporate_request(store, context, request_id, purpose)
    return _complete_request_from_client_approval_agent(store, existing, approval, context, purpose)


@app.get("/api/corporate/reviews/{token}", response_model=CorporateReviewResponse)
def get_client_review(token: str, store: TravelStore = Depends(get_store)) -> CorporateReviewResponse:
    request = _load_request_from_review_token(store, token)
    return _client_review_response(request)


@app.post("/api/corporate/reviews/{token}", response_model=CorporateReviewResponse)
def submit_client_review(
    token: str,
    submission: CorporateReviewSubmitRequest,
    store: TravelStore = Depends(get_store),
) -> CorporateReviewResponse:
    request = _load_request_from_review_token(store, token)
    context = demo_auth_context(os.getenv("TRAVEL_AI_REVIEW_ACTOR_EMAIL", "demo.agent@unipro.com"))
    purpose = "process signed client itinerary review"
    if submission.action == "approve":
        selected_index = submission.selected_option_index or 1
        completed = _complete_request_from_client_approval_agent(
            store,
            request,
            CorporateClientApprovalRequest(selected_option_index=selected_index),
            context,
            purpose,
        )
        return _client_review_response(completed)
    if submission.action == "cancel":
        cancelled = _cancel_request_from_client_review(store, request, context, purpose)
        return _client_review_response(cancelled)
    return _request_client_review_edits(store, request, submission, context, purpose)


def _cancel_request_from_client_review(
    store: TravelStore,
    request: CorporateTravelRequest,
    context: AuthContext,
    purpose: str,
) -> CorporateTravelRequest:
    if request.status in {"Completed", "Finalized"} or (request.client_review and request.client_review.status == "Approved"):
        audit_security_decision(store, context, "corporate.pipeline.client_cancel_blocked", "Approved itinerary cannot be cancelled from the client review link.", purpose, "deny", request.id)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Approved itinerary cannot be cancelled from this link")
    if not request.client_review:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Client review is not ready")

    now = datetime.now(timezone.utc)
    review = request.client_review.model_copy(update={
        "status": "Cancelled",
        "submitted_at": now,
    })
    saved = store.save_corporate_request(request.model_copy(update={
        "status": "Cancelled",
        "approval_status": "Rejected",
        "client_review": review,
        "client_review_history": [
            *request.client_review_history,
            CorporateClientReviewEvent(
                action="cancelled",
                revision_round=review.revision_round,
                created_at=now,
            ),
        ],
        "updated_at": now,
    }))
    _send_client_cancelled_email(store, saved, context, purpose)
    audit_security_decision(store, context, "corporate.pipeline.client_cancelled", "Client cancelled the itinerary review request.", purpose, "allow", saved.id)
    return saved


def _send_client_cancelled_email(store: TravelStore, request: CorporateTravelRequest, context: AuthContext, purpose: str) -> EmailEvent | None:
    to_email = _delivery_email_for_request(request)
    if not to_email:
        audit_security_decision(store, context, "corporate.pipeline.client_cancel_email_blocked", "Client email is missing; cancellation acknowledgement was not sent.", purpose, "deny", request.id)
        return None
    event = send_resend_notification(
        request.id,
        request,
        NotificationRequest(
            kind="client_cancelled",
            to=[to_email],
            note="Your travel request has been cancelled. The travel team will not finalize this itinerary unless you contact them again.",
        ),
    )
    saved = store.save_email_event(event)
    audit_security_decision(store, context, "corporate.pipeline.client_cancel_email_sent", saved.safe_message, purpose, "allow", request.id)
    return saved


def _request_client_review_edits(
    store: TravelStore,
    request: CorporateTravelRequest,
    submission: CorporateReviewSubmitRequest,
    context: AuthContext,
    purpose: str,
) -> CorporateReviewResponse:
    comment = (submission.edit_request_text or "").strip()
    if not comment:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Edit request text is required")
    if not request.client_review:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Client review is not ready")
    if request.status == "Cancelled" or request.client_review.status == "Cancelled":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Cancelled itinerary cannot be changed from this link")
    if request.status in {"Completed", "Finalized"} or request.client_review.status == "Approved":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Approved itinerary cannot be changed from this link")

    now = datetime.now(timezone.utc)
    current_round = request.client_review.revision_round
    history = [
        *request.client_review_history,
        CorporateClientReviewEvent(
            action="edits_requested",
            revision_round=current_round,
            edit_request_text=mask_sensitive_customer_text(comment),
            created_at=now,
        ),
    ]
    if current_round >= CLIENT_REVIEW_MAX_REVISIONS:
        review = request.client_review.model_copy(update={
            "status": "Agent Review Required",
            "edit_request_text": mask_sensitive_customer_text(comment),
            "submitted_at": now,
        })
        saved = store.save_corporate_request(request.model_copy(update={
            "client_review": review,
            "client_review_history": [
                *history,
                CorporateClientReviewEvent(
                    action="agent_review_required",
                    revision_round=current_round,
                    edit_request_text=mask_sensitive_customer_text(comment),
                    created_at=now,
                ),
            ],
            "updated_at": now,
        }))
        audit_security_decision(store, context, "corporate.pipeline.review_revision_blocked", "Client edit request reached the revision limit.", purpose, "deny", saved.id)
        return _client_review_response(saved)

    next_round = current_round + 1
    change_summary = _client_review_change_summary(next_round, comment)
    special_requests = [*request.special_requests]
    safe_edit_note = _client_safe_itinerary_edit_note(comment)
    if safe_edit_note:
        special_requests.append(f"Client edit request round {next_round}: {safe_edit_note}")
    request_for_revision = _request_with_client_edit_updates(request, comment).model_copy(update={
        "special_requests": special_requests,
    })
    policy_rows = store.list_reference_rows("company_policy")
    history_rows = store.list_reference_rows("traveller_history")
    visa_rows = store.list_reference_rows("visa_rules")
    planned_request, raw_plan = generate_corporate_plan_with_request(
        request_for_revision,
        policy_rows=policy_rows,
        traveller_history_rows=history_rows,
        visa_rule_rows=visa_rows,
    )
    plan = _customer_safe_plan(raw_plan)
    plan = plan.model_copy(update={"missing_information": [], "approval_status": "Processing"})
    revision_review = request.client_review.model_copy(update={
        "status": "Changes Requested",
        "revision_round": next_round,
        "edit_request_text": mask_sensitive_customer_text(comment),
        "change_summary": change_summary,
        "submitted_at": now,
    })
    revision_request = planned_request.model_copy(update={
        "generated_plan": plan,
        "status": "Processing",
        "approval_status": "Required",
        "client_review": revision_review,
        "client_review_history": history,
        "updated_at": now,
    })
    ready = _prepare_client_review_link(revision_request, change_summary)
    saved = store.save_corporate_request(ready)
    _send_review_link_email(store, saved, change_summary, context, purpose)
    audit_security_decision(store, context, "corporate.pipeline.review_revision_generated", f"Client review round {next_round} generated three revised options.", purpose, "allow", saved.id)
    return _client_review_response(saved)


def _request_with_client_edit_updates(request: CorporateTravelRequest, comment: str) -> CorporateTravelRequest:
    updated = _request_with_client_component_skips(request, comment)
    travel_updates = _client_travel_detail_updates(updated, comment)
    traveller_updates = _client_traveller_detail_updates(comment)
    preference_updates = _client_preference_updates(updated, comment)

    if travel_updates:
        updated = updated.model_copy(update={"travel_details": updated.travel_details.model_copy(update=travel_updates)})
    if traveller_updates:
        updated = updated.model_copy(update={"traveller_details": updated.traveller_details.model_copy(update=traveller_updates)})
    if preference_updates:
        updated = updated.model_copy(update={"preferences": updated.preferences.model_copy(update=preference_updates)})
    return updated


def _request_with_client_component_skips(request: CorporateTravelRequest, comment: str) -> CorporateTravelRequest:
    skip_hotel = _client_requested_component_skip(comment, CLIENT_REVIEW_HOTEL_SKIP_TERMS)
    skip_transfer = _client_requested_component_skip(comment, CLIENT_REVIEW_TRANSFER_SKIP_TERMS)
    if not (skip_hotel or skip_transfer):
        return request_with_component_dependencies(request)

    travel_updates: dict[str, object] = {}
    preference_updates: dict[str, object] = {}
    if skip_hotel:
        travel_updates["include_hotel"] = False
    if skip_transfer:
        travel_updates["include_ground_transfer"] = False
        preference_updates["airport_transfer_needed"] = False

    updated = request
    if travel_updates:
        updated = updated.model_copy(update={"travel_details": updated.travel_details.model_copy(update=travel_updates)})
    if preference_updates:
        updated = updated.model_copy(update={"preferences": updated.preferences.model_copy(update=preference_updates)})
    return request_with_component_dependencies(updated)


def _client_requested_component_skip(comment: str, terms: tuple[str, ...]) -> bool:
    for term in terms:
        term_pattern = re.escape(term)
        if re.search(rf"\b(?:do not|don't|dont)\s+(?:skip|remove|exclude)\b[^.?!\n]{{0,40}}\b{term_pattern}\b", comment, re.IGNORECASE):
            continue
        if re.search(rf"\b(?:skip|remove|exclude|without|no)\b[^.?!\n]{{0,40}}\b{term_pattern}\b", comment, re.IGNORECASE):
            return True
        if re.search(rf"\b{term_pattern}\s+(?:is\s+)?not\s+(?:needed|required)\b", comment, re.IGNORECASE):
            return True
    return False


def _client_travel_detail_updates(request: CorporateTravelRequest, comment: str) -> dict[str, object]:
    updates: dict[str, object] = {}
    updates.update(_client_date_updates(request, comment))
    text_fields = (
        ("origin", ("origin", "from city", "departure city"), 120),
        ("destination", ("destination", "to city", "arrival city"), 120),
        ("destination_country", ("destination country", "country"), 80),
        ("trip_purpose", ("trip purpose", "purpose"), 240),
        ("meeting_location", ("meeting location", "client office", "office"), 240),
    )
    for field, labels, max_length in text_fields:
        value = _client_edit_value(comment, labels, max_length)
        if value:
            updates[field] = value

    if "destination_country" not in updates:
        inferred_country = _infer_destination_country(
            str(updates.get("destination") or request.travel_details.destination or ""),
            str(updates.get("meeting_location") or request.travel_details.meeting_location or ""),
            request.preferences.hotel_preference or "",
        )
        if inferred_country:
            updates["destination_country"] = inferred_country
    return updates


def _client_traveller_detail_updates(comment: str) -> dict[str, object]:
    updates: dict[str, object] = {}
    text_fields = (
        ("traveler_name", ("traveler name", "traveller name", "passenger name"), 120),
        ("traveler_email", ("traveler email", "traveller email", "email"), 120),
        ("phone", ("phone", "mobile", "phone number"), 40),
        ("department", ("department",), 120),
        ("nationality", ("nationality",), 80),
        ("visa_status", ("visa status",), 120),
    )
    for field, labels, max_length in text_fields:
        value = _client_edit_value(comment, labels, max_length)
        if value:
            updates[field] = value

    for field, labels in (
        ("passport_expiry", ("passport expiry", "passport expiration")),
        ("visa_expiry", ("visa expiry", "visa expiration")),
    ):
        value = _client_edit_date_value(comment, labels)
        if value:
            updates[field] = value
    return updates


def _client_preference_updates(request: CorporateTravelRequest, comment: str) -> dict[str, object]:
    text_fields = [
        ("preferred_airline", ("preferred airline", "airline"), 120),
        ("flight_preference", ("flight preference", "flight timing", "flight time"), 160),
        ("meal_preference", ("meal preference", "meal"), 120),
        ("seat_preference", ("seat preference", "seat"), 80),
        ("timing_preference", ("timing preference", "timing"), 120),
    ]
    if request.travel_details.include_hotel:
        text_fields.extend([
            ("hotel_preference", ("hotel preference", "hotel"), 160),
            ("preferred_hotel_area", ("hotel area", "preferred hotel area"), 160),
        ])

    updates: dict[str, object] = {}
    for field, labels, max_length in text_fields:
        value = _client_edit_value(comment, labels, max_length)
        if value:
            updates[field] = value
    natural_flight = _client_natural_flight_preference(comment)
    if natural_flight:
        updates["flight_preference"] = natural_flight
    natural_hotel = _client_natural_hotel_preference(comment) if request.travel_details.include_hotel else None
    if natural_hotel:
        updates["hotel_preference"] = natural_hotel
    return updates


def _client_natural_flight_preference(comment: str) -> str | None:
    if re.search(r"\b(?:nonstop|non-stop|direct)\b", comment, re.IGNORECASE):
        return "Nonstop or direct flight"
    if re.search(r"\b(?:(?:only|just|at most|max(?:imum)?)\s+)?(?:one|1)[ -]?stop\b", comment, re.IGNORECASE):
        return "Only one-stop flight"
    return None


def _client_natural_hotel_preference(comment: str) -> str | None:
    if re.search(r"\b(?:different|another|alternate|alternative)\s+hotel\b|\bchange\s+(?:the\s+)?hotel\b", comment, re.IGNORECASE):
        return "Different hotel requested"
    return None


def _client_date_updates(request: CorporateTravelRequest, comment: str) -> dict[str, object]:
    updates: dict[str, date] = {}
    range_match = re.search(
        rf"\b(?:(?:please\s+)?(?:change|update|correct|set)\s+)?(?:the\s+)?(?:travel\s+)?dates?\s*(?:to|as|are|is|:|-)\s*"
        rf"(?P<depart>{CLIENT_REVIEW_EDIT_DATE_TOKEN_PATTERN})\s+(?:to|through|until|and|-)\s+"
        rf"(?P<return>{CLIENT_REVIEW_EDIT_DATE_TOKEN_PATTERN})",
        comment,
        re.IGNORECASE,
    )
    if range_match:
        depart = _client_edit_date_from_text(range_match.group("depart"))
        ret = _client_edit_date_from_text(range_match.group("return"))
        if depart and ret and ret >= depart:
            return {"depart_date": depart, "return_date": ret}

    relative_updates = _client_relative_day_date_updates(request, comment)
    if relative_updates:
        return relative_updates

    depart = _client_generic_depart_date(comment) or _client_specific_date(comment, (r"depart(?:ure)?", r"travel", r"start", r"leave"))
    ret = _client_specific_date(comment, (r"return", r"end"))
    if depart:
        updates["depart_date"] = depart
        if not ret and request.travel_details.depart_date and request.travel_details.return_date:
            updates["return_date"] = depart + (request.travel_details.return_date - request.travel_details.depart_date)
    if ret:
        updates["return_date"] = ret

    candidate_depart = updates.get("depart_date", request.travel_details.depart_date)
    candidate_return = updates.get("return_date", request.travel_details.return_date)
    if candidate_depart and candidate_return and candidate_return < candidate_depart:
        return {}
    return updates


def _client_relative_day_date_updates(request: CorporateTravelRequest, comment: str) -> dict[str, object]:
    if not request.travel_details.depart_date:
        return {}
    month = CLIENT_REVIEW_EDIT_MONTH_PATTERN
    match = re.search(
        rf"\b(?:change|update|correct|set)\s+(?:the\s+)?(?:travel\s+)?dates?\s+"
        rf"(?P<depart_month>{month})\s+(?P<old_depart>\d{{1,2}})\s+to\s+"
        rf"(?:(?P<new_depart_month>{month})\s+)?(?P<new_depart>\d{{1,2}})"
        rf"(?:\s+(?:and|,)\s+return\s+(?:(?P<return_month>{month})\s+)?(?P<old_return>\d{{1,2}})\s+to\s+"
        rf"(?:(?P<new_return_month>{month})\s+)?(?P<new_return>\d{{1,2}}))?",
        comment,
        re.IGNORECASE,
    )
    if not match:
        return {}
    depart_month = _client_edit_month_number(match.group("depart_month"))
    old_depart_day = _client_edit_day_number(match.group("old_depart"))
    new_depart_day = _client_edit_day_number(match.group("new_depart"))
    if not (depart_month and old_depart_day and new_depart_day):
        return {}
    current_depart = request.travel_details.depart_date
    if current_depart.month != depart_month or current_depart.day != old_depart_day:
        return {}
    new_depart = _client_month_day_date(
        match.group("new_depart_month") or match.group("depart_month"),
        new_depart_day,
        current_depart,
    )
    if not new_depart:
        return {}

    updates: dict[str, object] = {"depart_date": new_depart}
    current_return = request.travel_details.return_date
    if match.group("old_return") and current_return:
        return_month_text = match.group("return_month") or match.group("depart_month")
        old_return_month = _client_edit_month_number(return_month_text)
        old_return_day = _client_edit_day_number(match.group("old_return"))
        new_return_day = _client_edit_day_number(match.group("new_return"))
        if old_return_month and old_return_day and new_return_day and current_return.month == old_return_month and current_return.day == old_return_day:
            new_return = _client_month_day_date(match.group("new_return_month") or return_month_text, new_return_day, current_return)
            if new_return and new_return >= new_depart:
                updates["return_date"] = new_return
    elif current_return:
        updates["return_date"] = new_depart + (current_return - current_depart)
    return updates


def _client_edit_month_number(value: str | None) -> int | None:
    if not value:
        return None
    return CLIENT_REVIEW_EDIT_MONTHS.get(value.strip().lower())


def _client_edit_day_number(value: str | None) -> int | None:
    if not value:
        return None
    day = int(value)
    return day if 1 <= day <= 31 else None


def _client_month_day_date(month_text: str | None, day: int, reference: date) -> date | None:
    month = _client_edit_month_number(month_text)
    if not month:
        return None
    try:
        return date(reference.year, month, day)
    except ValueError:
        return None


def _client_generic_depart_date(comment: str) -> date | None:
    match = re.search(
        rf"\b(?:change|update|correct|set)\s+(?:the\s+)?date\s*(?:to|as|is|:|-)\s*(?P<date>{CLIENT_REVIEW_EDIT_DATE_TOKEN_PATTERN})",
        comment,
        re.IGNORECASE,
    )
    if not match:
        return None
    return _client_edit_date_from_text(match.group("date"))


def _client_specific_date(comment: str, labels: tuple[str, ...]) -> date | None:
    label_pattern = "|".join(labels)
    match = re.search(
        rf"\b(?:change|update|correct|set)?\s*(?:the\s+)?(?:{label_pattern})\s+date\s*(?:to|as|is|:|-)\s*(?P<date>{CLIENT_REVIEW_EDIT_DATE_TOKEN_PATTERN})",
        comment,
        re.IGNORECASE,
    )
    if not match:
        return None
    return _client_edit_date_from_text(match.group("date"))


def _client_edit_date_value(comment: str, labels: tuple[str, ...]) -> date | None:
    value = _client_edit_value(comment, labels, 40)
    if not value:
        return None
    return _client_edit_date_from_text(value)


def _client_edit_date_from_text(value: str) -> date | None:
    text = value.strip(" \"'`.,;:")
    iso_match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if iso_match:
        try:
            return date(int(iso_match.group(1)), int(iso_match.group(2)), int(iso_match.group(3)))
        except ValueError:
            return None

    month_first = re.fullmatch(
        rf"({CLIENT_REVIEW_EDIT_MONTH_PATTERN})\.?\s+(\d{{1,2}}),?\s+(\d{{4}})",
        text,
        re.IGNORECASE,
    )
    if month_first:
        return _client_date_from_parts(month_first.group(3), month_first.group(1), month_first.group(2))

    day_first = re.fullmatch(
        rf"(\d{{1,2}})\s+({CLIENT_REVIEW_EDIT_MONTH_PATTERN})\.?,?\s+(\d{{4}})",
        text,
        re.IGNORECASE,
    )
    if day_first:
        return _client_date_from_parts(day_first.group(3), day_first.group(2), day_first.group(1))
    return None


def _client_date_from_parts(year_text: str, month_text: str, day_text: str) -> date | None:
    month = _client_edit_month_number(month_text)
    day = _client_edit_day_number(day_text)
    if not month or not day:
        return None
    try:
        return date(int(year_text), month, day)
    except ValueError:
        return None


def _client_edit_value(comment: str, labels: tuple[str, ...], max_length: int) -> str | None:
    label_pattern = "|".join(re.escape(label) for label in labels)
    for pattern in (
        rf"\b(?:change|update|correct|set)\s+(?:the\s+)?(?:{label_pattern})\s*(?:to|as|is|:|-)\s+(.+?){CLIENT_REVIEW_EDIT_VALUE_LOOKAHEAD}",
        rf"\b(?:{label_pattern})\s*(?:to|as|is|are|:|-)\s+(.+?){CLIENT_REVIEW_EDIT_VALUE_LOOKAHEAD}",
    ):
        match = re.search(pattern, comment, re.IGNORECASE | re.DOTALL)
        if match:
            return _client_clean_edit_value(match.group(1), max_length)
    return None


def _client_clean_edit_value(value: str, max_length: int) -> str | None:
    protected_pattern = "|".join(re.escape(term) for term in CLIENT_REVIEW_PROTECTED_EDIT_TERMS)
    clean = re.split(rf"\s+(?:and\s+)?(?:{protected_pattern})\b", value, maxsplit=1, flags=re.IGNORECASE)[0]
    clean = re.sub(r"\s+", " ", clean).strip(" \"'`.,;:")
    clean = re.sub(r"^(?:please\s+)?(?:to|as|is|are)\s+", "", clean, flags=re.IGNORECASE).strip(" \"'`.,;:")
    if not clean:
        return None
    return clean[:max_length]


def _client_safe_itinerary_edit_note(comment: str) -> str | None:
    safe_parts = [
        part.strip()
        for part in re.split(r"(?<=[.!?])\s+|[;\n]+", mask_sensitive_customer_text(comment))
        if part.strip() and not _client_edit_mentions_protected_fields(part)
    ]
    safe_note = "; ".join(safe_parts).strip()
    if not safe_note:
        return None
    if len(safe_note) > 400:
        return f"{safe_note[:397]}..."
    return safe_note


def _client_edit_mentions_protected_fields(text: str) -> bool:
    for term in CLIENT_REVIEW_PROTECTED_EDIT_TERMS:
        if re.search(rf"\b{re.escape(term)}\b", text, re.IGNORECASE):
            return True
    return False


def _send_review_link_email(
    store: TravelStore,
    request: CorporateTravelRequest,
    change_summary: str | None,
    context: AuthContext,
    purpose: str,
) -> EmailEvent | None:
    to_email = _delivery_email_for_request(request)
    if not to_email or not request.client_review:
        audit_security_decision(store, context, "corporate.pipeline.review_link_blocked", "Client email or review link is missing.", purpose, "deny", request.id)
        return None
    event = send_resend_notification(
        request.id,
        request,
        NotificationRequest(
            kind="review_link",
            to=[to_email],
            note="Your requested itinerary edits are ready. Please review the refreshed dashboard options.",
            review_url=request.client_review.review_url,
            change_summary=change_summary,
            review_round=request.client_review.revision_round,
        ),
    )
    saved = store.save_email_event(event)
    audit_security_decision(store, context, "corporate.pipeline.review_revision_sent", saved.safe_message, purpose, "allow", request.id)
    return saved


def _client_review_change_summary(revision_round: int, comment: str) -> str:
    safe_comment = (_client_safe_itinerary_edit_note(comment) or "your itinerary corrections").rstrip(" .")
    if len(safe_comment) > 220:
        safe_comment = f"{safe_comment[:217]}..."
    protected_note = " Protected policy fields were left unchanged." if _client_edit_mentions_protected_fields(comment) else ""
    return f"Round {revision_round} options were regenerated around this request: {safe_comment}.{protected_note}"


def _client_review_response(request: CorporateTravelRequest) -> CorporateReviewResponse:
    plan = request.generated_plan
    review = request.client_review
    if not plan or not review:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Review link is not available")
    currency = _client_review_currency(request)
    options = [
        _client_review_option(request, option, index, currency)
        for index, option in enumerate(plan.travel_options[:CLIENT_REVIEW_OPTION_COUNT], start=1)
    ]
    if not options:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Client review options are not ready")
    return CorporateReviewResponse(
        request_id=request.id,
        traveler_name=request.traveller_details.traveler_name or "Traveler",
        company_name=request.company_details.company_name or "Company",
        route=f"{request.travel_details.origin or 'Origin pending'} to {request.travel_details.destination or 'Destination pending'}",
        depart_date=request.travel_details.depart_date,
        return_date=request.travel_details.return_date,
        status=review.status,
        revision_round=review.revision_round,
        expires_at=review.expires_at,
        submitted_at=review.submitted_at,
        change_summary=review.change_summary,
        special_request_notice=_special_request_notice(request),
        options=options,
        history=request.client_review_history,
    )


def _client_review_option(request: CorporateTravelRequest, option: TravelOption, index: int, currency: str) -> CorporateReviewOption:
    plan = request.generated_plan
    flight = next((offer for offer in plan.flight_offers if offer.id == option.flight_offer_id), None) if plan else None
    hotel = plan.hotel_offers[index - 1] if plan and index - 1 < len(plan.hotel_offers) else None
    transfer = next((offer for offer in plan.ground_transfer_offers if offer.id == option.ground_transfer_offer_id), None) if plan else None
    plan_currency = request.budgets.currency or "USD"
    return CorporateReviewOption(
        option_index=index,
        option_name=option.option_name,
        flight_summary=option.flight_summary,
        hotel_summary=option.hotel_summary,
        transfer_summary=option.transfer_summary,
        estimated_cost=_client_review_amount(option.estimated_cost, plan_currency, currency),
        currency=currency,
        maximum_budget=None,
        budget_delta=None,
        policy_status=option.policy_status,
        recommendation_reason=_customer_safe_review_reason(request, option, flight, hotel, transfer),
        reasoning_source_label=_review_reasoning_source_label(request),
        pros=option.pros,
        cons=option.cons,
        flight=_client_review_flight(flight, currency) if flight else None,
        hotel=_client_review_hotel(request, hotel, flight, currency) if hotel else None,
        transfer=_client_review_transfer(transfer, currency) if transfer else None,
    )


def _client_review_currency(request: CorporateTravelRequest) -> str:
    return origin_city_currency(request.travel_details.origin, request.budgets.currency or "USD")


def _client_review_amount(amount: int, source_currency: str | None, display_currency: str) -> int:
    return convert_planning_amount(int(amount or 0), source_currency or "USD", display_currency)  # type: ignore[arg-type]


def _client_review_flight(flight, display_currency: str) -> CorporateReviewFlight:
    return CorporateReviewFlight(
        id=flight.id,
        airline=flight.airline,
        airline_code=_flight_airline_code(flight),
        airline_logo_url=_flight_airline_logo_url(flight),
        summary=flight.summary,
        outbound=flight.outbound,
        return_leg=flight.return_leg,
        cabin=flight.cabin,
        total_amount=_client_review_amount(flight.total_amount, flight.currency, display_currency),
        currency=display_currency,
        source=flight.source,
        notes=flight.notes,
    )


def _client_review_hotel(request: CorporateTravelRequest, hotel, flight, display_currency: str) -> CorporateReviewHotel:
    return CorporateReviewHotel(
        id=hotel.id,
        name=hotel.name,
        summary=hotel.summary,
        address=hotel.address,
        check_in=hotel.check_in,
        check_out=hotel.check_out,
        check_in_starts_at=hotel.check_in_starts_at,
        checkout_time=hotel.checkout_time,
        room_notes=hotel.room_notes,
        cancellation_notes=hotel.cancellation_notes,
        unsent_special_requests=_review_hotel_special_requests(request, hotel, flight),
        total_amount=_client_review_amount(hotel.total_amount, hotel.currency, display_currency),
        currency=display_currency,
        image_url=hotel.image_url,
    )


def _client_review_transfer(transfer: CorporateGroundTransferOffer, display_currency: str) -> CorporateReviewTransfer:
    return CorporateReviewTransfer(
        id=transfer.id,
        pickup_airport_code=transfer.pickup_airport_code,
        pickup_time=transfer.pickup_time,
        dropoff_label=transfer.dropoff_label,
        dropoff_address=transfer.dropoff_address,
        service_type=transfer.service_type,
        vehicle_type=transfer.vehicle_type,
        passengers=transfer.passengers,
        baggage=transfer.baggage,
        total_amount=_client_review_amount(transfer.total_amount, transfer.currency, display_currency),
        currency=display_currency,
        cancellation_notes=transfer.cancellation_notes,
        notes=transfer.notes,
    )


def _special_request_notice(request: CorporateTravelRequest) -> str:
    items = _unique_text([
        *request.special_requests,
        *_generated_review_special_requests(request),
    ])
    if not items:
        return "No hotel or transfer service requests were captured for this itinerary."
    return "Special requests to include: " + "; ".join(
        mask_sensitive_customer_text(item) for item in items
    )


def _customer_safe_review_reason(request: CorporateTravelRequest, option: TravelOption, flight, hotel, transfer) -> str:
    context_reasons = _review_context_reasons(request, flight, hotel, transfer)
    if option.option_name in {"Best tier fit", "Best within budget"}:
        option_reason = "Option tradeoff: best hotel tier fit while still covering the flight, hotel, and airport transfer."
    elif option.option_name == "Fastest route":
        option_reason = "Option tradeoff: prioritizes arrival reliability when timing matters more than the lowest estimate."
    elif option.option_name == "Comfort-focused option":
        option_reason = "Option tradeoff: adds rest, luggage, and transfer buffers for long-haul travel or high-stakes meetings."
    else:
        option_reason = "Option tradeoff: balances cost, schedule, hotel fit, and airport transfer coverage."
    return "\n".join([*context_reasons, option_reason])


def _review_context_reasons(request: CorporateTravelRequest, flight, hotel, transfer) -> list[str]:
    reasons: list[str] = []
    source_label = _review_reasoning_source_label(request)
    route = _safe_route_text(request)
    if source_label == "Reason from your previous journeys":
        reasons.append(f"Route fit: previous journeys point to this type of {route} itinerary.")
    else:
        reasons.append(f"Route fit: the submitted request asks for a {route} itinerary.")

    preference_parts: list[str] = []
    preferred_airline = request.preferences.preferred_airline
    if preferred_airline and flight:
        if preferred_airline.casefold() in flight.airline.casefold():
            preference_parts.append("selected airline matches the captured airline preference")
        else:
            preference_parts.append("flight was weighed against the captured airline preference")

    hotel_context = request.preferences.preferred_hotel_area or request.preferences.hotel_preference or request.travel_details.meeting_location
    if hotel_context and hotel:
        preference_parts.append("hotel keeps the stay close to the requested business area")

    if request.preferences.timing_preference and flight:
        preference_parts.append("flight timing was weighed against the timing preference")

    if transfer:
        preference_parts.append("airport transfer is included through to the selected hotel")

    if flight and flight.cabin != "economy" and _traveller_designation_supports_cabin(request):
        preference_parts.append(f"travel policy allows {flight.cabin.replace('_', ' ')}")

    if preference_parts:
        reasons.append(_sentence_from_parts("Recommendation fit", preference_parts))
    return reasons


def _safe_route_text(request: CorporateTravelRequest) -> str:
    origin = request.travel_details.origin or "origin"
    destination = request.travel_details.destination or "destination"
    return f"{origin} to {destination}"


def _sentence_from_parts(label: str, parts: list[str]) -> str:
    if len(parts) == 1:
        return f"{label}: {parts[0]}."
    return f"{label}: {', '.join(parts[:-1])}, and {parts[-1]}."


def _traveller_designation_supports_cabin(request: CorporateTravelRequest) -> bool:
    text = " ".join([
        (request.traveller_details.employee_band or "").lower(),
        (request.traveller_details.employee_level or "").lower(),
        (request.company_details.approval_band or "").lower(),
        (request.company_details.policy_tier or "").lower(),
    ])
    return any(
        marker in text
        for marker in (
            "manager",
            "director",
            "vp",
            "vice president",
            "head of",
            "ceo",
            "chief",
            "founder",
            "president",
            "chair",
            "band m",
        )
    )


def _flight_airline_code(flight) -> str | None:
    code = _note_value(flight.notes, "Airline code:")
    if code:
        return code.upper()
    provider = flight.provider or ""
    if provider.startswith(("duffel-api/", "duffel-mcp/")):
        suffix = provider.split("/", 1)[1].strip()
        if suffix and suffix != "live" and len(suffix) <= 3:
            return suffix.upper()
    return None


def _flight_airline_logo_url(flight) -> str | None:
    return _note_value(flight.notes, "Airline logo:")


def _note_value(notes: list[str], prefix: str) -> str | None:
    for note in notes:
        if note.startswith(prefix):
            value = note.removeprefix(prefix).strip()
            return value or None
    return None


def _review_reasoning_source_label(request: CorporateTravelRequest) -> str:
    plan = request.generated_plan
    notes = " ".join(plan.agent_notes if plan else [])
    if request.preferences.past_hotel_preference or "Traveller history note:" in notes:
        return "Reason from your previous journeys"
    return "Reason from the submitted request"


def _review_hotel_special_requests(request: CorporateTravelRequest, hotel, flight) -> list[str]:
    early_check_in = _early_check_in_request(hotel, flight)
    return _unique_text([
        *hotel.unsent_special_requests,
        early_check_in,
        _breakfast_request(request, hotel),
        *_premium_service_requests(request, hotel, flight),
    ])


def _generated_review_special_requests(request: CorporateTravelRequest) -> list[str]:
    plan = request.generated_plan
    if not plan:
        return []
    requests: list[str] = []
    for index, option in enumerate(plan.travel_options[:CLIENT_REVIEW_OPTION_COUNT], start=0):
        flight = next((offer for offer in plan.flight_offers if offer.id == option.flight_offer_id), None)
        hotel = plan.hotel_offers[index] if index < len(plan.hotel_offers) else None
        if hotel:
            requests.extend(_review_hotel_special_requests(request, hotel, flight))
    return _unique_text(requests)


def _early_check_in_request(hotel, flight) -> str | None:
    arrival = _arrival_datetime_from_flight(flight)
    check_in_minutes = _time_minutes(hotel.check_in_starts_at)
    if not arrival or check_in_minutes is None:
        return None
    if hotel.check_in and arrival.date() != hotel.check_in:
        return None
    arrival_minutes = arrival.hour * 60 + arrival.minute
    if arrival_minutes >= check_in_minutes:
        return None
    arrival_label = arrival.strftime("%H:%M")
    check_in_label = hotel.check_in_starts_at or "standard hotel check-in"
    return f"Request early hotel check-in because destination arrival is {arrival_label} and standard hotel check-in starts at {check_in_label}."


def _breakfast_request(request: CorporateTravelRequest, hotel) -> str | None:
    text = " ".join([
        request.preferences.meal_preference or "",
        request.preferences.hotel_preference or "",
        request.preferences.preferred_hotel_area or "",
        hotel.room_notes or "",
        hotel.summary or "",
        " ".join(request.special_requests),
    ]).lower()
    if "breakfast" not in text:
        return None
    return "Request breakfast-inclusive rate or breakfast package for the stay."


def _premium_service_requests(request: CorporateTravelRequest, hotel, flight) -> list[str]:
    text = " ".join([
        request.traveller_details.employee_band or "",
        request.traveller_details.employee_level or "",
        request.company_details.approval_band or "",
        request.company_details.policy_tier or "",
        flight.cabin if flight else "",
        hotel.summary or "",
        hotel.room_notes or "",
        " ".join(request.special_requests),
    ]).lower()
    premium_context = any(marker in text for marker in ("ceo", "chief", "founder", "president", "chair", "executive", "business", "first", "premium", "suite", "5-star"))
    if not premium_context:
        return []
    return [
        "Request luggage assistance on arrival and departure.",
        "Request dedicated guest-relations or butler support if available for the premium stay.",
    ]


def _arrival_datetime_from_flight(flight) -> datetime | None:
    if not flight:
        return None
    matches = ISO_DATETIME_PATTERN.findall(flight.outbound or "")
    if not matches:
        return None
    try:
        return datetime.fromisoformat(matches[-1])
    except ValueError:
        return None


def _time_minutes(value: str | None) -> int | None:
    match = re.search(r"(\d{1,2}):(\d{2})", value or "")
    if not match:
        return None
    hour = int(match.group(1))
    minute = int(match.group(2))
    if hour > 23 or minute > 59:
        return None
    return hour * 60 + minute


def _unique_text(items: list[str | None]) -> list[str]:
    seen: set[str] = set()
    unique: list[str] = []
    for item in items:
        value = mask_sensitive_customer_text(item or "").strip()
        if not value:
            continue
        key = value.casefold()
        if key in seen:
            continue
        seen.add(key)
        unique.append(value)
    return unique


def _complete_request_from_client_approval_agent(
    store: TravelStore,
    existing: CorporateTravelRequest,
    approval: CorporateClientApprovalRequest,
    context: AuthContext,
    purpose: str,
) -> CorporateTravelRequest:
    if not existing.generated_plan or existing.status in {"Cancelled", "Pending Details"}:
        audit_security_decision(store, context, "corporate.pipeline.client_approval_blocked", "Approved itinerary is not ready.", purpose, "deny", existing.id)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Approved itinerary is not ready")

    selected_index = _selected_option_index(existing, approval)
    selected_option = existing.generated_plan.travel_options[selected_index]
    selected_hotel_id = existing.generated_plan.hotel_offers[selected_index].id if selected_index < len(existing.generated_plan.hotel_offers) else existing.generated_plan.selected_hotel_offer_id
    selected_transfer_id = selected_option.ground_transfer_offer_id or existing.generated_plan.selected_ground_transfer_offer_id
    now = datetime.now(timezone.utc)
    review = existing.client_review
    review_history = existing.client_review_history
    if review:
        review = review.model_copy(update={
            "status": "Approved",
            "selected_option_index": selected_index + 1,
            "submitted_at": now,
        })
        review_history = [
            *review_history,
            CorporateClientReviewEvent(
                action="approved",
                revision_round=review.revision_round,
                selected_option_index=selected_index + 1,
                created_at=now,
            ),
        ]
    final_plan = existing.generated_plan.model_copy(update={
        "selected_flight_offer_id": selected_option.flight_offer_id or existing.generated_plan.selected_flight_offer_id,
        "selected_hotel_offer_id": selected_hotel_id,
        "selected_ground_transfer_offer_id": selected_transfer_id,
        "customer_itinerary_draft": _customer_itinerary_from_selected_option(existing, selected_option),
        "approval_status": "Completed",
    })
    completed = existing.model_copy(update={
        "generated_plan": final_plan,
        "status": "Completed",
        "approval_status": "Received",
        "client_review": review,
        "client_review_history": review_history,
        "updated_at": now,
    })
    saved = store.save_corporate_request(completed)
    audit_security_decision(store, context, "corporate.pipeline.client_approved", f"Client approved itinerary option {selected_index + 1}.", purpose, "allow", saved.id)
    _send_final_itinerary_agent(store, saved, context, purpose)
    return saved


@app.get("/api/corporate/requests/{request_id}/export.xlsx")
def download_corporate_request_excel(
    request_id: str,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> StreamingResponse:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.request_export.denied")
    request = _load_authorized_corporate_request(store, context, request_id, purpose)
    if request.status not in {"Finalized", "Completed"} or not request.generated_plan:
        audit_security_decision(store, context, "corporate.request_export.blocked", "Final itinerary is not ready.", purpose, "deny", request.id)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Final itinerary is not ready")
    audit_security_decision(store, context, "corporate.request_export.allowed", "Corporate request Excel export generated.", purpose, "allow", request.id)
    return StreamingResponse(
        BytesIO(_corporate_request_workbook_bytes(request)),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{request.id}-final-itinerary.xlsx"'},
    )


@app.get("/api/corporate/requests/{request_id}/export.pdf")
def download_corporate_request_pdf(
    request_id: str,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> StreamingResponse:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.request_pdf_export.denied")
    request = _load_authorized_corporate_request(store, context, request_id, purpose)
    if request.status not in {"Finalized", "Completed"} or not request.generated_plan:
        audit_security_decision(store, context, "corporate.request_pdf_export.blocked", "Final itinerary is not ready.", purpose, "deny", request.id)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Final itinerary is not ready")
    audit_security_decision(store, context, "corporate.request_pdf_export.allowed", "Corporate request PDF itinerary generated.", purpose, "allow", request.id)
    return StreamingResponse(
        BytesIO(_corporate_request_pdf_bytes(request)),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{request.id}-final-itinerary.pdf"'},
    )


@app.get("/api/corporate/requests/{request_id}/options/{option_index}.pdf")
def download_corporate_option_pdf(
    request_id: str,
    option_index: int,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> StreamingResponse:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.option_pdf_export.denied")
    request = _load_authorized_corporate_request(store, context, request_id, purpose)
    if option_index not in {1, 2, 3} or not request.generated_plan:
        audit_security_decision(store, context, "corporate.option_pdf_export.blocked", "Itinerary option PDF is not ready.", purpose, "deny", request.id)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Itinerary option PDF is not ready")
    audit_security_decision(store, context, "corporate.option_pdf_export.allowed", f"Client itinerary option {option_index} PDF generated.", purpose, "allow", request.id)
    return StreamingResponse(
        BytesIO(_corporate_option_pdf_bytes(request, option_index)),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{request.id}-option-{option_index}.pdf"'},
    )


@app.get("/api/corporate/excel-template")
def download_corporate_excel_template(
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> StreamingResponse:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.excel_template.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.excel_template.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc

    audit_security_decision(store, context, "corporate.excel_template.allowed", "Corporate Excel workbook template generated.", purpose, "allow")
    return StreamingResponse(
        BytesIO(_corporate_template_workbook_bytes()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="corporate_travel_requests_template.xlsx"'},
    )


@app.get("/api/corporate/request-form.pdf")
def download_client_request_form_pdf(
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> StreamingResponse:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.request_form.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.request_form.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc

    audit_security_decision(store, context, "corporate.request_form.allowed", "Client request PDF generated.", purpose, "allow")
    return StreamingResponse(
        BytesIO(_client_request_pdf_bytes()),
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="client_travel_request_form.pdf"'},
    )


@app.get("/api/corporate/request-form.docx")
def download_client_request_form_docx(
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> StreamingResponse:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.request_form.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.request_form.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc

    audit_security_decision(store, context, "corporate.request_form.allowed", "Client request Word form generated.", purpose, "allow")
    return StreamingResponse(
        BytesIO(_client_request_docx_bytes()),
        media_type=DOCX_CONTENT_TYPE,
        headers={"Content-Disposition": 'attachment; filename="client_travel_request_form.docx"'},
    )


@app.post("/api/corporate/upload-excel", response_model=CorporateImportResponse)
@app.post("/api/corporate/requests/upload-excel", response_model=CorporateImportResponse)
async def upload_corporate_excel(
    file: UploadFile = File(...),
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CorporateImportResponse:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.excel.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.excel.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    content = await file.read()
    filename = (file.filename or "").lower()
    content_type = (file.content_type or "").lower()
    try:
        if _is_pdf_upload(filename, content_type):
            request_rows = _corporate_request_rows_from_pdf(content)
            policy_rows = []
            employee_rows = []
            history_rows = []
            visa_rows = []
        elif _is_docx_upload(filename, content_type):
            request_rows = _corporate_request_rows_from_docx(content)
            policy_rows = []
            employee_rows = []
            history_rows = []
            visa_rows = []
        else:
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(content), data_only=True)
            request_rows, policy_rows, employee_rows, history_rows, visa_rows = _corporate_profile_workbook_rows(workbook, file.filename or "company_profile.xlsx")
        _validated_corporate_request_rows(request_rows)
    except TravelFormValidationError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=exc.user_message) from exc
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid travel form") from exc

    created_ids: list[str] = []
    now = datetime.now(timezone.utc)
    policy_count = store.save_reference_rows("company_policy", policy_rows)
    history_count = store.save_reference_rows("traveller_history", history_rows)
    employee_count = _save_employee_profiles_from_rows(store, history_rows, now)
    visa_count = store.save_reference_rows("visa_rules", visa_rows)
    for row in request_rows:
        corporate_request = request_with_component_dependencies(_request_with_effective_budget(_corporate_request_from_row(row, context, now)))
        saved_request = store.save_corporate_request(corporate_request)
        processed_request = _run_automated_pipeline(store, saved_request, context, purpose)
        created_ids.append(processed_request.id)
    audit_security_decision(store, context, "corporate.excel.allowed", "Corporate Excel workbook imported.", purpose, "allow")
    return CorporateImportResponse(
        request_count=len(created_ids),
        policy_count=policy_count,
        traveller_history_count=history_count,
        employee_profile_count=employee_count,
        visa_rule_count=visa_count,
        created_request_ids=created_ids,
    )


@app.post("/api/corporate/company-policy/upload-pdf", response_model=CompanyPolicyImportResponse)
async def upload_company_policy_pdf(
    file: UploadFile = File(...),
    x_company_name: str | None = Header(default=None),
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CompanyPolicyImportResponse:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.policy_pdf.denied")
    try:
        ensure_scope(context, "admin:summary", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.policy_pdf.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    content = await file.read()
    filename = file.filename or "company_policy.pdf"
    content_type = (file.content_type or "").lower()
    if not (filename.lower().endswith(".pdf") or content_type == "application/pdf"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid company policy PDF")
    try:
        rows = _company_policy_rows_from_pdf(content, x_company_name, filename)
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid company policy PDF") from exc
    count = store.save_reference_rows("company_policy", rows)
    _save_policy_group_from_pdf_rows(store, rows, context.email)
    audit_security_decision(store, context, "corporate.policy_pdf.allowed", "Company policy PDF imported.", purpose, "allow")
    return CompanyPolicyImportResponse(
        company_name=str(rows[0].get("company_name") or "Company pending"),
        policy_count=count,
        rules=[str(row.get("rule") or row.get("approval_rule") or row.get("policy_text") or "") for row in rows],
    )


@app.get("/api/corporate/companies", response_model=list[CompanyPipelineStatus])
def list_company_pipeline_statuses(
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> list[CompanyPipelineStatus]:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.companies.denied")
    try:
        ensure_scope(context, "admin:summary", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.companies.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    audit_security_decision(store, context, "corporate.companies.allowed", "Company pipeline statuses read.", purpose, "allow")
    return _company_pipeline_statuses(store)


@app.get("/api/corporate/admin/summary", response_model=CorporateAdminSummary)
def corporate_admin_summary(
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> CorporateAdminSummary:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.admin.summary.denied")
    try:
        ensure_scope(context, "admin:summary", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.admin.summary.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    audit_security_decision(store, context, "corporate.admin.summary.allowed", "Corporate aggregate summary read.", purpose, "allow")
    return store.corporate_admin_summary()


@app.get("/api/corporate/admin/email-events", response_model=list[EmailEvent])
def corporate_admin_email_events(
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> list[EmailEvent]:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.admin.email_events.denied")
    try:
        ensure_any_scope(context, {"admin:audit", "travel:plan"}, purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.admin.email_events.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    audit_security_decision(store, context, "corporate.admin.email_events.allowed", "Corporate email events read.", purpose, "allow")
    return store.list_email_events()


@app.get("/api/corporate/requests/{request_id}/pipeline-events", response_model=list[CorporatePipelineEvent])
def corporate_request_pipeline_events(
    request_id: str,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> list[CorporatePipelineEvent]:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.pipeline_events.denied")
    _load_authorized_corporate_request(store, context, request_id, purpose)
    audit_security_decision(store, context, "corporate.pipeline_events.allowed", "Corporate pipeline events read.", purpose, "allow", request_id)
    return _corporate_pipeline_events(store, request_id)


@app.get("/api/corporate/requests/{request_id}/pipeline-events/stream")
async def stream_corporate_request_pipeline_events(
    request_id: str,
    follow_seconds: int = 0,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> StreamingResponse:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.pipeline_events_stream.denied")
    _load_authorized_corporate_request(store, context, request_id, purpose)
    audit_security_decision(store, context, "corporate.pipeline_events_stream.allowed", "Corporate pipeline event stream opened.", purpose, "allow", request_id)

    async def event_stream():
        seen_event_ids: set[str] = set()
        follow_window = min(max(follow_seconds, 0), PIPELINE_STREAM_MAX_FOLLOW_SECONDS)
        deadline = asyncio.get_running_loop().time() + follow_window
        while True:
            for event in _corporate_pipeline_events(store, request_id):
                if event.id in seen_event_ids:
                    continue
                seen_event_ids.add(event.id)
                yield _server_sent_event("pipeline.event", event.model_dump(mode="json"))
            if follow_window <= 0 or asyncio.get_running_loop().time() >= deadline:
                break
            await asyncio.sleep(PIPELINE_STREAM_POLL_SECONDS)
        yield _server_sent_event("pipeline.complete", {"request_id": request_id, "event_count": len(seen_event_ids)})

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


def _corporate_pipeline_events(store: TravelStore, request_id: str) -> list[CorporatePipelineEvent]:
    request = store.get_corporate_request(request_id)
    if not request:
        return []
    events = [
        CorporatePipelineEvent(
            id=f"request:{request.id}:{request.updated_at.isoformat()}",
            source="request",
            stage="request_state",
            status=request.status,
            message=f"Request is currently {request.status}.",
            created_at=request.updated_at,
        )
    ]
    for audit_event in store.list_audit_events():
        if audit_event.trip_id != request_id:
            continue
        events.append(CorporatePipelineEvent(
            id=f"audit:{audit_event.id}",
            source="audit",
            stage=_pipeline_stage_from_audit(audit_event.event_type),
            status=_pipeline_status_from_audit(audit_event.decision),
            message=audit_event.message,
            created_at=_datetime_value(audit_event.created_at),
        ))
    for email_event in store.list_email_events(request_id):
        events.append(CorporatePipelineEvent(
            id=f"email:{email_event.id}",
            source="email",
            stage=f"email.{email_event.kind or 'unknown'}",
            status=email_event.status,
            message=email_event.safe_message,
            created_at=email_event.created_at,
            review_round=email_event.review_round,
        ))
    return sorted(events, key=lambda event: event.created_at)


def _pipeline_stage_from_audit(event_type: str) -> str:
    stage = PIPELINE_STAGE_BY_EVENT_TYPE.get(event_type)
    if stage:
        return stage
    if "review_link" in event_type:
        return "review_link_email"
    if "review_revision" in event_type:
        return "revision_generation"
    if "client_approved" in event_type:
        return "client_approval"
    if "final" in event_type:
        return "final_itinerary"
    if "pipeline" in event_type:
        return "automated_pipeline"
    return event_type


def _pipeline_status_from_audit(decision: str) -> str:
    if decision == "allow":
        return "completed"
    if decision == "deny":
        return "blocked"
    return decision or "recorded"


def _server_sent_event(event_name: str, payload: dict[str, Any]) -> str:
    return f"event: {event_name}\ndata: {json.dumps(payload, default=str)}\n\n"


def _datetime_value(value: datetime | str) -> datetime:
    if isinstance(value, datetime):
        return value
    return datetime.fromisoformat(str(value).replace("Z", "+00:00"))


@app.get("/api/travelers", response_model=list[TravelerProfile])
def list_travelers(
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> list[TravelerProfile]:
    purpose = protected_purpose(x_travel_purpose, store, context, "travelers.list.denied")
    try:
        ensure_any_scope(context, {"traveler:read", "admin:summary"}, purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "travelers.list.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    audit_security_decision(store, context, "travelers.list.allowed", "Traveler roster read.", purpose, "allow")
    return store.list_travelers()


@app.get("/api/travelers/{traveler_id}", response_model=TravelerProfile)
def get_traveler(
    traveler_id: str,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> TravelerProfile:
    purpose = protected_purpose(x_travel_purpose, store, context, "travelers.read.denied")
    try:
        ensure_any_scope(context, {"traveler:read", "admin:summary"}, purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "travelers.read.denied", str(exc), purpose, "deny", traveler_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    traveler = store.get_traveler(traveler_id)
    if traveler is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Not found")
    audit_security_decision(store, context, "travelers.read.allowed", "Traveler dossier read.", purpose, "allow", traveler.id)
    return traveler


@app.put("/api/travelers/{traveler_id}", response_model=TravelerProfile)
def save_traveler_profile(
    traveler_id: str,
    traveler: TravelerProfile,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> TravelerProfile:
    purpose = protected_purpose(x_travel_purpose, store, context, "travelers.write.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "travelers.write.denied", str(exc), purpose, "deny", traveler_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    existing = store.get_traveler(traveler_id)
    now = datetime.now(timezone.utc)
    saved = traveler.model_copy(update={
        "id": traveler_id,
        "created_at": existing.created_at if existing else traveler.created_at,
        "updated_at": now,
    })
    saved = store.save_traveler(saved)
    audit_security_decision(store, context, "travelers.write.allowed", "Traveler profile saved.", purpose, "allow", saved.id)
    return saved


@app.get("/api/policies", response_model=list[PolicyGroup])
def list_policies(
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> list[PolicyGroup]:
    purpose = protected_purpose(x_travel_purpose, store, context, "policies.list.denied")
    try:
        ensure_any_scope(context, {"policy:read", "admin:summary"}, purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "policies.list.denied", str(exc), purpose, "deny")
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    audit_security_decision(store, context, "policies.list.allowed", "Policy groups read.", purpose, "allow")
    return store.list_policy_groups()


@app.get("/api/policies/{policy_id}", response_model=PolicyGroup)
def get_policy(
    policy_id: str,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> PolicyGroup:
    purpose = protected_purpose(x_travel_purpose, store, context, "policies.read.denied")
    policy = _load_policy_group(store, context, policy_id, purpose)
    audit_security_decision(store, context, "policies.read.allowed", "Policy group read.", purpose, "allow", policy.id)
    return policy


@app.get("/api/policies/{policy_id}/versions", response_model=list[PolicyRevision])
def list_policy_versions(
    policy_id: str,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> list[PolicyRevision]:
    purpose = protected_purpose(x_travel_purpose, store, context, "policies.versions.denied")
    policy = _load_policy_group(store, context, policy_id, purpose)
    audit_security_decision(store, context, "policies.versions.allowed", "Policy revisions read.", purpose, "allow", policy.id)
    return policy.revisions


@app.get("/api/policies/{policy_id}/activity", response_model=list[PolicyActivityEvent])
def list_policy_activity(
    policy_id: str,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> list[PolicyActivityEvent]:
    purpose = protected_purpose(x_travel_purpose, store, context, "policies.activity.denied")
    _load_policy_group(store, context, policy_id, purpose)
    audit_security_decision(store, context, "policies.activity.allowed", "Policy activity read.", purpose, "allow", policy_id)
    return store.list_policy_activity(policy_id)


@app.post("/api/policies/{policy_id}/revisions/{revision_id}/approve", response_model=PolicyGroup)
def approve_policy_revision(
    policy_id: str,
    revision_id: str,
    action: RevisionActionRequest | None = None,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> PolicyGroup:
    purpose = protected_purpose(x_travel_purpose, store, context, "policies.approve.denied")
    revision, policy = _load_policy_revision(store, context, policy_id, revision_id, purpose)
    approved = revision.model_copy(
        update={
            "status": "Approved",
            "reviewer_comments": [*revision.reviewer_comments, action.comment] if action and action.comment else revision.reviewer_comments,
            "updated_at": datetime.now(timezone.utc),
        }
    )
    saved = store.save_policy_revision(policy.id, approved, context.email, "Policy revision approved for publication.")
    if saved is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Not found")
    audit_security_decision(store, context, "policies.approve.allowed", "Policy revision approved.", purpose, "allow", policy.id)
    return saved


@app.post("/api/policies/{policy_id}/revisions/{revision_id}/request-changes", response_model=PolicyGroup)
def request_policy_revision_changes(
    policy_id: str,
    revision_id: str,
    action: RevisionActionRequest | None = None,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> PolicyGroup:
    purpose = protected_purpose(x_travel_purpose, store, context, "policies.request_changes.denied")
    revision, policy = _load_policy_revision(store, context, policy_id, revision_id, purpose)
    changed = revision.model_copy(
        update={
            "status": "Changes Requested",
            "reviewer_comments": [*revision.reviewer_comments, action.comment or "Changes requested."] if action else [*revision.reviewer_comments, "Changes requested."],
            "updated_at": datetime.now(timezone.utc),
        }
    )
    saved = store.save_policy_revision(policy.id, changed, context.email, "Policy revision changes requested.")
    if saved is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Not found")
    audit_security_decision(store, context, "policies.request_changes.allowed", "Policy revision changes requested.", purpose, "allow", policy.id)
    return saved


@app.post("/api/corporate/requests/{request_id}/notifications", response_model=EmailEvent)
def send_corporate_request_notification(
    request_id: str,
    notification: NotificationRequest,
    x_travel_purpose: str | None = Header(default=None),
    context: AuthContext = Depends(require_auth),
    store: TravelStore = Depends(get_store),
) -> EmailEvent:
    purpose = protected_purpose(x_travel_purpose, store, context, "corporate.notification.denied")
    try:
        ensure_scope(context, "travel:plan", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "corporate.notification.denied", str(exc), purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    request = _load_authorized_corporate_request(store, context, request_id, purpose)
    if notification.kind == "final_itinerary" and (request.status not in {"Finalized", "Completed"} or not request.generated_plan):
        audit_security_decision(store, context, "corporate.notification.blocked", "Final itinerary is not ready.", purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Final itinerary is not ready")
    attachment = _corporate_request_pdf_bytes(request) if notification.attach_itinerary else None
    event = send_resend_notification(
        request_id,
        request,
        notification,
        attachment,
        f"{request.id}-final-itinerary.pdf" if attachment else None,
        "application/pdf" if attachment else None,
    )
    saved = store.save_email_event(event)
    audit_security_decision(store, context, "corporate.notification.recorded", saved.safe_message, purpose, "allow", request_id)
    return saved


@app.post("/api/mail/resend/webhook", response_model=EmailEvent)
async def record_resend_webhook(request: Request, store: TravelStore = Depends(get_store)) -> EmailEvent:
    payload = await request.body()
    _verify_resend_webhook(payload, request)
    try:
        event = ResendWebhookEvent.model_validate_json(payload)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid webhook") from exc
    data = event.data or {}
    if event.type == "email.received":
        email_id = _resend_inbound_email_id(data)
        existing = _existing_received_email_event(store, email_id)
        if existing:
            return existing
        return store.save_email_event(_record_received_email(event, store))
    tags = data.get("tags") if isinstance(data.get("tags"), dict) else {}
    request_id = tags.get("request_id") if isinstance(tags, dict) else None
    kind = tags.get("kind") if isinstance(tags, dict) else None
    review_round = data.get("review_round") or tags.get("review_round") if isinstance(tags, dict) else None
    recorded = EmailEvent(
        request_id=str(request_id) if request_id else None,
        kind=kind if kind in {"approval_request", "review_link", "document_update", "final_itinerary", "client_cancelled"} else None,
        status="received",
        to=[str(item) for item in data.get("to", [])] if isinstance(data.get("to"), list) else [],
        subject=str(data.get("subject") or ""),
        provider_message_id=str(data.get("email_id") or ""),
        safe_message=f"Resend webhook received: {event.type}",
        review_round=_optional_int(review_round),
    )
    return store.save_email_event(recorded)


def _record_received_email(event: ResendWebhookEvent, store: TravelStore) -> EmailEvent:
    data = event.data or {}
    email_id = _resend_inbound_email_id(data)
    to = [str(item) for item in data.get("to", [])] if isinstance(data.get("to"), list) else []
    subject = str(data.get("subject") or "Inbound travel form")
    safe_message = "Inbound email received but no travel form could be parsed. Review manually."
    request_id: str | None = None
    requester_email: str | None = None
    attachments: list[tuple[str, str, bytes]] = []
    text = ""
    try:
        received = _resend_received_email(email_id)
        requester_email = _inbound_sender_email(data, received)
        text = str(received.get("text") or received.get("body") or _html_to_text(str(received.get("html") or ""))).strip()
        received_subject = str(received.get("subject") or subject)
        context = demo_auth_context(os.getenv("TRAVEL_AI_INBOUND_OWNER_EMAIL", "demo.agent@unipro.com"))
        approved = _try_complete_inbound_approval(store, context, f"{received_subject}\n{text}")
        if approved:
            return _received_email_event_with_source(approved, email_id, to, subject)
        attachments = _resend_received_attachment_bytes(email_id, data.get("attachments"))
        for filename, content_type, content in attachments:
            if _is_pdf_upload(filename, content_type):
                selected_pdf = _try_complete_selected_option_pdf(store, context, content)
                if selected_pdf:
                    return _received_email_event_with_source(selected_pdf, email_id, to, subject)
        rows: list[dict[str, object]] = []
        for filename, content_type, content in attachments:
            if _is_pdf_upload(filename, content_type):
                rows = _corporate_request_rows_from_pdf(content)
                break
            if _is_docx_upload(filename, content_type):
                rows = _corporate_request_rows_from_docx(content)
                break
        if not rows and text:
            rows = _corporate_request_rows_from_text(text)
        if rows:
            _validated_corporate_request_rows(rows)
            now = datetime.now(timezone.utc)
            saved_requests = [
                _run_automated_pipeline(
                    store,
                    store.save_corporate_request(request_with_component_dependencies(_corporate_request_from_row(_row_with_requester_email(row, requester_email), context, now))),
                    context,
                    "process inbound client travel request form",
                )
                for row in rows
            ]
            request_id = saved_requests[0].id
            safe_message = (
                f"Inbound email received and request {saved_requests[0].id} entered the automated pipeline."
                if len(saved_requests) == 1
                else f"Inbound email received and {len(saved_requests)} requests entered the automated pipeline."
            )
    except TravelFormValidationError as exc:
        safe_message = f"Inbound email received but the travel form needs correction: {exc.user_message}"
        if requester_email:
            _send_inbound_form_rejection_email(store, requester_email, subject, exc.user_message)
        log_internal_issue(
            logging.getLogger(INTERNAL_LOGGER_NAME),
            "corporate.inbound.validation_failed",
            "Inbound travel form failed validation.",
            level=logging.WARNING,
            validation_errors=exc.errors,
        )
    except Exception:
        safe_message = "Inbound email received but could not be imported. Review manually."
    return EmailEvent(
        request_id=request_id,
        kind=None,
        status="received",
        to=to,
        subject=subject,
        provider_message_id=email_id,
        safe_message=safe_message,
        body_text=mask_sensitive_customer_text(text)[:6000] if text else None,
        attachment_names=[filename for filename, _, _ in attachments],
    )


def _resend_inbound_email_id(data: dict[str, Any]) -> str:
    return str(data.get("email_id") or data.get("id") or "")


def _existing_received_email_event(store: TravelStore, email_id: str) -> EmailEvent | None:
    if not email_id:
        return None
    for event in store.list_email_events():
        if event.provider_message_id == email_id and event.status == "received":
            return event
    return None


def _received_email_event_with_source(event: EmailEvent, email_id: str, to: list[str], subject: str) -> EmailEvent:
    return event.model_copy(update={
        "provider_message_id": event.provider_message_id or email_id,
        "to": event.to or to,
        "subject": event.subject or subject,
    })


def _send_inbound_form_rejection_email(store: TravelStore, to_email: str, subject: str, reason: str) -> EmailEvent:
    correction_request = CorporateTravelRequest(
        id=f"corp_rejected_{uuid4().hex[:12]}",
        status="Cancelled",
        requester_email=to_email,
        traveller_details=TravellerDetails(traveler_name="travel request form"),
        travel_details=TravelDetails(origin="Submitted form", destination="Needs correction"),
    )
    event = send_resend_notification(
        correction_request.id,
        correction_request,
        NotificationRequest(
            kind="document_update",
            to=[to_email],
            note=(
                "We could not process your travel request form because it needs a correction. "
                f"{reason} Please update the form and resend it."
            ),
        ),
    )
    saved = store.save_email_event(event)
    log_internal_issue(
        logging.getLogger(INTERNAL_LOGGER_NAME),
        "corporate.inbound.validation_reply",
        "Inbound form correction email was recorded.",
        level=logging.WARNING,
        request_id=correction_request.id,
        delivery_status=saved.status,
        recipient_domain=to_email.rsplit("@", 1)[-1].lower() if "@" in to_email else "",
    )
    return saved


def _inbound_sender_email(webhook_data: dict[str, Any], received: dict[str, Any]) -> str | None:
    for value in (webhook_data.get("from"), received.get("from"), received.get("sender")):
        email = _email_address(value)
        if email:
            return email
    return None


def _email_address(value: object) -> str | None:
    if isinstance(value, dict):
        value = value.get("email") or value.get("address") or value.get("from")
    if not value:
        return None
    _, email = parseaddr(str(value))
    return email.strip() or None


def _row_with_requester_email(row: dict[str, object], requester_email: str | None) -> dict[str, object]:
    if not requester_email:
        return row
    return {**row, "requester_email": requester_email}


def _is_pdf_upload(filename: str, content_type: str) -> bool:
    return filename.lower().endswith(".pdf") or content_type.lower() == "application/pdf"


def _is_docx_upload(filename: str, content_type: str) -> bool:
    return filename.lower().endswith(".docx") or content_type.lower() == DOCX_CONTENT_TYPE


def _try_complete_inbound_approval(store: TravelStore, context: AuthContext, text: str) -> EmailEvent | None:
    match = re.search(r"\b(corp_req_[a-z0-9]+)\b", text, flags=re.IGNORECASE)
    if not match or not re.search(r"\b(approve|approved|selected|go ahead)\b", text, flags=re.IGNORECASE):
        return None
    request = store.get_corporate_request(match.group(1))
    if not request or not request.generated_plan:
        return None
    option_match = re.search(r"\boption\s*([1-3])\b", text, flags=re.IGNORECASE)
    completed = _complete_request_from_client_approval_agent(
        store,
        request,
        CorporateClientApprovalRequest(selected_option_index=int(option_match.group(1)) if option_match else 1),
        context,
        "process inbound client itinerary approval",
    )
    to_email = _delivery_email_for_request(completed)
    return EmailEvent(
        request_id=completed.id,
        kind="final_itinerary",
        status="received",
        to=[to_email] if to_email else [],
        subject="Client approved itinerary",
        safe_message=f"Client approved itinerary for request {completed.id}; final itinerary pipeline completed.",
        body_text=mask_sensitive_customer_text(text)[:6000],
    )


def _try_complete_selected_option_pdf(store: TravelStore, context: AuthContext, content: bytes) -> EmailEvent | None:
    try:
        text = _extract_pdf_text(content)
    except Exception:
        return None
    request_match = re.search(r"\bRequest:\s*(corp_req_[a-z0-9]+)\b", text, flags=re.IGNORECASE)
    option_match = re.search(r"\bItinerary\s+Option\s+([1-3])\b", text, flags=re.IGNORECASE)
    if not request_match or not option_match:
        return None
    request = store.get_corporate_request(request_match.group(1))
    if not request or not request.generated_plan:
        return None
    completed = _complete_request_from_client_approval_agent(
        store,
        request,
        CorporateClientApprovalRequest(selected_option_index=int(option_match.group(1))),
        context,
        "process returned selected itinerary pdf",
    )
    to_email = _delivery_email_for_request(completed)
    return EmailEvent(
        request_id=completed.id,
        kind="final_itinerary",
        status="received",
        to=[to_email] if to_email else [],
        subject="Selected itinerary PDF received",
        safe_message=f"Selected itinerary PDF recognized for request {completed.id}; final itinerary pipeline completed.",
        body_text=mask_sensitive_customer_text(text)[:6000],
    )


def _resend_headers() -> dict[str, str]:
    api_key = os.getenv("RESEND_API_KEY")
    if not api_key:
        raise RuntimeError("Email provider is not configured")
    return {"Authorization": f"Bearer {api_key}"}


def _resend_json(email_id: str, paths: list[str]) -> Any:
    base_url = os.getenv("RESEND_API_BASE_URL", "https://api.resend.com").rstrip("/")
    last_error: Exception | None = None
    for path in paths:
        try:
            response = httpx.get(f"{base_url}{path}", headers=_resend_headers(), timeout=30.0)
            response.raise_for_status()
            body = response.json()
            if isinstance(body, dict) and isinstance(body.get("data"), dict):
                return body["data"]
            return body
        except Exception as exc:
            last_error = exc
    raise RuntimeError(f"Could not retrieve received email {email_id}") from last_error


def _resend_received_email(email_id: str) -> dict[str, Any]:
    if not email_id:
        raise RuntimeError("Missing received email id")
    body = _resend_json(email_id, [
        f"/emails/receiving/{email_id}",
        f"/inbound/emails/{email_id}",
        f"/api/emails/received/{email_id}",
    ])
    if not isinstance(body, dict):
        raise RuntimeError("Invalid received email response")
    return body


def _resend_received_attachment_bytes(email_id: str, webhook_attachments: Any) -> list[tuple[str, str, bytes]]:
    raw_items: Any = webhook_attachments
    fetched_remote = False
    if not isinstance(raw_items, list) or not raw_items:
        list_body = _resend_json(email_id, [
            f"/emails/receiving/{email_id}/attachments",
            f"/inbound/emails/{email_id}/attachments",
            f"/api/emails/received/{email_id}/attachments",
        ])
        raw_items = list_body.get("data") if isinstance(list_body, dict) else list_body
        fetched_remote = True
    attachment_items = raw_items if isinstance(raw_items, list) else []
    results: list[tuple[str, str, bytes]] = []
    for item in attachment_items:
        if not isinstance(item, dict):
            continue
        filename = str(item.get("filename") or "attachment")
        content_type = str(item.get("content_type") or item.get("contentType") or "")
        content = _attachment_content(item)
        if content is None and item.get("download_url"):
            content = _download_attachment_url(str(item["download_url"]))
        if content is None and item.get("id"):
            content = _received_attachment_detail(email_id, str(item["id"]))
        if content is not None:
            results.append((filename, content_type, content))
    if not results and isinstance(webhook_attachments, list) and webhook_attachments and not fetched_remote:
        return _resend_received_attachment_bytes(email_id, None)
    return results


def _received_attachment_detail(email_id: str, attachment_id: str) -> bytes | None:
    body = _resend_json(email_id, [
        f"/emails/receiving/{email_id}/attachments/{attachment_id}",
        f"/inbound/emails/{email_id}/attachments/{attachment_id}",
        f"/api/emails/received/{email_id}/attachments/{attachment_id}",
    ])
    if not isinstance(body, dict):
        return None
    content = _attachment_content(body)
    if content is not None:
        return content
    download_url = body.get("download_url")
    if download_url:
        return _download_attachment_url(str(download_url))
    return None


def _download_attachment_url(download_url: str) -> bytes:
    response = httpx.get(download_url, timeout=30.0)
    response.raise_for_status()
    return response.content


def _attachment_content(item: dict[str, Any]) -> bytes | None:
    encoded = item.get("content") or item.get("file")
    if not encoded:
        return None
    if isinstance(encoded, bytes):
        return encoded
    try:
        return base64.b64decode(str(encoded), validate=False)
    except Exception:
        return None


def _html_to_text(value: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "\n", value)).strip()


def _verify_resend_webhook(payload: bytes, request: Request) -> None:
    secret = os.getenv("RESEND_WEBHOOK_SECRET")
    if not secret:
        return
    svix_id = request.headers.get("svix-id")
    timestamp = request.headers.get("svix-timestamp")
    signature = request.headers.get("svix-signature")
    if not svix_id or not timestamp or not signature:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid webhook")
    secret_bytes = _decode_svix_secret(secret)
    signed_payload = b".".join([svix_id.encode("utf-8"), timestamp.encode("utf-8"), payload])
    expected = base64.b64encode(hmac.new(secret_bytes, signed_payload, hashlib.sha256).digest()).decode("ascii")
    signatures = [part.removeprefix("v1,") for part in signature.split()]
    if not any(hmac.compare_digest(expected, candidate) for candidate in signatures):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid webhook")


def _decode_svix_secret(secret: str) -> bytes:
    if not secret.startswith("whsec_"):
        return secret.encode("utf-8")
    try:
        return base64.b64decode(secret.removeprefix("whsec_"), validate=True)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid webhook") from exc


def _load_authorized_corporate_request(store: TravelStore, context: AuthContext, request_id: str, purpose: str) -> CorporateTravelRequest:
    request = store.get_corporate_request(request_id)
    if request is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Not found")
    is_admin = "admin:summary" in context.scopes
    if not is_admin and request.owner_id != context.user_id:
        audit_security_decision(store, context, "corporate.request.denied", "Corporate request access denied.", purpose, "deny", request_id)
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Not found")
    return request


def _load_policy_group(store: TravelStore, context: AuthContext, policy_id: str, purpose: str) -> PolicyGroup:
    try:
        ensure_any_scope(context, {"policy:read", "policy:write", "admin:summary"}, purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "policies.access.denied", str(exc), purpose, "deny", policy_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    policy = store.get_policy_group(policy_id)
    if policy is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Not found")
    return policy


def _load_policy_revision(store: TravelStore, context: AuthContext, policy_id: str, revision_id: str, purpose: str) -> tuple[PolicyRevision, PolicyGroup]:
    try:
        ensure_scope(context, "policy:write", purpose)
    except SecurityError as exc:
        audit_security_decision(store, context, "policies.write.denied", str(exc), purpose, "deny", policy_id)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden") from exc
    policy = _load_policy_group(store, context, policy_id, purpose)
    for revision in policy.revisions:
        if revision.id == revision_id:
            return revision, policy
    raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Not found")


def _sheet_rows(workbook: Any, sheet_name: str) -> list[dict[str, object]]:
    if sheet_name not in workbook.sheetnames:
        return []
    sheet = workbook[sheet_name]
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [_header_name(value) for value in values[0]]
    rows: list[dict[str, object]] = []
    for raw_row in values[1:]:
        row: dict[str, object] = {}
        for index, value in enumerate(raw_row):
            if index >= len(headers) or not headers[index] or value is None:
                continue
            if isinstance(value, datetime):
                value = value.date()
            row[headers[index]] = value
        if row:
            rows.append(row)
    return rows


def _sheet_rows_any(workbook: Any, *sheet_names: str) -> list[dict[str, object]]:
    aliases = {_header_name(name) for name in sheet_names}
    for sheet_name in workbook.sheetnames:
        if _header_name(sheet_name) in aliases:
            return _sheet_rows(workbook, sheet_name)
    return []


def _corporate_profile_workbook_rows(workbook: Any, filename: str) -> tuple[list[dict[str, object]], list[dict[str, object]], list[dict[str, object]], list[dict[str, object]], list[dict[str, object]]]:
    request_rows = _sheet_rows_any(workbook, "Travel Requests", "Requests")
    policy_rows = _sheet_rows_any(workbook, "Company Policy", "Policies", "Travel Policy")
    old_employee_rows = _sheet_rows_any(workbook, "Employee Details", "Employees")
    old_history_rows = _sheet_rows_any(workbook, "Traveller History", "Traveler History", "Travel History")
    old_visa_rows = _sheet_rows_any(workbook, "Visa Rules")

    profile_rows = _sheet_rows_any(workbook, "Traveler Profiles", "Traveller Profiles")
    preference_rows = _sheet_rows_any(workbook, "Travel Preferences")
    flight_rows = _sheet_rows_any(workbook, "Flight Preferences")
    visa_record_rows = _sheet_rows_any(workbook, "Visa Passport Records", "Visa Records", "Passport Records")
    hotel_rows = _sheet_rows_any(workbook, "Preferred Hotels")
    history_rows = _sheet_rows_any(workbook, "Past Travel History")
    hotel_insight_rows = _sheet_rows_any(workbook, "Hotel Insights")

    company_name = _workbook_company_name(filename, [*request_rows, *policy_rows, *old_employee_rows, *profile_rows])
    normalized_profiles = _normalized_traveler_profile_rows(
        profile_rows,
        preference_rows,
        flight_rows,
        visa_record_rows,
        history_rows,
        hotel_rows,
        hotel_insight_rows,
        company_name,
    )
    normalized_history = _normalized_past_history_rows(history_rows, profile_rows, company_name)
    normalized_visa = _normalized_visa_record_rows(visa_record_rows, profile_rows, company_name)
    return (
        [_row_with_company(row, company_name) for row in request_rows],
        [_row_with_company(row, company_name) for row in policy_rows],
        [_row_with_company(row, company_name) for row in old_employee_rows],
        [
            *[_row_with_company(row, company_name) for row in old_employee_rows],
            *[_row_with_company(row, company_name) for row in old_history_rows],
            *normalized_profiles,
            *normalized_history,
        ],
        [*[_row_with_company(row, company_name) for row in old_visa_rows], *normalized_visa],
    )


def _normalized_traveler_profile_rows(
    profiles: list[dict[str, object]],
    preferences: list[dict[str, object]],
    flights: list[dict[str, object]],
    visas: list[dict[str, object]],
    history: list[dict[str, object]],
    hotels: list[dict[str, object]],
    hotel_insights: list[dict[str, object]],
    company_name: str,
) -> list[dict[str, object]]:
    preferences_by_id = _rows_by_key(preferences, "traveler_id")
    flights_by_id = _rows_by_key(flights, "traveler_id")
    visas_by_id = _rows_grouped_by_key(visas, "traveler_id")
    history_by_id = _rows_grouped_by_key(history, "traveler_id")
    rows: list[dict[str, object]] = []
    for profile in profiles:
        traveler_id = _str(profile, "traveler_id")
        if not traveler_id:
            continue
        preference = preferences_by_id.get(_row_key(traveler_id), {})
        flight = flights_by_id.get(_row_key(traveler_id), {})
        visa = (visas_by_id.get(_row_key(traveler_id)) or [{}])[0]
        trips = history_by_id.get(_row_key(traveler_id), [])
        row = {
            "traveler_id": traveler_id,
            "traveler_name": _str(profile, "name", "traveler_name"),
            "traveler_email": _str(profile, "email", "traveler_email"),
            "company_name": _str(profile, "company_name", "company") or _company_from_email(_str(profile, "email")) or company_name,
            "profile_type": _str(profile, "employee_band", "traveler_band", "band", "profile_type"),
            "department": _str(profile, "department"),
            "employee_level": _str(profile, "employee_band", "traveler_band", "band", "profile_type"),
            "employee_band": _str(profile, "employee_band", "traveler_band", "band", "profile_type"),
            "location": _str(profile, "base_location"),
            "home_airport": _str(profile, "home_airport"),
            "phone": _str(profile, "phone"),
            "nationality": _str(profile, "nationality"),
            "passport_number": _str(visa, "passport_number"),
            "passport_expiry": _date_value(visa.get("passport_expiry")),
            "visa_status": _visa_status_text(visa),
            "visa_expiry": _date_value(visa.get("visa_expiry")),
            "preferred_airline": _str(flight, "preferred_airlines"),
            "flight_preference": _joined_values(flight, "layover_rule", "timing_preference", "cabin_rule", "notes"),
            "hotel_preference": _str(preference, "hotel_tier_preference") or _hotel_fit_note(hotels, profile),
            "preferred_hotel_area": _str(preference, "room_preference"),
            "meal_preference": _str(preference, "meal_preference"),
            "seat_preference": _str(preference, "seat_preference"),
            "timing_preference": _str(flight, "timing_preference"),
            "airport_transfer_needed": _transport_needs_transfer(_str(preference, "ground_transport")),
            "policy_notes": _str(profile, "approval_notes"),
            "notes": _joined_values(profile, "core_travel_behavior", "personal_preferences", "approval_notes"),
            "recent_trips": _recent_trip_summary(trips),
            "past_hotel_preference": _hotel_insight_note(hotel_insights, profile),
        }
        rows.append({key: value for key, value in row.items() if value not in {None, ""}})
    return rows


def _normalized_past_history_rows(history: list[dict[str, object]], profiles: list[dict[str, object]], company_name: str) -> list[dict[str, object]]:
    profiles_by_id = _rows_by_key(profiles, "traveler_id")
    rows: list[dict[str, object]] = []
    for item in history:
        traveler_id = _str(item, "traveler_id")
        profile = profiles_by_id.get(_row_key(traveler_id), {})
        row = {
            "traveler_id": traveler_id,
            "traveler_name": _str(item, "traveler_name") or _str(profile, "name"),
            "traveler_email": _str(profile, "email"),
            "company_name": _company_from_email(_str(profile, "email")) or company_name,
            "profile_type": _str(item, "profile_type") or _str(profile, "profile_type"),
            "destination": _str(item, "destination_city"),
            "destination_country": _str(item, "country"),
            "preferred_airline": _str(item, "airline"),
            "hotel_preference": _str(item, "hotel_name"),
            "preferred_hotel_area": _str(item, "hotel_category"),
            "past_trips": _joined_values(item, "destination_city", "country", "purpose"),
            "recent_trips": _joined_values(item, "destination_city", "country", "purpose"),
            "notes": _joined_values(item, "hotel_feedback_comment", "preference_inferred", "trip_outcome"),
            "meal_preference": _str(profile, "personal_preferences"),
            "seat_preference": None,
        }
        rows.append({key: value for key, value in row.items() if value not in {None, ""}})
    return rows


def _normalized_visa_record_rows(records: list[dict[str, object]], profiles: list[dict[str, object]], company_name: str) -> list[dict[str, object]]:
    profiles_by_id = _rows_by_key(profiles, "traveler_id")
    rows: list[dict[str, object]] = []
    for record in records:
        traveler_id = _str(record, "traveler_id")
        profile = profiles_by_id.get(_row_key(traveler_id), {})
        row = {
            "traveler_id": traveler_id,
            "traveler_name": _str(record, "traveler_name") or _str(profile, "name"),
            "traveler_email": _str(profile, "email"),
            "company_name": _company_from_email(_str(profile, "email")) or company_name,
            "from_country": _str(record, "nationality") or _str(profile, "nationality"),
            "destination_country": _str(record, "country_region"),
            "visa_required": "yes",
            "visa_status": _visa_status_text(record),
            "visa_expiry": _date_value(record.get("visa_expiry")),
            "passport_number": _str(record, "passport_number"),
            "passport_expiry": _date_value(record.get("passport_expiry")),
            "entry_type": _str(record, "entry_type"),
            "notes": _str(record, "notes"),
        }
        if row.get("destination_country"):
            rows.append({key: value for key, value in row.items() if value not in {None, ""}})
    return rows


def _corporate_request_rows_from_pdf(content: bytes) -> list[dict[str, object]]:
    text = _extract_pdf_text(content)
    return _corporate_request_rows_from_text(text)


def _corporate_request_rows_from_docx(content: bytes) -> list[dict[str, object]]:
    text = _extract_docx_text(content)
    return _corporate_request_rows_from_text(text)


def _corporate_request_rows_from_text(text: str) -> list[dict[str, object]]:
    sections = _travel_form_sections(text)
    rows = [_corporate_request_row_from_text(section) for section in sections]
    return rows


def _validated_corporate_request_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    errors: list[str] = []
    for index, row in enumerate(rows, start=1):
        prefix = f"Request {index}: " if len(rows) > 1 else ""
        errors.extend(f"{prefix}{error}" for error in _travel_form_validation_errors(row))
    if errors:
        raise TravelFormValidationError(errors)
    return rows


def _validate_corporate_request_route(request: CorporateTravelRequest) -> None:
    errors = _travel_form_validation_errors({
        "origin": request.travel_details.origin,
        "destination": request.travel_details.destination,
    })
    if errors:
        raise TravelFormValidationError(errors)


def _travel_form_validation_errors(row: dict[str, object]) -> list[str]:
    errors: list[str] = []
    for field, label in (("origin", "Origin city / airport"), ("destination", "Destination city / airport")):
        endpoint_error = _route_endpoint_validation_error(label, row.get(field))
        if endpoint_error:
            errors.append(endpoint_error)
    origin_key = _row_key(row.get("origin"))
    destination_key = _row_key(row.get("destination"))
    if origin_key and destination_key and origin_key == destination_key:
        errors.append("Origin and destination must be different city or airport values.")
    return errors


def _route_endpoint_validation_error(label: str, value: object) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    key = _row_key(text)
    if key not in COUNTRY_ONLY_ROUTE_KEYS:
        return None
    example = COUNTRY_ROUTE_EXAMPLES.get(key, "a specific city or airport")
    return f"{label} must be a city or airport, not a country. For {text}, use {example} and put the country in Destination country."


def _travel_form_sections(text: str) -> list[str]:
    normalized = text.replace("\r", "\n")
    starts = [match.start() for match in re.finditer(r"(?im)^\s*travell?er\s+name\s*:", normalized)]
    if len(starts) <= 1:
        return [normalized]
    starts.append(len(normalized))
    return [normalized[starts[index]:starts[index + 1]].strip() for index in range(len(starts) - 1)]


def _corporate_request_row_from_text(text: str) -> dict[str, object]:
    row: dict[str, object] = {
        "traveler_name": _pdf_field(text, "traveller name", "traveler name", "name"),
        "traveler_email": _pdf_field(text, "traveller email", "traveler email", "email"),
        "employee_id": _pdf_field(text, "employee id"),
        "employee_band": _pdf_field(text, "employee band", "traveller band", "traveler band", "policy band", "profile type", "employee level", "designation", "role", "traveller level", "traveler level"),
        "department": _pdf_field(text, "department"),
        "nationality": _pdf_field(text, "nationality"),
        "passport_expiry": _pdf_field(text, "passport expiry"),
        "origin": _pdf_field(text, "origin city / airport", "origin", "from"),
        "destination": _pdf_field(text, "destination city / airport", "destination", "to"),
        "destination_country": _pdf_field(text, "destination country", "country"),
        "depart_date": _pdf_field(text, "departure date", "depart date"),
        "return_date": _pdf_field(text, "return date"),
        "trip_purpose": _pdf_field(text, "purpose / meeting location", "travel purpose", "purpose"),
        "meeting_location": _pdf_field(text, "meeting office / location", "meeting location"),
        "cabin": _pdf_field(text, "cabin"),
        "flight_preference": _pdf_field(text, "flight preference"),
        "hotel_preference": _pdf_field(text, "hotel area / location preference", "hotel area / star rating", "hotel preference", "hotel location preference", "office or location preference", "office location preference", "preferred hotel area"),
        "include_outbound_flight": _pdf_field(text, "include outbound flight", "outbound flight needed", "origin flight needed", "include origin flight", "departure flight needed"),
        "include_return_flight": _pdf_field(text, "include return flight", "return flight needed", "return flight required"),
        "include_hotel": _pdf_field(text, "include hotel", "hotel needed", "book hotel", "hotel required"),
        "include_ground_transfer": _pdf_field(text, "include ground transfer", "ground transfer needed", "airport transfer needed", "cab needed"),
        "airport_transfer_needed": _pdf_field(text, "airport transfer needed"),
        "meal_preference": _pdf_field(text, "meal preference"),
        "special_requests": _pdf_field(text, "trip-specific requests", "trip specific requests", "special requests"),
    }
    row = {key: value for key, value in row.items() if value not in {None, ""}}
    if not row.get("destination_country"):
        inferred_country = _infer_destination_country(str(row.get("destination") or ""), str(row.get("meeting_location") or ""), str(row.get("hotel_preference") or ""))
        if inferred_country:
            row["destination_country"] = inferred_country
    if not row.get("traveler_name") or not (row.get("origin") and row.get("destination")):
        raise ValueError("PDF travel form is missing traveler or route details")
    return row


def _infer_destination_country(*values: str) -> str | None:
    text = " ".join(values).lower()
    city_country = {
        "san jose": "USA",
        "sanjose": "USA",
        "new york": "USA",
        "chicago": "USA",
        "berlin": "Germany",
        "munich": "Germany",
        "singapore": "Singapore",
        "london": "United Kingdom",
        "johannesburg": "South Africa",
        "dubai": "United Arab Emirates",
        "tokyo": "Japan",
    }
    for city, country in city_country.items():
        if city in text:
            return country
    return None


def _extract_pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if text.strip():
            return text
    except Exception:
        pass
    return _extract_simple_pdf_text(content)


def _extract_docx_text(content: bytes) -> str:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            document_xml = archive.read("word/document.xml")
    except (KeyError, zipfile.BadZipFile) as exc:
        raise ValueError("Invalid Word travel form") from exc

    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = ElementTree.fromstring(document_xml)
    lines: list[str] = []
    for paragraph in root.findall(".//w:p", namespace):
        parts = [node.text or "" for node in paragraph.findall(".//w:t", namespace)]
        line = "".join(parts).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def _extract_simple_pdf_text(content: bytes) -> str:
    parts = re.findall(rb"\(((?:\\.|[^\\)])*)\)\s*Tj", content)
    lines = []
    for part in parts:
        value = part.replace(rb"\(", b"(").replace(rb"\)", b")").replace(rb"\\", b"\\")
        lines.append(value.decode("utf-8", errors="ignore"))
    return "\n".join(lines)


def _pdf_field(text: str, *labels: str) -> str | None:
    normalized = text.replace("\r", "\n")
    for label in labels:
        pattern = re.compile(rf"(?im)^\s*{re.escape(label)}(?:\s*\([^:\n]*\))?\s*:\s*(.+?)\s*$")
        match = pattern.search(normalized)
        if match:
            value = match.group(1).strip(" _")
            return value or None
    return None


def _save_employee_profiles_from_rows(store: TravelStore, rows: list[dict[str, object]], now: datetime) -> int:
    saved = 0
    seen_emails: set[str] = set()
    for row in rows:
        email = _str(row, "traveler_email", "traveller_email", "employee_email", "email")
        name = _str(row, "traveler_name", "traveller_name", "employee_name", "name")
        if not email or not name:
            continue
        email_key = email.lower()
        if email_key in seen_emails:
            continue
        seen_emails.add(email_key)
        passport_expiry = _date_value(row.get("passport_expiry"))
        visa_expiry = _date_value(row.get("visa_expiry"))
        documents = [
            {
                "document_type": "passport",
                "label": "Passport",
                "status": _document_status(passport_expiry),
                "expires_at": passport_expiry,
                "redacted_value": _redacted(row.get("passport_number")),
            }
        ]
        visa_status = _str(row, "visa_status", "visa")
        if visa_status or visa_expiry:
            documents.append(
                {
                    "document_type": "visa",
                    "label": "Visa",
                    "status": _document_status(visa_expiry) if visa_expiry else "Needs Review",
                    "expires_at": visa_expiry,
                    "redacted_value": visa_status or "On file",
                }
            )
        traveler = TravelerProfile(
            id=f"traveler_{hashlib.sha1(email.lower().encode('utf-8')).hexdigest()[:12]}",
            name=name,
            email=email,
            company=_str(row, "company_name", "company") or "Company pending",
            department=_str(row, "department"),
            vip_level=_str(row, "vip_level", "employee_level"),
            status=_traveler_status(documents),
            location=_str(row, "location", "base_location", "city"),
            seat_preference=_str(row, "seat_preference"),
            meal_preference=_str(row, "meal_preference"),
            hotel_preference=_str(row, "hotel_preference", "hotel_notes", "preferred_hotel_area"),
            policy_notes=[note for note in [_str(row, "policy_notes", "notes")] if note],
            loyalty_programs=_loyalty_programs(row),
            documents=documents,
            recent_trips=[trip for trip in [_str(row, "recent_trips", "past_trips")] if trip],
            created_at=now,
            updated_at=now,
        )
        store.save_traveler(traveler)
        saved += 1
    return saved


def _company_policy_rows_from_pdf(content: bytes, company_name: str | None, filename: str) -> list[dict[str, object]]:
    text = _extract_pdf_text(content)
    extracted_company = company_name or _pdf_field(text, "company", "client", "organization", "organisation")
    company = (extracted_company or _company_from_filename(filename) or "Company pending").strip()
    allowed_cabins = _pdf_field(text, "allowed cabins", "cabin policy", "flight cabin")
    hotel_tier = _pdf_field(text, "hotel tier", "hotel star tier", "hotel stars")
    approval_rule = _pdf_field(text, "approval rule", "approval requirement", "approvals")
    policy_tier = _pdf_field(text, "policy tier", "tier")
    notes = _pdf_field(text, "notes", "policy notes")
    rules = _policy_rules_from_text(text)
    if allowed_cabins:
        rules.append(f"Allowed cabins: {allowed_cabins}")
    if hotel_tier:
        rules.append(f"Hotel tier: {hotel_tier}")
    if approval_rule:
        rules.append(f"Approval rule: {approval_rule}")
    if notes:
        rules.append(f"Notes: {notes}")
    row = {
        "company_name": company,
        "policy_tier": policy_tier or "Standard",
        "allowed_cabins": _cabin_list_from_text(allowed_cabins),
        "hotel_tier": hotel_tier,
        "approval_rule": approval_rule,
        "notes": notes,
        "policy_text": _compact_text(text)[:1600],
        "rule": "; ".join(dict.fromkeys(rule for rule in rules if rule))[:1600],
        "source_file": filename,
    }
    if not (row["rule"] or row["policy_text"]):
        raise ValueError("Company policy PDF has no readable policy text")
    return [row]


def _save_policy_group_from_pdf_rows(store: TravelStore, rows: list[dict[str, object]], actor_email: str) -> None:
    if not rows:
        return
    company = str(rows[0].get("company_name") or "Company pending")
    policy_id = f"policy_company_{hashlib.sha1(company.lower().encode('utf-8')).hexdigest()[:12]}"
    active_rules = []
    for row in rows:
        for label, value in [
            ("Allowed cabins", row.get("allowed_cabins")),
            ("Hotel tier", row.get("hotel_tier")),
            ("Approval rule", row.get("approval_rule")),
            ("Policy notes", row.get("notes") or row.get("rule")),
        ]:
            if value:
                active_rules.append({"label": label, "value": str(value), "status": "Active"})
    if not active_rules:
        active_rules.append({"label": "Policy text", "value": str(rows[0].get("policy_text") or "Uploaded policy PDF"), "status": "Active"})
    policy = PolicyGroup(
        id=policy_id,
        client_name=company,
        business_unit="Corporate Travel",
        status="Active",
        active_rules=active_rules,
        compliance_score=100,
    )
    store.save_policy_group(policy)
    store.save_policy_activity(
        PolicyActivityEvent(policy_id=policy.id, actor=actor_email, activity="Imported company policy PDF.", status="SUCCESS")
    )


def _company_pipeline_statuses(store: TravelStore) -> list[CompanyPipelineStatus]:
    history_rows = store.list_reference_rows("traveller_history")
    visa_rows = store.list_reference_rows("visa_rules")
    policy_rows = store.list_reference_rows("company_policy")
    companies: dict[str, dict[str, object]] = {}

    def bucket(company_name: str | None) -> dict[str, object]:
        name = (company_name or "Company pending").strip() or "Company pending"
        return companies.setdefault(name, {"travellers": set(), "history": 0, "visa": 0, "policy": 0})

    for row in history_rows:
        data = bucket(_row_company(row))
        data["history"] = int(data["history"]) + 1
        traveller_ref = _str(row, "traveler_email", "traveller_email", "email", "employee_email", "traveler_id", "employee_id", "traveler_name", "name")
        if traveller_ref:
            travellers = data["travellers"]
            assert isinstance(travellers, set)
            travellers.add(_row_key(traveller_ref))

    for row in visa_rows:
        data = bucket(_row_company(row))
        data["visa"] = int(data["visa"]) + 1
        traveller_ref = _str(row, "traveler_email", "traveller_email", "email", "traveler_id", "traveler_name")
        if traveller_ref:
            travellers = data["travellers"]
            assert isinstance(travellers, set)
            travellers.add(_row_key(traveller_ref))

    for row in policy_rows:
        data = bucket(_row_company(row))
        data["policy"] = int(data["policy"]) + 1

    statuses: list[CompanyPipelineStatus] = []
    for company_name, data in sorted(companies.items()):
        travellers = data["travellers"]
        assert isinstance(travellers, set)
        traveler_count = len(travellers)
        policy_count = int(data["policy"])
        statuses.append(
            CompanyPipelineStatus(
                company_name=company_name,
                traveler_count=traveler_count,
                traveler_list_status="Updated" if traveler_count else "Missing",
                policy_status="Uploaded" if policy_count else "Missing",
                policy_count=policy_count,
                visa_record_count=int(data["visa"]),
                history_row_count=int(data["history"]),
            )
        )
    return statuses


def _enrich_request_from_reference_data(
    request: CorporateTravelRequest,
    history_rows: list[dict[str, object]],
    visa_rows: list[dict[str, object]],
    traveler_matches: list[dict[str, object]] | None = None,
) -> CorporateTravelRequest:
    matches = traveler_matches if traveler_matches is not None else _matching_traveler_rows(request, history_rows)
    if not matches:
        return request
    merged = _merged_reference_row(matches)
    visa = _matching_visa_reference(request, visa_rows, matches) or {}

    traveller = request.traveller_details
    company = request.company_details
    travel = request.travel_details
    preferences = request.preferences
    traveller_update = {
        "traveler_name": traveller.traveler_name or _str(merged, "traveler_name", "name", "employee_name"),
        "traveler_email": traveller.traveler_email or _str(merged, "traveler_email", "email", "employee_email"),
        "phone": traveller.phone or _str(merged, "phone", "phone_number", "mobile"),
        "employee_id": traveller.employee_id or _str(merged, "employee_id", "traveler_id"),
        "employee_level": traveller.employee_level or _str(merged, "employee_band", "traveler_band", "profile_type", "employee_level"),
        "employee_band": traveller.employee_band or _str(merged, "employee_band", "traveler_band", "profile_type", "employee_level"),
        "department": traveller.department or _str(merged, "department"),
        "nationality": traveller.nationality or _str(visa, "from_country", "nationality") or _str(merged, "nationality"),
        "passport_number": traveller.passport_number or _str(visa, "passport_number") or _str(merged, "passport_number"),
        "passport_expiry": traveller.passport_expiry or _date_value(visa.get("passport_expiry")) or _date_value(merged.get("passport_expiry")),
        "visa_status": traveller.visa_status or _visa_status_text(visa) or _str(merged, "visa_status", "visa"),
        "visa_expiry": traveller.visa_expiry or _date_value(visa.get("visa_expiry")) or _date_value(merged.get("visa_expiry")),
        "medical_notes": traveller.medical_notes or _str(merged, "medical_notes"),
        "accessibility_notes": traveller.accessibility_notes or _str(merged, "accessibility_notes"),
    }
    company_update = {
        "company_name": company.company_name or _row_company(merged),
        "cost_center": company.cost_center or _str(merged, "cost_center"),
        "approving_manager": company.approving_manager or _str(merged, "approving_manager", "manager"),
        "approval_manager_email": company.approval_manager_email or _str(merged, "approval_manager_email", "manager_email"),
        "approval_band": company.approval_band or _str(merged, "approval_band", "approver_band"),
        "policy_tier": company.policy_tier or _str(merged, "policy_tier", "approval_band", "profile_type"),
    }
    travel_update = {
        "origin": travel.origin or _str(merged, "home_airport", "origin", "base_location", "location"),
        "destination": travel.destination,
        "destination_country": travel.destination_country,
        "depart_date": travel.depart_date,
        "return_date": travel.return_date,
        "trip_purpose": travel.trip_purpose,
        "meeting_location": travel.meeting_location,
        "flexible_dates": travel.flexible_dates,
        "include_outbound_flight": travel.include_outbound_flight,
        "include_return_flight": travel.include_return_flight,
        "include_hotel": travel.include_hotel,
        "include_ground_transfer": travel.include_ground_transfer,
        "travelers": travel.travelers,
        "cabin": travel.cabin,
    }
    preference_update = {
        "preferred_airline": preferences.preferred_airline or _str(merged, "preferred_airline", "preferred_airlines"),
        "flight_preference": preferences.flight_preference or _str(merged, "flight_preference", "layover_rule"),
        "hotel_preference": preferences.hotel_preference or _str(merged, "hotel_preference", "hotel_notes"),
        "preferred_hotel_area": preferences.preferred_hotel_area or _str(merged, "preferred_hotel_area"),
        "hotel_star_rating": preferences.hotel_star_rating or _str(merged, "hotel_star_rating"),
        "past_hotel_preference": preferences.past_hotel_preference or _str(merged, "past_hotel_preference", "hotel_name"),
        "airport_transfer_needed": preferences.airport_transfer_needed or _bool_value(merged.get("airport_transfer_needed")),
        "meal_preference": preferences.meal_preference or _str(merged, "meal_preference"),
        "seat_preference": preferences.seat_preference or _str(merged, "seat_preference"),
        "timing_preference": preferences.timing_preference or _str(merged, "timing_preference"),
    }
    return request.model_copy(
        update={
            "traveller_details": traveller.model_copy(update=traveller_update),
            "company_details": company.model_copy(update=company_update),
            "travel_details": travel.model_copy(update=travel_update),
            "preferences": preferences.model_copy(update=preference_update),
        }
    )


def _has_traveler_profile_rows(history_rows: list[dict[str, object]]) -> bool:
    return any(
        _str(row, "traveler_email", "traveller_email", "email", "employee_email", "employee_id", "traveler_id", "traveler_name", "traveller_name", "employee_name", "name")
        for row in history_rows
    )


def _traveler_profile_warning(history_rows: list[dict[str, object]], traveler_matches: list[dict[str, object]]) -> str | None:
    if traveler_matches or not _has_traveler_profile_rows(history_rows):
        return None
    return NEW_TRAVELER_PROFILE_WARNING


def _plan_with_profile_warning(plan: CorporateTravelPlan, warning: str | None) -> CorporateTravelPlan:
    return _plan_with_agent_note(plan, warning)


def _plan_with_agent_note(plan: CorporateTravelPlan, note: str | None) -> CorporateTravelPlan:
    if not note or note in plan.agent_notes:
        return plan
    return plan.model_copy(update={"agent_notes": [note, *plan.agent_notes]})


def _document_status(expiry: date | None) -> str:
    if not expiry:
        return "Needs Review"
    days = (expiry - date.today()).days
    if days < 0:
        return "Missing"
    if days <= 180:
        return "Expiring Soon"
    return "Ready"


def _rows_by_key(rows: list[dict[str, object]], key: str) -> dict[str, dict[str, object]]:
    mapped: dict[str, dict[str, object]] = {}
    for row in rows:
        value = _row_key(row.get(key))
        if value:
            mapped.setdefault(value, row)
    return mapped


def _rows_grouped_by_key(rows: list[dict[str, object]], key: str) -> dict[str, list[dict[str, object]]]:
    grouped: dict[str, list[dict[str, object]]] = {}
    for row in rows:
        value = _row_key(row.get(key))
        if value:
            grouped.setdefault(value, []).append(row)
    return grouped


def _row_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def _workbook_company_name(filename: str, rows: list[dict[str, object]]) -> str:
    for row in rows:
        company = _row_company(row)
        if company and company != "Company pending":
            return company
        inferred = _company_from_email(_str(row, "traveler_email", "traveller_email", "email", "employee_email"))
        if inferred:
            return inferred
    return _company_from_filename(filename) or "Company pending"


def _row_company(row: dict[str, object]) -> str:
    return _str(row, "company_name", "company", "client_name", "organization", "organisation") or _company_from_email(_str(row, "traveler_email", "traveller_email", "email", "employee_email")) or "Company pending"


def _row_with_company(row: dict[str, object], company_name: str) -> dict[str, object]:
    if _row_company(row) != "Company pending":
        return row
    return {**row, "company_name": company_name}


def _company_from_email(email: str | None) -> str | None:
    if not email or "@" not in email:
        return None
    domain = email.split("@", 1)[1].split(".", 1)[0]
    if not domain:
        return None
    return domain.replace("-", " ").replace("_", " ").title()


def _company_from_filename(filename: str) -> str | None:
    stem = filename.rsplit("/", 1)[-1].rsplit(".", 1)[0]
    clean = re.sub(r"(?i)(company|corporate|travel|profile|dataset|revised|policy|context|workbook|forms?)", " ", stem)
    clean = re.sub(r"[^A-Za-z0-9]+", " ", clean).strip()
    return clean.title() if clean else None


def _visa_status_text(row: dict[str, object]) -> str | None:
    parts: list[str] = []
    for part in [
        _str(row, "visa_status", "status"),
        _str(row, "visa_permit_type", "visa_type"),
        _str(row, "country_region", "destination_country", "country"),
    ]:
        if not part:
            continue
        if any(part == existing or part in existing for existing in parts):
            continue
        parts.append(part)
    return " - ".join(part for part in parts if part) or None


def _joined_values(row: dict[str, object], *keys: str) -> str | None:
    values = [_str(row, key) for key in keys]
    return "; ".join(value for value in values if value) or None


def _recent_trip_summary(rows: list[dict[str, object]]) -> str | None:
    summaries = []
    for row in rows[:4]:
        destination = _str(row, "destination_city", "destination")
        country = _str(row, "country", "destination_country")
        purpose = _str(row, "purpose")
        if destination or country or purpose:
            summaries.append(" - ".join(item for item in [destination, country, purpose] if item))
    return "; ".join(summaries) or None


def _hotel_fit_note(hotels: list[dict[str, object]], profile: dict[str, object]) -> str | None:
    profile_type = _row_key(_str(profile, "profile_type"))
    for hotel in hotels:
        if profile_type and profile_type in _row_key(_str(hotel, "best_fit_profile")):
            return _joined_values(hotel, "hotel_name", "area", "business_comment")
    return None


def _hotel_insight_note(insights: list[dict[str, object]], profile: dict[str, object]) -> str | None:
    profile_type = _row_key(_str(profile, "profile_type"))
    for insight in insights:
        if profile_type and profile_type in _row_key(_str(insight, "best_fit_profile")):
            return _joined_values(insight, "hotel_name", "recommendation", "watch_outs")
    return None


def _transport_needs_transfer(value: str | None) -> bool:
    text = (value or "").lower()
    return any(marker in text for marker in ("private", "airport", "taxi", "transfer", "ride"))


def _policy_rules_from_text(text: str) -> list[str]:
    rules: list[str] = []
    for raw_line in text.splitlines():
        line = raw_line.strip(" -\t")
        if not line or len(line) < 8:
            continue
        if ":" in line:
            label, value = line.split(":", 1)
            if value.strip() and label.strip().lower() not in {"company", "client", "organization", "organisation"}:
                rules.append(f"{label.strip()}: {value.strip()}")
        elif "budget" not in line.lower() and any(marker in line.lower() for marker in ("allowed", "approval", "cabin", "hotel", "passport", "visa")):
            rules.append(line)
    return rules[:12]


def _compact_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _cabin_list_from_text(value: str | None) -> str | None:
    if not value:
        return None
    cabins = []
    text = value.lower().replace("-", " ")
    for cabin in ("economy", "premium economy", "business", "first"):
        if cabin in text:
            cabins.append(cabin.replace(" ", "_"))
    return ",".join(dict.fromkeys(cabins)) or value


def _money_int(value: object) -> int | None:
    if value is None:
        return None
    match = re.search(r"\d[\d,]*(?:\.\d+)?", str(value))
    if not match:
        return None
    return int(float(match.group(0).replace(",", "")))


def _money_display(value: object, currency: object) -> str | None:
    amount = _money_int(value)
    if not amount:
        return None
    return f"{currency or 'USD'} {amount}"


def _request_with_band_policy(
    request: CorporateTravelRequest,
    policy_rows: list[dict[str, object]],
) -> tuple[CorporateTravelRequest, list[dict[str, object]], str | None]:
    matches = _matching_band_policy_rows(request, policy_rows)
    if not matches:
        guess = _band_policy_guess_row(request)
        if not guess:
            return request, [], None
        guessed_request, _, _ = _request_with_band_policy(request, [guess])
        band_label = request.traveller_details.employee_band or request.traveller_details.employee_level or "submitted band"
        note = f"No uploaded Company Policy matched; used {band_label} for hotel tier and cabin policy."
        return guessed_request, [guess], note

    policy = _merged_reference_row(matches)
    traveller = request.traveller_details
    company = request.company_details
    policy_band = _str(policy, "employee_band", "traveler_band", "traveller_band", "band", "profile_type")

    traveller_update: dict[str, object] = {}
    if policy_band and not traveller.employee_band:
        traveller_update["employee_band"] = policy_band
    if policy_band and not traveller.employee_level:
        traveller_update["employee_level"] = policy_band

    company_update: dict[str, object] = {}
    policy_company = _row_company(policy)
    if not company.company_name and policy_company != "Company pending":
        company_update["company_name"] = policy_company
    approval_band = _str(policy, "approval_band", "approver_band")
    if approval_band and not company.approval_band:
        company_update["approval_band"] = approval_band
    policy_tier = _str(policy, "policy_tier", "tier")
    request_band_key = _row_key(traveller.employee_band or traveller.employee_level)
    current_tier_key = _row_key(company.policy_tier)
    if policy_tier and (not company.policy_tier or current_tier_key == request_band_key):
        company_update["policy_tier"] = policy_tier

    policy_notes = _str(policy, "approval_rule", "policy_notes", "notes")

    updated = request.model_copy(
        update={
            "traveller_details": traveller.model_copy(update=traveller_update),
            "company_details": company.model_copy(update=company_update),
            "budgets": request.budgets.model_copy(update={"policy_notes": policy_notes}) if policy_notes and not request.budgets.policy_notes else request.budgets,
        }
    )
    band_label = traveller.employee_band or traveller.employee_level or policy_band or "submitted band"
    return updated, matches, f"Band policy matched for {band_label}; hotel tier and cabin policy were applied from Company Policy."


def _band_policy_guess_row(request: CorporateTravelRequest) -> dict[str, object] | None:
    tier = None
    for value in (request.traveller_details.employee_band, request.traveller_details.employee_level):
        tier = _band_policy_tier_from_text(value)
        if tier:
            break
    if not tier:
        return None
    band_label = request.traveller_details.employee_band or request.traveller_details.employee_level
    return {
        **BAND_POLICY_GUESS_ROWS[tier],
        "employee_band": band_label,
        "approval_band": request.company_details.approval_band or band_label,
    }


def _band_policy_tier_from_text(value: object) -> str | None:
    compact = _row_key(value)
    if not compact:
        return None
    if compact in {"1", "band1"} or ("band" in compact and compact.endswith("1")):
        return "employee"
    if compact in {"2", "band2"} or ("band" in compact and compact.endswith("2")):
        return "manager"
    if compact in {"3", "band3"} or ("band" in compact and compact.endswith("3")):
        return "executive"
    if any(marker in compact for marker in ("ceo", "chief", "founder", "president", "chair", "executive", "principal", "bandc", "bandx")):
        return "executive"
    if any(marker in compact for marker in ("director", "vp", "vicepresident", "head", "bandm3", "bandm4", "bandm5")):
        return "senior"
    if any(marker in compact for marker in ("manager", "bandm1", "bandm2")):
        return "manager"
    if any(marker in compact for marker in ("employee", "associate", "analyst", "lead", "staff", "individualcontributor", "bande", "bandic")):
        return "employee"
    return None


def _matching_band_policy_rows(request: CorporateTravelRequest, rows: list[dict[str, object]]) -> list[dict[str, object]]:
    band_key = _row_key(request.traveller_details.employee_band or request.traveller_details.employee_level)
    if band_key:
        matches = [
            row for row in rows
            if _row_key(_str(row, "employee_band", "traveler_band", "traveller_band", "band", "profile_type")) == band_key
        ]
        if matches:
            company_key = _row_key(request.company_details.company_name)
            if company_key:
                company_matches = [row for row in matches if _row_key(_row_company(row)) == company_key]
                if company_matches:
                    return company_matches
            return matches
    tier_key = _row_key(request.company_details.policy_tier)
    if not tier_key:
        return []
    return [
        row for row in rows
        if _row_key(_str(row, "policy_tier", "tier")) == tier_key
    ]


def _matching_traveler_rows(request: CorporateTravelRequest, rows: list[dict[str, object]]) -> list[dict[str, object]]:
    email = _row_key(request.traveller_details.traveler_email)
    employee_id = _row_key(request.traveller_details.employee_id)
    name = _row_key(request.traveller_details.traveler_name)
    matches = []
    for row in rows:
        row_email = _row_key(_str(row, "traveler_email", "traveller_email", "email", "employee_email"))
        row_employee = _row_key(_str(row, "employee_id", "traveler_id"))
        row_name = _row_key(_str(row, "traveler_name", "traveller_name", "employee_name", "name"))
        if (email and row_email == email) or (employee_id and row_employee == employee_id) or (name and row_name == name):
            matches.append(row)
    return matches


def _merged_reference_row(rows: list[dict[str, object]]) -> dict[str, object]:
    merged: dict[str, object] = {}
    for row in rows:
        for key, value in row.items():
            if key not in merged and value not in {None, ""}:
                merged[key] = value
    return merged


def _matching_visa_reference(
    request: CorporateTravelRequest,
    visa_rows: list[dict[str, object]],
    matched_history_rows: list[dict[str, object]],
) -> dict[str, object] | None:
    identifiers = {
        _row_key(request.traveller_details.traveler_email),
        _row_key(request.traveller_details.employee_id),
        _row_key(request.traveller_details.traveler_name),
    }
    for row in matched_history_rows:
        identifiers.update(
            {
                _row_key(_str(row, "traveler_email", "traveller_email", "email", "employee_email")),
                _row_key(_str(row, "employee_id", "traveler_id")),
                _row_key(_str(row, "traveler_name", "traveller_name", "employee_name", "name")),
            }
        )
    identifiers.discard("")
    destination_country = _row_key(request.travel_details.destination_country)
    destination = _row_key(request.travel_details.destination)
    fallback: dict[str, object] | None = None
    for row in visa_rows:
        row_ids = {
            _row_key(_str(row, "traveler_email", "traveller_email", "email")),
            _row_key(_str(row, "employee_id", "traveler_id")),
            _row_key(_str(row, "traveler_name", "traveller_name", "name")),
        }
        if not identifiers.intersection(row_ids):
            continue
        fallback = fallback or row
        row_destination = _row_key(_str(row, "destination_country", "country_region", "country"))
        if not destination_country or row_destination in {destination_country, destination}:
            return row
    return fallback


def _traveler_status(documents: list[dict[str, object]]) -> str:
    statuses = {str(document.get("status")) for document in documents}
    if "Missing" in statuses:
        return "Missing Passport"
    if "Expiring Soon" in statuses:
        return "Passport Expiring"
    if "Needs Review" in statuses:
        return "Document Update Required"
    return "Compliant"


def _redacted(value: object) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    return _masked_identifier(text)


def _masked_identifier(value: object) -> str:
    text = str(value or "").strip()
    if len(text) <= 4:
        return "*" * len(text) if text else ""
    return f"{'*' * (len(text) - 4)}{text[-4:]}"


def _loyalty_programs(row: dict[str, object]) -> list[dict[str, str | None]]:
    programs = []
    airline = _str(row, "loyalty_airline", "preferred_airline")
    if airline:
        programs.append({"provider": airline, "tier": _str(row, "airline_tier"), "account_ref": _str(row, "airline_account_ref") or "On file"})
    hotel = _str(row, "loyalty_hotel", "hotel_loyalty")
    if hotel:
        programs.append({"provider": hotel, "tier": _str(row, "hotel_tier"), "account_ref": _str(row, "hotel_account_ref") or "On file"})
    return programs


def _corporate_template_workbook_bytes() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    requests = workbook.active
    requests.title = "Travel Requests"
    requests.append(
        [
            "traveler_name",
            "traveler_email",
            "employee_id",
            "department",
            "nationality",
            "origin",
            "phone",
            "employee_band",
            "destination",
            "destination_country",
            "depart_date",
            "return_date",
            "trip_purpose",
            "meeting_location",
            "flexible_dates",
            "cabin",
            "passport_expiry",
            "visa_status",
            "visa_expiry",
            "approval_band",
            "preferred_airline",
            "flight_preference",
            "hotel_preference",
            "preferred_hotel_area",
            "past_hotel_preference",
            "airport_transfer_needed",
            "meal_preference",
            "seat_preference",
            "timing_preference",
            "medical_notes",
            "accessibility_notes",
            "extra_baggage_notes",
            "special_requests",
        ]
    )
    requests.append(
        [
            "Anika Rao",
            "anika.rao@example.com",
            "E-101",
            "Sales",
            "India",
            "Hyderabad",
            "+91 98765 43210",
            "2",
            "Johannesburg",
            "South Africa",
            "2026-06-10",
            "2026-06-17",
            "Client workshops",
            "Sandton client office",
            "yes",
            "economy",
            "2028-01-15",
            "Valid visa on file",
            "2027-08-20",
            "Band A2",
            "Qatar Airways",
            "Short layover",
            "Business hotel near client office",
            "Sandton",
            "Garden Court",
            "yes",
            "Vegetarian",
            "Aisle",
            "Morning arrival",
            "None",
            "None",
            "One sample kit",
            "Airport pickup; late check-in",
        ]
    )

    policy = workbook.create_sheet("Company Policy")
    policy.append(["policy_tier", "employee_band", "approval_band", "allowed_cabins", "hotel_tier", "approval_rule", "notes"])
    policy.append(["standard", "2", "Band A2", "economy,premium_economy", "Manager tier: 3/4-star hotels", "approval required for cabin or document policy exceptions", "Prefer refundable fares for client-facing trips"])

    employees = workbook.create_sheet("Employee Details")
    employees.append([
        "employee_id",
        "employee_name",
        "employee_email",
        "company_name",
        "department",
        "employee_band",
        "location",
        "nationality",
        "passport_number",
        "passport_expiry",
        "visa_status",
        "visa_expiry",
        "seat_preference",
        "meal_preference",
        "hotel_preference",
        "loyalty_airline",
        "airline_tier",
        "loyalty_hotel",
        "hotel_tier",
        "policy_notes",
        "recent_trips",
    ])
    employees.append([
        "E-101",
        "Anika Rao",
        "anika.rao@example.com",
        "Unipro",
        "Sales",
        "2",
        "Hyderabad",
        "India",
        "Z1234567",
        "2028-01-15",
        "Valid visa on file",
        "2027-08-20",
        "Aisle",
        "Vegetarian",
        "Business hotel near client office",
        "Qatar Airways",
        "Gold",
        "Marriott Bonvoy",
        "Gold",
        "Business class requires approval unless flight exceeds policy threshold",
        "Hyderabad to Singapore; Hyderabad to London",
    ])

    history = workbook.create_sheet("Traveller History")
    history.append(["traveler_email", "preferred_airline", "meal_preference", "seat_preference", "hotel_notes", "notes"])
    history.append(["anika.rao@example.com", "Qatar Airways", "Vegetarian", "Aisle", "Prefers hotels near office", "Past trips favored short layovers"])

    visa = workbook.create_sheet("Visa Rules")
    visa.append(["from_country", "destination_country", "visa_required", "transit_notes", "notes"])
    visa.append(["India", "South Africa", "yes", "Check transit visa if route changes", "Business visa must be valid for full stay"])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _final_total_for_origin_currency(request: CorporateTravelRequest) -> tuple[int, str] | None:
    plan = request.generated_plan
    if not plan:
        return None
    selected_option = _selected_travel_option_for_request(request)
    estimated_cost = selected_option.estimated_cost if selected_option else plan.budget_policy_check.estimated_cost
    final_currency = origin_city_currency(request.travel_details.origin, request.budgets.currency)
    final_total = convert_final_amount(estimated_cost, request.budgets.currency or "USD", final_currency).amount
    return int(round(final_total)), final_currency


def _origin_currency_amount(amount: int, request: CorporateTravelRequest, source_currency: str | None = None) -> tuple[int, str]:
    final_currency = origin_city_currency(request.travel_details.origin, request.budgets.currency)
    final_total = convert_final_amount(amount, source_currency or request.budgets.currency or "USD", final_currency).amount
    return int(round(final_total)), final_currency


def _origin_currency_component_display(amount: int, currency: str | None, request: CorporateTravelRequest) -> str:
    converted, final_currency = _origin_currency_amount(amount, request, currency)
    return f"{converted} {final_currency}"


def _final_flight_summary(request: CorporateTravelRequest, option: TravelOption, flight) -> str:
    if not flight:
        return option.flight_summary
    fare = _origin_currency_component_display(flight.total_amount, flight.currency, request)
    cabin = flight.cabin.replace("_", " ")
    return f"{flight.airline}; {cabin}; {flight.outbound}; fare {fare}"


def _final_hotel_summary(request: CorporateTravelRequest, option: TravelOption, hotel) -> str:
    if not hotel:
        return option.hotel_summary
    rate = _origin_currency_component_display(hotel.total_amount, hotel.currency, request)
    stay = f"{hotel.nights} night(s)" if hotel.nights else "stay dates pending"
    return f"{hotel.name}; {stay}; rate {rate}"


def _final_transfer_summary(request: CorporateTravelRequest, option: TravelOption, transfer) -> str:
    if not transfer:
        return option.transfer_summary
    rate = _origin_currency_component_display(transfer.total_amount, transfer.currency, request)
    service = transfer.service_type.replace("_", " ").title()
    return f"{service} transfer from {transfer.pickup_airport_code} to {transfer.dropoff_label}; rate {rate}"


def _final_customer_itinerary_note(request: CorporateTravelRequest, option: TravelOption | None = None) -> str:
    if not request.generated_plan:
        return ""
    stored_note = request.generated_plan.customer_itinerary_draft
    final_currency = origin_city_currency(request.travel_details.origin, request.budgets.currency)
    if stored_note and not _has_non_origin_currency(stored_note, final_currency):
        return stored_note
    selected_option = option or _selected_travel_option_for_request(request)
    if selected_option:
        return _customer_itinerary_from_selected_option(request, selected_option)
    return stored_note


def _has_non_origin_currency(text: str, final_currency: str) -> bool:
    codes = set(re.findall(r"\b[A-Z]{3}\b", text))
    if any(code in DISPLAY_CURRENCY_CODES and code != final_currency for code in codes):
        return True
    return final_currency != "USD" and "$" in text


def _corporate_request_workbook_bytes(request: CorporateTravelRequest) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    final_total = _final_total_for_origin_currency(request)
    summary = workbook.active
    summary.title = "Request"
    summary.append(["field", "value"])
    summary.append(["request_id", request.id])
    summary.append(["traveler_name", request.traveller_details.traveler_name or ""])
    summary.append(["traveler_email", request.traveller_details.traveler_email or ""])
    summary.append(["passport_number", _masked_identifier(request.traveller_details.passport_number or "")])
    summary.append(["passport_expiry", request.traveller_details.passport_expiry.isoformat() if request.traveller_details.passport_expiry else ""])
    summary.append(["company_name", request.company_details.company_name or ""])
    summary.append(["origin", request.travel_details.origin or ""])
    summary.append(["destination", request.travel_details.destination or ""])
    summary.append(["depart_date", request.travel_details.depart_date.isoformat() if request.travel_details.depart_date else ""])
    summary.append(["return_date", request.travel_details.return_date.isoformat() if request.travel_details.return_date else ""])
    summary.append(["trip_purpose", request.travel_details.trip_purpose or ""])
    if final_total:
        summary.append(["final_total", final_total[0]])
        summary.append(["final_total_currency", final_total[1]])
    summary.append(["status", request.status])

    plans = workbook.create_sheet("Recommended Plans")
    plans.append(["option_name", "flight_summary", "hotel_summary", "transfer_summary", "estimated_cost", "policy_status", "recommendation_reason"])
    if request.generated_plan:
        for option in request.generated_plan.travel_options:
            plans.append(
                [
                    option.option_name,
                    option.flight_summary,
                    option.hotel_summary,
                    option.transfer_summary,
                    option.estimated_cost,
                    option.policy_status,
                    option.recommendation_reason,
                ]
            )

    itinerary = workbook.create_sheet("Final Itinerary")
    itinerary.append(["section", "value"])
    if request.generated_plan:
        if final_total:
            itinerary.append(["final_total", f"{final_total[0]} {final_total[1]}"])
        itinerary.append(["customer_message", mask_sensitive_customer_text(request.generated_plan.customer_message_draft)])
        itinerary.append(["final_itinerary", mask_sensitive_customer_text(_final_customer_itinerary_note(request))])
        itinerary.append(["agent_note", request.generated_plan.agent_note])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _customer_safe_plan(plan):
    return plan.model_copy(
        update={
            "customer_message_draft": mask_sensitive_customer_text(plan.customer_message_draft),
            "customer_itinerary_draft": mask_sensitive_customer_text(plan.customer_itinerary_draft),
        }
    )


def _corporate_request_pdf_bytes(request: CorporateTravelRequest) -> bytes:
    plan = request.generated_plan
    traveler = request.traveller_details
    travel = request.travel_details
    final_total = _final_total_for_origin_currency(request)
    final_total_display = f"{final_total[0]} {final_total[1]}" if final_total else "Pending"
    selected_flight = None
    selected_hotel = None
    selected_transfer = None
    selected_option = None
    if plan:
        selected_flight = next((offer for offer in plan.flight_offers if offer.id == plan.selected_flight_offer_id), None)
        selected_hotel = next((offer for offer in plan.hotel_offers if offer.id == plan.selected_hotel_offer_id), None)
        selected_transfer = next((offer for offer in plan.ground_transfer_offers if offer.id == plan.selected_ground_transfer_offer_id), None)
        selected_option = _selected_travel_option_for_request(request)

    pdf = _PdfPageBuilder()
    pdf.header(
        "Unipro Travel Operations",
        "Final Travel Itinerary",
        f"Request {request.id} | Generated {_date_text(date.today())}",
    )
    pdf.badge_row([
        ("Traveler", traveler.traveler_name or "Traveler pending"),
        ("Route", f"{travel.origin or 'Origin pending'} to {travel.destination or 'Destination pending'}"),
        ("Final total", final_total_display),
    ])
    pdf.table("Traveler Details", [
        ("Name", traveler.traveler_name or "Traveler pending"),
        ("Email", traveler.traveler_email or "Email pending"),
        ("Nationality", traveler.nationality or "Not captured"),
        ("Passport expiry", _date_text(traveler.passport_expiry)),
    ])
    pdf.table("Trip Details", [
        ("Origin", travel.origin or "Origin pending"),
        ("Destination", travel.destination or "Destination pending"),
        ("Destination country", travel.destination_country or "Not captured"),
        ("Travel dates", f"{_date_text(travel.depart_date)} to {_date_text(travel.return_date)}"),
        ("Purpose", travel.trip_purpose or "Business travel"),
        ("Meeting location", travel.meeting_location or "Not captured"),
        ("Hotel/location preference", request.preferences.hotel_preference or request.preferences.preferred_hotel_area or "Not captured"),
        ("Flight preference", request.preferences.flight_preference or request.preferences.preferred_airline or "Not captured"),
    ])
    if plan:
        pdf.table("Readiness And Policy", [
            ("Passport", plan.travel_readiness.passport_status),
            ("Visa", plan.travel_readiness.visa_status),
            ("Transit", plan.travel_readiness.transit_warning or "No transit warning"),
            ("Policy status", plan.budget_policy_check.policy_status),
            ("Final total", final_total_display),
        ])
        if selected_flight:
            pdf.table("Selected Flight", [
                ("Airline", selected_flight.airline),
                ("Summary", selected_flight.summary),
                ("Outbound departure / landing", selected_flight.outbound),
                ("Return departure / landing", selected_flight.return_leg or "Not captured"),
                ("Cabin", selected_flight.cabin),
                ("Fare", _origin_currency_component_display(selected_flight.total_amount, selected_flight.currency, request)),
            ])
        if selected_hotel:
            pdf.table("Selected Hotel", [
                ("Hotel", selected_hotel.name),
                ("Summary", selected_hotel.summary),
                ("Address", selected_hotel.address or "Not captured"),
                ("Stay", f"{_date_text(selected_hotel.check_in)} to {_date_text(selected_hotel.check_out)}"),
                ("Check-in starts", selected_hotel.check_in_starts_at or "Not captured"),
                ("Checkout time", selected_hotel.checkout_time or "Not captured"),
                ("Rooms / guests", f"{selected_hotel.rooms} room(s), {selected_hotel.guests} guest(s)"),
                ("Room notes", selected_hotel.room_notes or "Not captured"),
                ("Cancellation", selected_hotel.cancellation_notes or "Must be verified before confirmation"),
                ("Rate", _origin_currency_component_display(selected_hotel.total_amount, selected_hotel.currency, request)),
            ])
        if selected_transfer:
            pdf.table("Airport Transfer", [
                ("Cab service source", "Ground transport"),
                ("Pickup airport", selected_transfer.pickup_airport_code),
                ("Pickup time", selected_transfer.pickup_time or "Not captured"),
                ("Dropoff", selected_transfer.dropoff_label),
                ("Dropoff address", selected_transfer.dropoff_address or "Not captured"),
                ("Cab service", selected_transfer.service_type),
                ("Vehicle", selected_transfer.vehicle_type or "Not captured"),
                ("Passengers", str(selected_transfer.passengers)),
                ("Baggage", selected_transfer.baggage or "Not captured"),
                ("Cancellation", selected_transfer.cancellation_notes or "Must be verified before confirmation"),
                ("Rate", _origin_currency_component_display(selected_transfer.total_amount, selected_transfer.currency, request)),
            ])
        pdf.note("Special Requests", _special_request_notice(request))
        pdf.note("Client Note", mask_sensitive_customer_text(_final_customer_itinerary_note(request, selected_option)))
    else:
        pdf.note("Recommended Travel Plan", "Plan not generated.")
    return _styled_pdf(pdf.pages())


def _corporate_option_pdf_bytes(request: CorporateTravelRequest, option_index: int) -> bytes:
    plan = request.generated_plan
    if not plan:
        return _simple_pdf(["Itinerary Option", "Plan not generated."])
    index = max(0, min(option_index - 1, len(plan.travel_options) - 1))
    option = plan.travel_options[index]
    flight = next((offer for offer in plan.flight_offers if offer.id == option.flight_offer_id), None)
    hotel = plan.hotel_offers[index] if index < len(plan.hotel_offers) else None
    transfer = next((offer for offer in plan.ground_transfer_offers if offer.id == option.ground_transfer_offer_id), None)
    final_total, final_currency = _origin_currency_amount(option.estimated_cost, request)
    lines = [
        f"Itinerary Option {index + 1}: {option.option_name}",
        f"Request: {request.id}",
        f"Traveller: {request.traveller_details.traveler_name or 'Traveller pending'}",
        f"Route: {request.travel_details.origin or 'Origin pending'} to {request.travel_details.destination or 'Destination pending'}",
        f"Dates: {request.travel_details.depart_date or 'Date pending'} to {request.travel_details.return_date or 'Return pending'}",
        f"Flight: {_final_flight_summary(request, option, flight)}",
        f"Hotel: {_final_hotel_summary(request, option, hotel)}",
        f"Transfer: {_final_transfer_summary(request, option, transfer)}",
        f"Final total: {final_total} {final_currency}",
        f"Policy fit: {option.policy_status}",
        f"Why this option: {option.recommendation_reason}",
    ]
    return _simple_pdf(lines)


def _selected_option_index(request: CorporateTravelRequest, approval: CorporateClientApprovalRequest) -> int:
    plan = request.generated_plan
    if not plan:
        return 0
    if approval.selected_option_index:
        return min(max(approval.selected_option_index - 1, 0), len(plan.travel_options) - 1)
    selected_name = (approval.selected_option_name or "").strip().lower()
    if selected_name:
        for index, option in enumerate(plan.travel_options):
            if option.option_name.lower() == selected_name:
                return index
    return 0


def _selected_travel_option_for_request(request: CorporateTravelRequest) -> TravelOption | None:
    plan = request.generated_plan
    if not plan or not plan.travel_options:
        return None
    if request.client_review and request.client_review.selected_option_index:
        index = request.client_review.selected_option_index - 1
        if 0 <= index < len(plan.travel_options):
            return plan.travel_options[index]
    selected_flight_id = plan.selected_flight_offer_id
    for option in plan.travel_options:
        if selected_flight_id and option.flight_offer_id == selected_flight_id:
            return option
    return plan.travel_options[0]


def _customer_itinerary_from_selected_option(request: CorporateTravelRequest, option: TravelOption) -> str:
    plan = request.generated_plan
    flight = None
    hotel = None
    transfer = None
    if plan:
        hotel_index = plan.travel_options.index(option) if option in plan.travel_options else 0
        flight = next((offer for offer in plan.flight_offers if offer.id == option.flight_offer_id), None)
        hotel = plan.hotel_offers[hotel_index] if hotel_index < len(plan.hotel_offers) else None
        transfer = next((offer for offer in plan.ground_transfer_offers if offer.id == option.ground_transfer_offer_id), None)
    final_total, final_currency = _origin_currency_amount(option.estimated_cost, request)
    lines = [
        f"Approved itinerary for {request.traveller_details.traveler_name or 'traveller'}",
        f"Route: {request.travel_details.origin or 'Origin pending'} to {request.travel_details.destination or 'Destination pending'}",
        f"Dates: {request.travel_details.depart_date or 'Date pending'} to {request.travel_details.return_date or 'Return pending'}",
        f"Selected option: {option.option_name}",
        f"Flight: {_final_flight_summary(request, option, flight)}",
        f"Hotel: {_final_hotel_summary(request, option, hotel)}",
        f"Airport transfer: {_final_transfer_summary(request, option, transfer)}",
        f"Final total: {final_total} {final_currency}",
    ]
    if flight:
        lines.extend([
            f"Outbound flight departure / landing: {flight.outbound}",
            f"Return flight departure / landing: {flight.return_leg or 'Not captured'}",
            f"Flight cabin: {flight.cabin}",
        ])
    if hotel:
        lines.extend([
            f"Hotel stay dates: {hotel.check_in or 'Not captured'} to {hotel.check_out or 'Not captured'}",
            f"Hotel check-in starts: {hotel.check_in_starts_at or 'Not captured'}",
            f"Hotel checkout time: {hotel.checkout_time or 'Not captured'}",
            f"Hotel address: {hotel.address or 'Not captured'}",
        ])
    if transfer:
        lines.extend([
            "Cab service source: Ground transport",
            f"Cab service: {transfer.service_type} / {transfer.vehicle_type or 'vehicle pending'}",
            f"Cab pickup: {transfer.pickup_airport_code} at {transfer.pickup_time or 'time pending'}",
            f"Cab dropoff: {transfer.dropoff_label} ({transfer.dropoff_address or 'address pending'})",
        ])
    lines.extend([
        _special_request_notice(request),
    ])
    return "\n".join(lines)


class _PdfPageBuilder:
    def __init__(self) -> None:
        self._pages: list[list[str]] = []
        self._commands: list[str] = []
        self._y = 748

    def pages(self) -> list[str]:
        if self._commands:
            self._pages.append(self._commands)
        return ["\n".join(page) for page in self._pages]

    def _new_page(self) -> None:
        self._pages.append(self._commands)
        self._commands = []
        self._y = 748

    def _ensure(self, height: int) -> None:
        if self._y - height < 54:
            self._new_page()

    def _text(self, x: int, y: int, text: str, size: int = 9, color: str = "0.10 0.12 0.14") -> None:
        self._commands.append(f"{color} rg BT /F1 {size} Tf {x} {y} Td ({_pdf_escape(str(text))}) Tj ET")

    def _rect(self, x: int, y: int, width: int, height: int, fill: str | None = None, stroke: str | None = None) -> None:
        if fill:
            self._commands.append(f"{fill} rg {x} {y} {width} {height} re f")
        if stroke:
            self._commands.append(f"{stroke} RG {x} {y} {width} {height} re S")

    def header(self, brand: str, title: str, meta: str) -> None:
        self._rect(36, 704, 540, 64, "0.93 0.98 0.97", "0.72 0.84 0.82")
        self._text(54, 744, brand, 10, "0.00 0.31 0.28")
        self._text(54, 724, title, 20, "0.00 0.24 0.22")
        self._text(54, 710, meta, 8, "0.32 0.42 0.43")
        self._y = 686

    def badge_row(self, badges: list[tuple[str, str]]) -> None:
        self._ensure(54)
        width = 172
        for index, (label, value) in enumerate(badges):
            x = 36 + (index * 184)
            self._rect(x, self._y - 44, width, 44, "0.98 0.99 0.99", "0.84 0.89 0.88")
            self._text(x + 10, self._y - 17, label.upper(), 7, "0.35 0.44 0.45")
            self._text(x + 10, self._y - 32, _pdf_trim(value, 27), 10, "0.08 0.12 0.13")
        self._y -= 62

    def table(self, title: str, rows: list[tuple[str, str]]) -> None:
        wrapped_rows = [(label, _wrap_pdf_text(str(value), 62)) for label, value in rows]
        height = 30 + sum(max(22, 14 * len(value_lines) + 8) for _, value_lines in wrapped_rows)
        self._ensure(height + 12)
        top = self._y
        self._rect(36, top - height, 540, height, "1 1 1", "0.82 0.88 0.87")
        self._rect(36, top - 26, 540, 26, "0.95 0.98 0.98")
        self._text(50, top - 18, title, 11, "0.00 0.31 0.28")
        y = top - 44
        for label, value_lines in wrapped_rows:
            row_height = max(22, 14 * len(value_lines) + 8)
            line_y = y - row_height + 10
            self._commands.append(f"0.86 0.91 0.90 RG 36 {line_y} m 576 {line_y} l S")
            self._text(50, y, label, 8, "0.35 0.44 0.45")
            for line_index, line in enumerate(value_lines):
                self._text(180, y - (line_index * 14), line, 9, "0.11 0.14 0.15")
            y -= row_height
        self._y = top - height - 14

    def note(self, title: str, body: str) -> None:
        lines = _wrap_pdf_text(body or "No note provided.", 84)[:8]
        height = 32 + (len(lines) * 14)
        self._ensure(height + 12)
        top = self._y
        self._rect(36, top - height, 540, height, "0.98 0.99 0.99", "0.82 0.88 0.87")
        self._text(50, top - 18, title, 11, "0.00 0.31 0.28")
        for index, line in enumerate(lines):
            self._text(50, top - 38 - (index * 14), line, 9, "0.11 0.14 0.15")
        self._y = top - height - 14

    def footer_note(self, text: str) -> None:
        self._ensure(44)
        self._rect(36, self._y - 34, 540, 34, "0.99 0.98 0.94", "0.90 0.84 0.65")
        self._text(50, self._y - 21, _pdf_trim(text, 104), 8, "0.43 0.35 0.13")
        self._y -= 46


def _styled_pdf(streams: list[str]) -> bytes:
    if not streams:
        streams = [""]
    font_id = 3
    page_refs = [5 + index * 2 for index in range(len(streams))]
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{page_id} 0 R' for page_id in page_refs)}] /Count {len(streams)} >>".encode("ascii"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for index, stream_text in enumerate(streams):
        content_id = 4 + index * 2
        page_id = 5 + index * 2
        stream = stream_text.encode("utf-8")
        objects.append(b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode("ascii")
        )
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode("ascii"))
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")
    xref_start = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_start}\n%%EOF\n".encode("ascii"))
    return bytes(pdf)


def _wrap_pdf_text(value: str, limit: int) -> list[str]:
    text = " ".join(str(value).split())
    if not text:
        return [""]
    lines: list[str] = []
    current = ""
    for word in text.split():
        candidate = f"{current} {word}".strip()
        if len(candidate) > limit and current:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines or [""]


def _pdf_trim(value: str, limit: int) -> str:
    text = " ".join(str(value).split())
    if len(text) <= limit:
        return text
    return f"{text[: max(0, limit - 3)].rstrip()}..."


def _wrapped_pdf_lines(lines: list[str], limit: int = 86) -> list[str]:
    wrapped: list[str] = []
    for line in lines:
        if len(line) <= limit:
            wrapped.append(line)
            continue
        words = line.split()
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if len(candidate) > limit and current:
                wrapped.append(current)
                current = word
            else:
                current = candidate
        if current:
            wrapped.append(current)
    return wrapped


def _date_text(value: date | None) -> str:
    return value.isoformat() if value else "TBD"


def _client_request_form_lines() -> list[str]:
    lines = ["Unipro Travel Operations", "Client Travel Request Form", ""]
    for title, fields in _client_request_form_sections():
        lines.extend([title, *fields, ""])
    lines.append("Hotel tier is automatic from band: CEO 4/5-star, Manager 3/4-star, Employee 2/3-star.")
    lines.append("Excel remains available as a backup import for operations teams.")
    lines.append("Passport, visa, loyalty, seating, meal, and saved preferences are pulled from the traveler roster.")
    return lines


def _client_request_form_sections() -> list[tuple[str, list[str]]]:
    return [
        ("Traveler matching", [
            "Traveller name: ________________________________",
            "Employee band / role (CEO, Manager, Employee, or band 1/2/3): _______________________",
        ]),
        ("Trip details", [
            "Origin city / airport: _______________",
            "Destination city / airport: __________",
            "Destination country: __________________",
            "Departure date: ______________________",
            "Return date: _________________________",
            "Meeting office / location: ____________",
            "Flexible dates: Yes / No",
        ]),
        ("Preferences and requests", [
            "Flight preference: ___________________",
            "Hotel area / location preference: ____",
            "Airport transfer needed: Yes / No",
            "Meal preference: _____________________",
            "Trip-specific requests: ______________",
        ]),
    ]


def _client_request_pdf_bytes() -> bytes:
    return _simple_pdf(_client_request_form_lines())


def _client_request_docx_bytes() -> bytes:
    return _structured_request_docx()


def _structured_request_docx() -> bytes:
    document_xml = "\n".join([
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">',
        "<w:body>",
        _docx_paragraph("Unipro Travel Operations", bold=True, size=34, color="174A5C", align="center"),
        _docx_paragraph("Client Travel Request Form", bold=True, size=28, color="0F172A", align="center"),
        _docx_paragraph("Fill the editable fields below. Hotel tier is automatic from band; Excel remains available as a backup import.", size=20, color="64748B", align="center"),
        *[_docx_section(title, fields) for title, fields in _client_request_form_sections()],
        _docx_paragraph("Roster note: Passport, visa, loyalty, seating, meal, and saved preferences are pulled from the traveler roster.", size=20, color="64748B"),
        '<w:sectPr><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr>',
        "</w:body>",
        "</w:document>",
    ])
    content_types = "\n".join([
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">',
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>',
        '<Default Extension="xml" ContentType="application/xml"/>',
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>',
        "</Types>",
    ])
    relationships = "\n".join([
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">',
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>',
        "</Relationships>",
    ])
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document_xml)
    return buffer.getvalue()


def _docx_section(title: str, fields: list[str]) -> str:
    rows = "\n".join(_docx_table_row(field) for field in fields)
    return "\n".join([
        _docx_paragraph(title, bold=True, size=24, color="174A5C"),
        '<w:tbl><w:tblPr><w:tblW w:w="5000" w:type="pct"/><w:tblBorders>'
        '<w:top w:val="single" w:sz="8" w:color="D6DEE3"/>'
        '<w:left w:val="single" w:sz="8" w:color="D6DEE3"/>'
        '<w:bottom w:val="single" w:sz="8" w:color="D6DEE3"/>'
        '<w:right w:val="single" w:sz="8" w:color="D6DEE3"/>'
        '<w:insideH w:val="single" w:sz="6" w:color="E5E7EB"/>'
        '</w:tblBorders></w:tblPr>',
        rows,
        "</w:tbl>",
        _docx_paragraph(""),
    ])


def _docx_table_row(field: str) -> str:
    return "\n".join([
        "<w:tr>",
        '<w:tc><w:tcPr><w:tcW w:w="5000" w:type="pct"/><w:shd w:fill="F8FAFC"/></w:tcPr>',
        _docx_field_paragraph(field, size=21, color="0F172A"),
        "</w:tc>",
        "</w:tr>",
    ])


def _docx_field_paragraph(line: str, *, size: int = 22, color: str | None = None) -> str:
    if ":" not in line:
        return _docx_paragraph(line, bold=True, size=size, color=color)
    label, value = line.split(":", 1)
    p_props = '<w:pPr><w:spacing w:after="120"/></w:pPr>'
    label_props = "".join([
        "<w:b/>",
        f'<w:color w:val="{color}"/>' if color else "",
        f'<w:sz w:val="{size}"/>',
    ])
    value_props = "".join([
        f'<w:color w:val="{color}"/>' if color else "",
        f'<w:sz w:val="{size}"/>',
    ])
    return "".join([
        f"<w:p>{p_props}",
        f'<w:r><w:rPr>{label_props}</w:rPr><w:t xml:space="preserve">{xml_escape(label)}:</w:t></w:r>',
        f'<w:r><w:rPr>{value_props}</w:rPr><w:t xml:space="preserve">{xml_escape(value)}</w:t></w:r>',
        "</w:p>",
    ])


def _docx_paragraph(line: str, *, bold: bool = False, size: int = 22, color: str | None = None, align: str | None = None) -> str:
    if not line:
        return "<w:p/>"
    alignment = f'<w:jc w:val="{align}"/>' if align else ""
    p_props = f'<w:pPr>{alignment}<w:spacing w:after="120"/></w:pPr>'
    r_props = "".join([
        "<w:b/>" if bold else "",
        f'<w:color w:val="{color}"/>' if color else "",
        f'<w:sz w:val="{size}"/>',
    ])
    return f'<w:p>{p_props}<w:r><w:rPr>{r_props}</w:rPr><w:t xml:space="preserve">{xml_escape(line)}</w:t></w:r></w:p>'


def _simple_pdf(lines: list[str]) -> bytes:
    content_lines = ["BT", "/F1 18 Tf", "50 780 Td", f"({_pdf_escape(lines[0])}) Tj", "/F1 11 Tf", "0 -30 Td"]
    for line in lines[1:]:
        safe_line = _pdf_escape(line)
        content_lines.append(f"({safe_line}) Tj" if safe_line else "() Tj")
        content_lines.append("0 -20 Td")
    content_lines.append("ET")
    stream = "\n".join(content_lines).encode("utf-8")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode("ascii"))
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")
    xref_start = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_start}\n%%EOF\n".encode("ascii"))
    return bytes(pdf)


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _corporate_request_from_row(row: dict[str, object], context: AuthContext, now: datetime) -> CorporateTravelRequest:
    special = row.get("special_requests") or row.get("special_request") or ""
    travel_details = {
        "origin": _str(row, "origin", "from"),
        "destination": _str(row, "destination", "to"),
        "destination_country": _str(row, "destination_country", "country"),
        "depart_date": _date_value(row.get("depart_date") or row.get("departure_date")),
        "return_date": _date_value(row.get("return_date")),
        "trip_purpose": _str(row, "trip_purpose", "purpose"),
        "meeting_location": _str(row, "meeting_location"),
        "flexible_dates": _bool_value(row.get("flexible_dates")),
        "include_outbound_flight": _component_bool(row, "include_outbound_flight", "outbound_flight_needed", "origin_flight_needed", default=True),
        "include_return_flight": _component_bool(row, "include_return_flight", "return_flight_needed", default=True),
        "include_hotel": _component_bool(row, "include_hotel", "hotel_needed", "book_hotel", "hotel_required", default=True),
        "include_ground_transfer": _component_bool(row, "include_ground_transfer", "ground_transfer_needed", "airport_transfer_needed", "cab_needed", default=True),
        "travelers": _int_value(row.get("travelers") or row.get("travellers") or 1, default=1),
    }
    submitted_cabin = _str(row, "cabin")
    if submitted_cabin:
        travel_details["cabin"] = _cabin_value(submitted_cabin)
    return CorporateTravelRequest(
        owner_id=context.user_id,
        owner_department=context.department,
        status="New",
        requester_email=_str(row, "requester_email", "sender_email", "submitter_email"),
        traveller_details=TravellerDetails(
            traveler_name=_str(row, "traveler_name", "traveller_name", "name"),
            traveler_email=_str(row, "traveler_email", "traveller_email", "email"),
            phone=_str(row, "phone", "phone_number", "mobile"),
            employee_id=_str(row, "employee_id"),
            employee_level=_str(row, "employee_band", "traveler_band", "traveller_band", "band", "employee_level", "level"),
            employee_band=_str(row, "employee_band", "traveler_band", "traveller_band", "band"),
            department=_str(row, "department"),
            nationality=_str(row, "nationality", "citizenship"),
            passport_number=_str(row, "passport_number"),
            passport_expiry=_date_value(row.get("passport_expiry")),
            visa_status=_str(row, "visa_status", "visa"),
            visa_expiry=_date_value(row.get("visa_expiry")),
            medical_notes=_str(row, "medical_notes", "medical_accessibility_notes"),
            accessibility_notes=_str(row, "accessibility_notes"),
        ),
        company_details=CompanyDetails(
            company_name=_company_name_from_request_row(row),
            cost_center=_str(row, "cost_center"),
            approving_manager=_str(row, "approving_manager", "manager"),
            approval_manager_email=_str(row, "approval_manager_email", "approving_manager_email", "manager_email"),
            approval_band=_str(row, "approval_band", "approver_band"),
            policy_tier=_str(row, "policy_tier", "approval_band", "employee_band"),
        ),
        travel_details=TravelDetails(**travel_details),
        preferences=TravelPreferences(
            preferred_airline=_str(row, "preferred_airline"),
            flight_preference=_str(row, "flight_preference"),
            hotel_preference=_str(row, "hotel_preference"),
            preferred_hotel_area=_str(row, "preferred_hotel_area"),
            hotel_star_rating=_str(row, "hotel_star_rating"),
            past_hotel_preference=_str(row, "past_hotel_preference"),
            airport_transfer_needed=_bool_value(row.get("airport_transfer_needed")),
            meal_preference=_str(row, "meal_preference"),
            seat_preference=_str(row, "seat_preference"),
            timing_preference=_str(row, "timing_preference"),
        ),
        budgets=TravelBudget(
            total_budget=_optional_int(row.get("total_budget") or row.get("budget_usd") or row.get("budget")),
            currency=_str(row, "currency") or origin_city_currency(str(travel_details.get("origin") or ""), "USD"),
            daily_budget=_optional_int(row.get("daily_budget") or row.get("per_day_budget") or row.get("daily_allowance")),
            flight_budget_per_hour=_optional_int(row.get("flight_budget_per_hour") or row.get("flight_hourly_budget")),
            estimated_flight_hours=_optional_int(row.get("estimated_flight_hours") or row.get("flight_hours")),
            max_flight_budget=_optional_int(row.get("max_flight_budget")),
            max_hotel_budget=_optional_int(row.get("max_hotel_budget")),
            policy_notes=_str(row, "policy_notes"),
            extra_baggage_notes=_str(row, "extra_baggage_notes"),
        ),
        special_requests=[item.strip() for item in str(special).replace(";", ",").split(",") if item and item.strip()],
        created_at=now,
        updated_at=now,
    )


def _company_name_from_request_row(row: dict[str, object]) -> str | None:
    return _str(row, "company_name") or _company_from_email(
        _str(row, "requester_email", "sender_email", "submitter_email", "traveler_email", "traveller_email", "email", "employee_email")
    )


def _header_name(value: object) -> str:
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower())).strip("_")


def _str(row: dict[str, object], *keys: str) -> str | None:
    for key in keys:
        value = row.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return None


def _date_value(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None or not str(value).strip():
        return None
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError:
        return None


def _optional_int(value: object) -> int | None:
    if value is None or not str(value).strip():
        return None
    return _int_value(value, default=0) or None


def _bool_value(value: object) -> bool:
    return str(value or "").strip().lower() in {"true", "yes", "y", "1"}


def _component_bool(row: dict[str, object], *keys: str, default: bool) -> bool:
    for key in keys:
        value = row.get(key)
        if value is None or not str(value).strip():
            continue
        normalized = str(value).strip().lower()
        if normalized in {"false", "no", "n", "0", "skip", "skipped", "not needed", "not required", "none", "exclude", "excluded"}:
            return False
        if normalized in {"true", "yes", "y", "1", "required", "needed", "include", "included"}:
            return True
    return default


def _int_value(value: object, default: int) -> int:
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return default


def _cabin_value(value: object) -> str:
    clean = str(value or "economy").strip().lower().replace(" ", "_")
    return clean if clean in {"economy", "premium_economy", "business", "first"} else "economy"
